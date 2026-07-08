# -*- coding: utf-8 -*-
"""
KUBIOS TODO-EN-UNO ESTABLE v41
--------------------------------
Novedades respecto a v20:
- Mantiene el flujo completo de extracción OCR, Excel acumulativo, gráficos y Word
- Añade un módulo de mensaje emocional basado exclusivamente en los datos longitudinales
- Aplica filtros mínimos de validez antes de emitir el mensaje emocional
- Genera un nivel de confianza del mensaje emocional
- Integra el mensaje emocional en el informe del participante
"""

import os
import re
import sys
import ast
import platform
import shutil
import subprocess
import traceback
import unicodedata
from pathlib import Path

import cv2
import numpy as np
import pandas as pd
import pytesseract
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


from config import (
    PARTICIPANTES_DIR_NAMES,
    IMAGE_EXTENSIONS,
    DATE_PATTERNS,
    SESSION_GROUP_GAP_SECONDS,
    RESULTADOS_DIR_NAME,
    ASSETS_DIR_NAME,
    SHARED_IMAGES_SUBDIR,
    PROJECT_PROGRESS,
    PROJECT_OVERALL_PROGRESS,
    PROJECT_PROGRESS_NOTES,
    PROJECT_VERSION,
    PROJECT_NEXT_OBJECTIVE,
)

# Módulo de participantes (v1.3.5)
# Import robusto para ejecución desde Windows/macOS/Linux.
try:
    PROJECT_ROOT_FOR_IMPORT = Path(__file__).resolve().parent
    if str(PROJECT_ROOT_FOR_IMPORT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT_FOR_IMPORT))
    from modules.participants import (
        show_participants_manager,
        participants_summary,
        selected_paths_from_active_selection,
        print_active_selection,
        load_active_selection,
    )
    PARTICIPANTS_IMPORT_ERROR = None
except Exception as exc:
    show_participants_manager = None
    participants_summary = None
    selected_paths_from_active_selection = None
    print_active_selection = None
    load_active_selection = None
    PARTICIPANTS_IMPORT_ERROR = str(exc)

# Scientific Validation Framework (v1.4.0-a)
try:
    from modules.validation import show_validation_menu, print_validation_dashboard, ensure_validation_structure
    VALIDATION_IMPORT_ERROR = None
except Exception as exc:
    show_validation_menu = None
    print_validation_dashboard = None
    ensure_validation_structure = None
    VALIDATION_IMPORT_ERROR = str(exc)

SCRIPT_DIR = Path(__file__).resolve().parent
BASE_DIR = SCRIPT_DIR
IMAGES_DIR = BASE_DIR
INFORMES_DIR = SCRIPT_DIR / RESULTADOS_DIR_NAME
SHARED_IMAGES_DIR = SCRIPT_DIR / ASSETS_DIR_NAME / SHARED_IMAGES_SUBDIR

def _resolver_ruta(path_obj):
    path_obj = Path(path_obj).expanduser()
    try:
        return path_obj.resolve()
    except Exception:
        return path_obj.absolute()


def resolver_directorio_participantes(root_dir):
    """Localiza la carpeta de datos de forma compatible Windows/macOS/Linux.

    Prioridad:
    1) Datos/
    2) datos/
    3) DATOS/
    4) carpeta raíz del proyecto
    """
    root_dir = _resolver_ruta(root_dir)

    for nombre in ("Datos", "datos", "DATOS"):
        datos_dir = _resolver_ruta(root_dir / nombre)
        if datos_dir.exists() and datos_dir.is_dir():
            return datos_dir

    return root_dir




def listar_carpetas_participantes(root_dir):
    root_dir = Path(root_dir)
    participantes = []
    excluidas = {
        "00_imagenes", "00 imagenes", "00-imagenes", "informes", "resultados", "imagenes", "imágenes",
        "assets", "excel_acumulado", "__pycache__", ".git", ".venv", "venv", "env",
    }

    try:
        items = sorted(root_dir.iterdir(), key=lambda p: p.name.lower())
    except Exception:
        items = []

    for p in items:
        if not p.is_dir():
            continue
        nombre = p.name.strip()
        nombre_norm = nombre.lower()
        if nombre_norm in excluidas:
            continue
        if nombre_norm.startswith(("00_", "00-", "00 ")):
            continue
        participantes.append(p)
    return participantes


def safe_simple_folder_name(texto):
    texto = str(texto or "sin_nombre").strip()
    texto = re.sub(r"^(imagenes?|imágenes?)[_\- ]+", "", texto, flags=re.IGNORECASE).strip()
    texto = re.sub(r"^\d+[_\- ]*", "", texto).strip()
    texto = re.sub(r"[^\w\- ]+", "", texto, flags=re.UNICODE)
    texto = re.sub(r"\s+", "_", texto)
    return texto or "sin_nombre"


def nombre_sujeto_desde_carpeta(carpeta):
    nombre = Path(carpeta).name.strip()
    nombre = re.sub(r"^(imagenes?|imágenes?)[_\- ]+", "", nombre, flags=re.IGNORECASE).strip()
    nombre = re.sub(r"^\d+[_\- ]*", "", nombre).strip()
    nombre = nombre.replace("_", " ").replace("-", " ")
    nombre = re.sub(r"\s+", " ", nombre).strip()
    return nombre or Path(carpeta).name.strip() or "Sin nombre"


def _norm_nombre_para_match(texto):
    texto = _strip_accents_robust(str(texto or "").lower()) if "_strip_accents_robust" in globals() else unicodedata.normalize("NFKD", str(texto or "").lower())
    if not isinstance(texto, str):
        texto = str(texto)
    texto = "".join(ch for ch in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(ch))
    texto = re.sub(r"^(imagenes?|imagenes?)[_\- ]+", "", texto, flags=re.IGNORECASE).strip()
    texto = re.sub(r"^\d+[_\- ]*", "", texto).strip()
    texto = re.sub(r"[^a-z0-9]+", "", texto)
    return texto


def obtener_carpetas_imagenes_participante(participant_dir):
    """Devuelve rutas candidatas para imágenes de un participante.

    Soporta estructuras como:
    datos/imagenes_judith/
    datos/Judith/imagenes_judith/
    datos/Judith/imagenes/
    """
    participant_dir = _resolver_ruta(participant_dir)
    nombre_participante = nombre_sujeto_desde_carpeta(participant_dir)
    nombre_norm = _norm_nombre_para_match(nombre_participante)
    candidatos = []
    vistos = set()

    def add(p):
        if p is None:
            return
        try:
            rp = Path(p).expanduser().resolve()
        except Exception:
            rp = Path(p).expanduser().absolute()
        key = str(rp).lower()
        if key not in vistos:
            vistos.add(key)
            candidatos.append(rp)

    # 1. La propia carpeta puede ser datos/imagenes_judith.
    add(participant_dir)

    # 2. Subcarpetas internas habituales.
    for nombre in ("imagenes", "imágenes", "Imagenes", "Imágenes", "images", "Images"):
        add(participant_dir / nombre)

    for child in sorted(participant_dir.iterdir(), key=lambda p: p.name.lower()) if participant_dir.exists() and participant_dir.is_dir() else []:
        if child.is_dir() and re.match(r"^(imagenes?|imágenes?|images)[_\- ]*", child.name, flags=re.IGNORECASE):
            add(child)

    # 3. Carpetas hermanas tipo datos/imagenes_judith cuando el participante es datos/Judith.
    parent = participant_dir.parent
    if parent.exists() and parent.is_dir():
        for child in sorted(parent.iterdir(), key=lambda p: p.name.lower()):
            if not child.is_dir():
                continue
            child_norm = _norm_nombre_para_match(child.name)
            child_name_low = child.name.lower()
            if child_name_low.startswith(("imagenes_", "imagenes-", "imagenes ", "imagen_", "imagen-", "imagen ", "imágenes_", "imágenes-", "imágenes ", "images_", "images-", "images ")):
                if not nombre_norm or nombre_norm in child_norm or child_norm in {"imagenes" + nombre_norm, "imagen" + nombre_norm, "images" + nombre_norm}:
                    add(child)

    return candidatos


def resolver_carpeta_imagenes_participante(participant_dir):
    """Elige la mejor carpeta de imágenes para el participante seleccionado."""
    candidatos = obtener_carpetas_imagenes_participante(participant_dir)
    mejor_dir = _resolver_ruta(participant_dir)
    mejor_files = []
    for cand in candidatos:
        files = get_image_files(cand, recursive=True) if "get_image_files" in globals() else []
        if len(files) > len(mejor_files):
            mejor_dir = cand
            mejor_files = files
    return mejor_dir, mejor_files, candidatos


def seleccionar_carpetas_participantes(root_dir):
    root_dir = _resolver_ruta(root_dir)
    directorio_participantes = resolver_directorio_participantes(root_dir)
    participantes = listar_carpetas_participantes(directorio_participantes)

    # v1.3.3: si existe una selección activa creada desde el gestor de participantes,
    # se puede reutilizar directamente para informes/análisis.
    if selected_paths_from_active_selection is not None:
        active_paths = selected_paths_from_active_selection(root_dir)
        active_paths = [p for p in active_paths if p in participantes]
        if active_paths:
            print("\nSe ha detectado una selección activa de participantes:")
            for idx, carpeta in enumerate(active_paths, start=1):
                print(f"{idx:>3}. {nombre_sujeto_desde_carpeta(carpeta)}")
            usar = input("\n¿Usar esta selección activa? [S/n]: ").strip().lower()
            if usar in {"", "s", "si", "sí", "y", "yes"}:
                return active_paths

    if not participantes:
        print("\nNo se han encontrado participantes.")
        print("\nCopia las carpetas de los participantes dentro de:")
        print(f"\n   {directorio_participantes}")
        print("\nEjemplo:")
        print("Datos/")
        print("   Juan/")
        print("   María/")
        print("   Pedro/")
        input("\nPulsa ENTER para salir...")
        raise SystemExit

    if directorio_participantes != root_dir:
        print(f"Carpeta de datos detectada automáticamente: {directorio_participantes}")

    print("\n" + "=" * 50)
    print(" PARTICIPANTES DETECTADOS")
    print("=" * 50)
    print(f"\nParticipantes encontrados: {len(participantes)}\n")

    for i, carpeta in enumerate(participantes, start=1):
        print(f"{i:>3}. {nombre_sujeto_desde_carpeta(carpeta)}")

    print("\n" + "-" * 50)
    print("  0. Todos los participantes")
    print("-" * 50)
    print("\nPuedes seleccionar:")
    print(" • Un participante      -> 1")
    print(" • Varios               -> 1,3,5")
    print(" • Intervalo            -> 5-12")
    print(" • Todos                -> 0")

    while True:
        entrada = input("\nSelección: ").strip().lower()

        if entrada == "":
            return [participantes[0]]

        if entrada in {"0", "todos", "todas", "all", "*"}:
            return participantes

        if re.fullmatch(r"\d+\s*-\s*\d+", entrada):
            inicio_txt, fin_txt = re.split(r"\s*-\s*", entrada)
            inicio = int(inicio_txt)
            fin = int(fin_txt)

            if 1 <= inicio <= fin <= len(participantes):
                return participantes[inicio - 1:fin]

            print("Intervalo no válido. Ejemplo: 5-12")
            continue

        partes = [x.strip() for x in re.split(r"[,\s;]+", entrada) if x.strip()]

        if not partes:
            print("Entrada no válida.")
            continue

        indices = []
        valido = True

        for parte in partes:
            if not re.fullmatch(r"\d+", parte):
                valido = False
                break

            idx = int(parte)

            if not (1 <= idx <= len(participantes)):
                valido = False
                break

            if idx not in indices:
                indices.append(idx)

        if not valido:
            print("Entrada no válida.")
            continue

        return [participantes[i - 1] for i in indices]


def configurar_directorios_base(participant_dir):
    global BASE_DIR, IMAGES_DIR, INFORMES_DIR

    BASE_DIR = _resolver_ruta(participant_dir)
    imagenes_detectadas_dir, _, _ = resolver_carpeta_imagenes_participante(BASE_DIR)
    IMAGES_DIR = imagenes_detectadas_dir
    INFORMES_DIR = SCRIPT_DIR / RESULTADOS_DIR_NAME / safe_simple_folder_name(nombre_sujeto_desde_carpeta(BASE_DIR))
    INFORMES_DIR.mkdir(parents=True, exist_ok=True)

    return BASE_DIR


# ============================================================
# BASE ROBUSTA DE RECONOCIMIENTO OCR KUBIOS - V32
# Sustituye la base de reconocimiento textual sin alterar el flujo
# general del script: carpetas, menús, Excel, Word, gráficos e informes.
# ============================================================

KUBIOS_RECOGNITION_VERSION = "V32_BASE_ROBUSTA_VARIABLES_KUBIOS"

ROBUST_VARIABLE_MAP = {
    "RR_MEAN": {
        "labels": ["Mean RR", "Mean R R", "Mean R-R", "RR mean", "Average RR", "AvRR", "RR medio"],
        "range": (250, 2000),
    },
    "SDNN": {
        "labels": ["SDNN", "SDNN ms", "SDNN (ms)"],
        "range": (0, 500),
    },
    "SD1": {
        "labels": ["Poincaré SD1", "Poincare SD1", "Poincaré SD1", "SD1"],
        "range": (0, 350),
    },
    "SD2": {
        "labels": ["Poincaré SD2", "Poincare SD2", "Poincaré SD2", "SD2"],
        "range": (0, 700),
    },
    "STRESS_INDEX": {
        "labels": ["Stress index", "Stress Index", "Stress", "Índice de estrés", "Indice de estres"],
        "range": (0, 3000),
    },
    "BREATH_RATE": {
        "labels": ["Respiratory rate", "Respiration rate", "Breathing rate", "Resp rate", "Frecuencia respiratoria"],
        "range": (3, 60),
    },
    "LF_POWER": {
        "labels": ["LF power", "LF power ms2", "LF power ms²", "LF ms2", "LF ms²"],
        "range": (0, 300000),
    },
    "HF_POWER": {
        "labels": ["HF power", "HF power ms2", "HF power ms²", "HF ms2", "HF ms²"],
        "range": (0, 300000),
    },
    "LF_NU": {
        "labels": ["LF power (n.u.)", "LF power n.u.", "LF power nu", "LF n.u.", "LF nu"],
        "range": (0, 100),
    },
    "HF_NU": {
        "labels": ["HF power (n.u.)", "HF power n.u.", "HF power nu", "HF n.u.", "HF nu"],
        "range": (0, 100),
    },
    "LF_HF": {
        "labels": ["LF/HF ratio", "LF HF ratio", "LF/HF", "LF-HF", "LF HF"],
        "range": (0, 100),
    },
    "RMSSD": {
        "labels": ["RMSSD", "RMSSD ms", "RMSSD (ms)"],
        "range": (0, 500),
    },
    "HR": {
        "labels": ["Heartrate", "Heart rate", "HR", "Pulse", "Frecuencia cardiaca"],
        "range": (25, 220),
    },
    "READINESS": {
        "labels": ["Today readiness", "Readiness"],
        "range": (0, 100),
    },
}

CORE_OCR_VARIABLES = [
    "RR_MEAN", "SDNN", "SD1", "SD2", "STRESS_INDEX", "BREATH_RATE",
    "LF_POWER", "HF_POWER", "LF_NU", "HF_NU", "LF_HF",
]


def _strip_accents_robust(txt):
    txt = str(txt)
    txt = unicodedata.normalize("NFKD", txt)
    txt = "".join(ch for ch in txt if not unicodedata.combining(ch))
    return txt


def normalize_ocr_for_matching(txt):
    txt = _strip_accents_robust(txt).lower()
    replacements = {
        "|": " ", "[": " ", "]": " ", "{": " ", "}": " ",
        "—": "-", "–": "-", "·": ".", "•": ".", ",": ".",
        "ms²": "ms2", "ms^2": "ms2", "m s2": "ms2",
        "n u": "nu", "n.u": "nu", "n.u.": "nu",
        "poincaré": "poincare", "poincaré": "poincare",
        "heartrate": "heart rate",
    }
    for a, b in replacements.items():
        txt = txt.replace(a, b)
    txt = re.sub(r"(?<=\d)\s+(?=\.\d)", "", txt)
    txt = re.sub(r"(?<=\d)\s+(?=\d)", " ", txt)
    txt = re.sub(r"[ \t]+", " ", txt)
    return txt


def normalize_label_for_matching(txt):
    txt = normalize_ocr_for_matching(txt)
    txt = txt.replace("/", " ").replace("-", " ").replace(".", " ")
    txt = re.sub(r"[^a-z0-9]+", " ", txt)
    txt = re.sub(r"\s+", " ", txt).strip()
    return txt


def _numbers_from_text_robust(txt):
    nums = []
    for m in re.finditer(r"[-+]?\d+(?:[\.,]\d+)?", str(txt)):
        val = clean_num(m.group(0))
        if val is not None:
            nums.append((m.start(), val, m.group(0)))
    return nums


def _valid_range(key, value):
    if value is None:
        return False
    lo, hi = ROBUST_VARIABLE_MAP.get(key, {}).get("range", (-np.inf, np.inf))
    try:
        return lo <= float(value) <= hi
    except Exception:
        return False


def _choose_best_candidate(key, candidates):
    valid = [c for c in candidates if _valid_range(key, c["value"])]
    if not valid:
        return None
    valid = sorted(valid, key=lambda c: (c.get("score", 999), c.get("distance", 999)))
    return valid[0]["value"]


def extract_metric_from_lines_robust(text, key):
    """Extrae un valor numérico buscando etiquetas equivalentes en líneas cercanas.
    Diseñado para tolerar errores OCR, acentos, n.u., ms² y separadores variables.
    """
    spec = ROBUST_VARIABLE_MAP.get(key)
    if not spec:
        return None

    raw_lines = [ln.strip() for ln in str(text).splitlines() if ln and ln.strip()]
    norm_lines = [normalize_ocr_for_matching(ln) for ln in raw_lines]
    labels = [(label, normalize_label_for_matching(label)) for label in spec["labels"]]
    candidates = []

    for i, line in enumerate(norm_lines):
        line_label = normalize_label_for_matching(line)
        for raw_label, norm_label in labels:
            if not norm_label:
                continue
            found = norm_label in line_label
            fuzzy = False
            if not found:
                # Coincidencia aproximada para errores típicos de OCR en etiquetas cortas/largas.
                tokens_label = set(norm_label.split())
                tokens_line = set(line_label.split())
                overlap = len(tokens_label & tokens_line)
                fuzzy = overlap >= max(1, min(len(tokens_label), 2)) and any(t in tokens_line for t in tokens_label)
            if not found and not fuzzy:
                continue

            window_lines = [line]
            if i + 1 < len(norm_lines):
                window_lines.append(norm_lines[i + 1])
            if i + 2 < len(norm_lines):
                window_lines.append(norm_lines[i + 2])
            window = "  ".join(window_lines)

            label_pos = normalize_label_for_matching(window).find(norm_label)
            nums = _numbers_from_text_robust(window)
            for pos, val, raw_num in nums:
                # Evitar coger números de la propia etiqueta: SD1, SD2, LF/HF, n.u.
                label_words = norm_label.split()
                if key == "SD1" and str(raw_num).strip() == "1":
                    continue
                if key == "SD2" and str(raw_num).strip() == "2":
                    continue
                if key in {"LF_HF", "LF_NU", "HF_NU"} and val in {1, 2} and raw_num in {"1", "2"}:
                    continue
                distance = abs(pos - label_pos) if label_pos >= 0 else pos
                score = distance + (0 if found else 20)
                candidates.append({"value": val, "distance": distance, "score": score, "line": raw_lines[i]})

    return _choose_best_candidate(key, candidates)


def extract_metric_by_regex_robust(text, key):
    norm = normalize_ocr_for_matching(text)
    specs = ROBUST_VARIABLE_MAP.get(key, {})
    candidates = []
    for label in specs.get("labels", []):
        lab = normalize_label_for_matching(label)
        if not lab:
            continue
        lab_pattern = r"\s*".join(re.escape(part) for part in lab.split())
        pattern = rf"{lab_pattern}[^0-9\-+]{{0,35}}([-+]?\d+(?:[\.,]\d+)?)"
        for m in re.finditer(pattern, norm, flags=re.IGNORECASE):
            val = clean_num(m.group(1))
            if val is not None:
                candidates.append({"value": val, "distance": 0, "score": 0})
    return _choose_best_candidate(key, candidates)


def extract_metric_robust(text, key):
    value = extract_metric_from_lines_robust(text, key)
    if value is not None:
        return value
    return extract_metric_by_regex_robust(text, key)


def build_ocr_variants(img):
    """Variantes rápidas de OCR global. No analiza burbujas ni añade flujo nuevo."""
    variants = []
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    variants.append(gray)
    try:
        variants.append(cv2.resize(gray, None, fx=2, fy=2, interpolation=cv2.INTER_CUBIC))
        variants.append(cv2.threshold(cv2.resize(gray, None, fx=2, fy=2, interpolation=cv2.INTER_CUBIC), 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1])
        clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(gray)
        variants.append(cv2.resize(clahe, None, fx=2, fy=2, interpolation=cv2.INTER_CUBIC))
    except Exception:
        pass
    return variants


def extract_global_text_robust(img):
    texts = []
    for variant in build_ocr_variants(img):
        for psm in (6, 4):
            try:
                txt = ocr_text(variant, psm=psm)
                if txt and txt.strip():
                    texts.append(txt)
            except Exception:
                continue
    if not texts:
        return extract_global_text(img)
    # Unir variantes aumenta la probabilidad de que una etiqueta salga bien sin cambiar el resto del pipeline.
    return "\n".join(dict.fromkeys(texts))


def robust_metric_pack_from_image(img):
    text = extract_global_text_robust(img)
    metrics = {}
    for key in ROBUST_VARIABLE_MAP:
        metrics[key] = extract_metric_robust(text, key)
    return metrics, text


APA_REFERENCES = {
    "task_force_1996": "Task Force of the European Society of Cardiology and the North American Society of Pacing and Electrophysiology. (1996). Heart rate variability: Standards of measurement, physiological interpretation and clinical use. European Heart Journal, 17(3), 354–381.",
    "shaffer_ginsberg_2017": "Shaffer, F., & Ginsberg, J. P. (2017). An overview of heart rate variability metrics and norms. Frontiers in Public Health, 5, 258.",
    "plews_2013": "Plews, D. J., Laursen, P. B., Stanley, J., Kilding, A. E., & Buchheit, M. (2013). Training adaptation and heart rate variability in elite endurance athletes: Opening the door to effective monitoring. Sports Medicine, 43(9), 773–781.",
    "stanley_2013": "Stanley, J., Peake, J. M., & Buchheit, M. (2013). Cardiac parasympathetic reactivation following exercise: Implications for training prescription. Sports Medicine, 43(12), 1259–1277.",
    "michael_2017": "Michael, S., Graham, K. S., & Davis, G. M. O. (2017). Cardiac autonomic responses during exercise and post-exercise recovery using heart rate variability and systolic time intervals—A review. Frontiers in Physiology, 8, 301.",
}

VARIABLE_META = {
    "HR": {
        "nombre": "Frecuencia cardiaca (HR)",
        "explicacion": "La frecuencia cardiaca en reposo es una variable complementaria a la variabilidad cardiaca y permite contextualizar el estado de activación fisiológica del organismo. Su interpretación mejora cuando se combina con variables vagales como RMSSD y con índices de estrés. (Task Force, 1996; Michael et al., 2017)",
        "interpretacion": "Valores más altos de lo habitual en reposo pueden ser compatibles con mayor activación fisiológica, fatiga acumulada, estrés o recuperación insuficiente. Valores más bajos, si el contexto general es estable, suelen ser compatibles con una recuperación más favorable.",
        "recomendaciones": "Si la HR aparece sostenidamente elevada junto con una reducción de RMSSD o PNS, conviene revisar la carga de entrenamiento, la calidad del sueño, el estado de hidratación y la presencia de estrés no relacionado con el ejercicio.",
        "refs": ["task_force_1996", "michael_2017"],
    },
    "RR_MEAN": {
        "nombre": "RR medio",
        "explicacion": "El RR medio representa la duración media de los intervalos entre latidos normales consecutivos y es la expresión temporal directa de la frecuencia cardiaca. Cuanto mayor es el RR medio, menor suele ser la frecuencia cardiaca. (Task Force, 1996)",
        "interpretacion": "Su utilidad es principalmente contextual. Debe interpretarse junto con HR y con variables de HRV, ya que por sí solo no describe de manera completa el estado autonómico.",
        "recomendaciones": "Cambios llamativos en RR medio deben revisarse junto con HR, RMSSD y la calidad de la medición antes de extraer conclusiones aplicadas.",
        "refs": ["task_force_1996"],
    },
    "RMSSD": {
        "nombre": "RMSSD",
        "explicacion": "El RMSSD es uno de los índices de HRV de corta duración más utilizados para monitorizar la actividad vagal cardiaca y la recuperación. Su uso longitudinal está ampliamente extendido en deporte y monitorización diaria. (Task Force, 1996; Shaffer & Ginsberg, 2017; Plews et al., 2013)",
        "interpretacion": "Valores relativamente más altos suelen asociarse con mayor influencia parasimpática y mejor recuperación. Valores bajos o una tendencia descendente mantenida pueden ser compatibles con fatiga, estrés fisiológico o menor capacidad de recuperación.",
        "recomendaciones": "Cuando RMSSD cae varios días seguidos y coincide con mayor HR o peor sensación subjetiva, suele ser recomendable reducir o modular la carga, priorizar el sueño y revisar el contexto general de recuperación.",
        "refs": ["task_force_1996", "shaffer_ginsberg_2017", "plews_2013"],
    },
    "LnRMSSD": {
        "nombre": "LnRMSSD",
        "explicacion": "El LnRMSSD es la transformación logarítmica natural del RMSSD y se utiliza con frecuencia en el seguimiento longitudinal porque reduce la asimetría de la distribución y facilita una lectura más estable de la respuesta autonómica. (Plews et al., 2013; Stanley et al., 2013)",
        "interpretacion": "Una tendencia ascendente y estable suele ser compatible con una adaptación favorable. Una tendencia descendente mantenida puede sugerir reducción del tono vagal o peor recuperación.",
        "recomendaciones": "El LnRMSSD es especialmente útil para el control diario. Conviene valorarlo junto con HR, PNS y SNS para interpretar la carga y la recuperación.",
        "refs": ["plews_2013", "stanley_2013"],
    },
    "SDNN": {
        "nombre": "SDNN",
        "explicacion": "La SDNN expresa la desviación estándar de los intervalos NN y refleja variabilidad global de la señal cardiaca. Es un índice clásico de HRV que resume la variabilidad total del registro. (Task Force, 1996; Shaffer & Ginsberg, 2017)",
        "interpretacion": "Valores más altos suelen indicar mayor variabilidad cardiaca global. Valores bajos pueden reflejar menor variabilidad total y deben interpretarse dentro del contexto de HR, RMSSD y condiciones de registro.",
        "recomendaciones": "La SDNN resulta más útil cuando se valora de forma longitudinal. Conviene evitar conclusiones basadas en un único valor aislado.",
        "refs": ["task_force_1996", "shaffer_ginsberg_2017"],
    },
    "SD1": {
        "nombre": "SD1",
        "explicacion": "El SD1 es un descriptor del gráfico de Poincaré y representa la variabilidad de corto plazo del sistema cardiaco. Se relaciona estrechamente con el componente vagal de la modulación autonómica. (Task Force, 1996; Shaffer & Ginsberg, 2017)",
        "interpretacion": "Valores más altos suelen ser compatibles con mayor variabilidad de corto plazo y mejor modulación vagal. Valores más bajos pueden sugerir menor flexibilidad autonómica inmediata.",
        "recomendaciones": "Su lectura gana valor cuando se interpreta junto con RMSSD, SD2 y la evolución temporal entre sesiones.",
        "refs": ["task_force_1996", "shaffer_ginsberg_2017"],
    },
    "SD2": {
        "nombre": "SD2",
        "explicacion": "El SD2 es otro descriptor del gráfico de Poincaré y representa la variabilidad a más largo plazo. Aporta información complementaria a SD1 y ayuda a describir la estructura global de la señal. (Task Force, 1996; Shaffer & Ginsberg, 2017)",
        "interpretacion": "Su interpretación aislada es limitada, pero combinada con SD1 puede aportar información sobre el equilibrio entre variabilidad de corto y de largo plazo.",
        "recomendaciones": "Es recomendable valorarlo siempre junto con SD1 y el patrón general del resto de variables de HRV.",
        "refs": ["task_force_1996", "shaffer_ginsberg_2017"],
    },
    "STRESS_INDEX": {
        "nombre": "Índice de estrés",
        "explicacion": "El índice de estrés es un descriptor compuesto de activación fisiológica y puede complementar la interpretación de HR y HRV cuando se analiza de forma repetida en el tiempo. (Shaffer & Ginsberg, 2017; Michael et al., 2017)",
        "interpretacion": "Valores elevados o una tendencia creciente pueden ser compatibles con mayor carga interna o activación fisiológica sostenida.",
        "recomendaciones": "Si el índice de estrés sube de forma mantenida y coincide con peor recuperación subjetiva, es razonable revisar la carga externa y el descanso.",
        "refs": ["shaffer_ginsberg_2017", "michael_2017"],
    },
    "LF_HF": {
        "nombre": "Relación LF/HF",
        "explicacion": "La relación LF/HF se ha utilizado tradicionalmente como un descriptor del balance autonómico, aunque su interpretación debe hacerse con prudencia y nunca de forma aislada. (Task Force, 1996; Shaffer & Ginsberg, 2017)",
        "interpretacion": "Valores relativamente altos pueden sugerir mayor activación simpática relativa, pero siempre deben contextualizarse con HR, RMSSD, PNS, SNS y las condiciones del registro.",
        "recomendaciones": "No conviene tomar decisiones aplicadas usando solo esta variable. Debe integrarse con el resto del patrón autonómico y con la evolución temporal.",
        "refs": ["task_force_1996", "shaffer_ginsberg_2017"],
    },
    "PNS_INDEX": {
        "nombre": "Índice PNS",
        "explicacion": "El índice PNS resume de forma compuesta aspectos vinculados con la actividad parasimpática y puede facilitar la lectura longitudinal del estado vagal. (Shaffer & Ginsberg, 2017; Plews et al., 2013)",
        "interpretacion": "Valores más altos suelen ser compatibles con mayor influencia parasimpática y mejor estado de recuperación autonómica.",
        "recomendaciones": "Cuando el PNS disminuye de forma sostenida, conviene revisar la recuperación, el sueño y la acumulación de carga.",
        "refs": ["shaffer_ginsberg_2017", "plews_2013"],
    },
    "SNS_INDEX": {
        "nombre": "Índice SNS",
        "explicacion": "El índice SNS resume de forma compuesta aspectos compatibles con activación simpática relativa y añade contexto cuando se interpreta con HR, RMSSD y el índice de estrés. (Shaffer & Ginsberg, 2017; Michael et al., 2017)",
        "interpretacion": "Valores altos pueden asociarse con mayor activación fisiológica o mayor carga acumulada, sobre todo si se acompañan de HR más elevada y descenso de variables vagales.",
        "recomendaciones": "Una elevación mantenida del SNS sugiere monitorizar de cerca el estado de recuperación y, si procede, ajustar el entrenamiento o el estrés general.",
        "refs": ["shaffer_ginsberg_2017", "michael_2017"],
    },
    "PHYSIO_AGE": {
        "nombre": "Edad fisiológica",
        "explicacion": "La edad fisiológica es un indicador estimado por la aplicación a partir del patrón autonómico registrado y debe interpretarse como una referencia orientativa, no como un marcador diagnóstico. (Shaffer & Ginsberg, 2017)",
        "interpretacion": "Su valor aislado tiene utilidad limitada. Resulta más razonable interpretarlo como un complemento descriptivo del perfil autonómico general.",
        "recomendaciones": "No conviene usar esta variable de forma aislada para tomar decisiones clínicas o de entrenamiento.",
        "refs": ["shaffer_ginsberg_2017"],
    },
}

SHORT_VARIABLE_INFO = {
    "HR": "Frecuencia cardiaca en reposo. Ayuda a contextualizar el nivel general de activación fisiológica.",
    "RR_MEAN": "Tiempo medio entre latidos normales consecutivos. Es la expresión temporal de la frecuencia cardiaca.",
    "RMSSD": "Indicador clave de la variabilidad de corto plazo y de la recuperación vagal cardiaca.",
    "LnRMSSD": "Versión logarítmica del RMSSD, útil para un seguimiento longitudinal más estable.",
    "SDNN": "Índice de variabilidad cardiaca global que resume la variabilidad total del registro.",
    "SD1": "Descriptor del gráfico de Poincaré asociado a la variabilidad de corto plazo.",
    "SD2": "Descriptor del gráfico de Poincaré asociado a la variabilidad de más largo plazo.",
    "STRESS_INDEX": "Indicador compuesto de activación fisiológica o carga interna.",
    "LF_HF": "Relación entre componentes espectrales; debe interpretarse con mucha prudencia y nunca de forma aislada.",
    "PNS_INDEX": "Índice compuesto relacionado con la actividad parasimpática y la recuperación autonómica.",
    "SNS_INDEX": "Índice compuesto relacionado con activación simpática relativa.",
    "PHYSIO_AGE": "Estimación orientativa basada en el patrón autonómico; no debe interpretarse como un marcador diagnóstico.",
}

SIMPLE_INTERPRETATION = {
    "HR": "En términos sencillos, si está más alta de lo habitual puede reflejar mayor activación o menor recuperación; si está más baja y estable, suele asociarse a una recuperación más favorable.",
    "RR_MEAN": "En términos sencillos, valores más altos suelen acompañarse de una frecuencia cardiaca más baja, algo que en muchos casos es compatible con un estado más relajado.",
    "RMSSD": "En términos sencillos, valores más altos suelen asociarse a mejor capacidad de recuperación; valores más bajos pueden aparecer cuando el organismo está más exigido o cansado.",
    "LnRMSSD": "En términos sencillos, ayuda a ver mejor la evolución en el tiempo: si aumenta de forma progresiva, suele ser una señal positiva de adaptación.",
    "SDNN": "En términos sencillos, valores más altos suelen reflejar mayor variabilidad global; valores bajos pueden indicar menor flexibilidad del sistema y deben leerse con prudencia.",
    "SD1": "En términos sencillos, valores más altos suelen relacionarse con mejor variabilidad de corto plazo y una respuesta vagal más favorable.",
    "SD2": "En términos sencillos, aporta contexto sobre la variabilidad a más largo plazo y conviene interpretarlo junto con SD1 y el resto de variables.",
    "STRESS_INDEX": "En términos sencillos, valores elevados pueden reflejar mayor carga o estrés fisiológico; si desciende con el tiempo, suele ser una señal de mejor recuperación.",
    "LF_HF": "En términos sencillos, puede sugerir cambios en el equilibrio autonómico, pero nunca debería interpretarse por sí sola.",
    "PNS_INDEX": "En términos sencillos, valores más altos suelen ser compatibles con un mejor estado de recuperación autonómica.",
    "SNS_INDEX": "En términos sencillos, valores más altos pueden reflejar mayor activación fisiológica o mayor carga acumulada.",
    "PHYSIO_AGE": "En términos sencillos, es una referencia orientativa del perfil autonómico, no una medida real de la edad ni un diagnóstico.",
}


plt.rcParams["figure.dpi"] = 170
plt.rcParams["savefig.dpi"] = 200
plt.rcParams["font.size"] = 10


def set_tesseract():
    candidates = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]
    for p in candidates:
        if os.path.exists(p):
            pytesseract.pytesseract.tesseract_cmd = p
            return p
    return getattr(pytesseract.pytesseract, "tesseract_cmd", None)


def ocr_text(img, psm=6, whitelist=None):
    cfg = f"--psm {psm}"
    if whitelist:
        cfg += f" -c tessedit_char_whitelist={whitelist}"
    return pytesseract.image_to_string(img, config=cfg)


def clean_num(s):
    if s is None:
        return None
    s = str(s).strip().replace(",", ".")
    s = re.sub(r"[^0-9\.\-]", "", s)
    if not s or s in {"-", ".", "-."}:
        return None
    try:
        return float(s)
    except Exception:
        return None


def detect_bubbles(img):
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    mask = (gray > 220).astype("uint8")
    n, labels, stats, cent = cv2.connectedComponentsWithStats(mask, 8)
    bubbles = []
    for i in range(1, n):
        x, y, w, h, area = stats[i]
        if 2500 <= area <= 7000 and 90 <= w <= 140 and 60 <= h <= 95:
            if y < 300:
                continue
            bubbles.append((x, y, w, h, area))
    return sorted(bubbles, key=lambda t: t[1])[:3]


def bubble_ocr(img, bubble, top_px, threshold, scale):
    x, y, w, h, _ = bubble
    crop = img[y:y + h, x:x + w].copy()
    crop = crop[:top_px, :]
    gray = cv2.cvtColor(crop, cv2.COLOR_BGR2GRAY)
    th = cv2.threshold(gray, threshold, 255, cv2.THRESH_BINARY)[1]
    big = cv2.resize(th, None, fx=scale, fy=scale, interpolation=cv2.INTER_CUBIC)
    txt = ocr_text(big, psm=7, whitelist="-0123456789.")
    return clean_num(txt), txt


def read_bubble_value(img, bubble, kind):
    if kind == "PNS_INDEX":
        val, raw = bubble_ocr(img, bubble, top_px=42, threshold=150, scale=10)
        if val is not None:
            return float(int(val * 10) / 10), raw
        val, raw = bubble_ocr(img, bubble, top_px=45, threshold=170, scale=10)
        return (float(int(val * 10) / 10) if val is not None else None), raw
    if kind == "SNS_INDEX":
        val, raw = bubble_ocr(img, bubble, top_px=38, threshold=150, scale=10)
        if val is not None:
            return round(val, 2), raw
        val, raw = bubble_ocr(img, bubble, top_px=40, threshold=150, scale=10)
        return (round(val, 2) if val is not None else None), raw
    if kind == "PHYSIO_AGE":
        val, raw = bubble_ocr(img, bubble, top_px=34, threshold=150, scale=8)
        if val is not None:
            return int(round(val)), raw
        val, raw = bubble_ocr(img, bubble, top_px=38, threshold=150, scale=8)
        return (int(round(val)) if val is not None else None), raw
    return None, None


def extract_global_text(img):
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    big = cv2.resize(gray, None, fx=2, fy=2, interpolation=cv2.INTER_CUBIC)
    return ocr_text(big, psm=6)


def extract_by_patterns(text, patterns):
    for pattern in patterns:
        m = re.search(pattern, text, re.IGNORECASE)
        if m:
            return clean_num(m.group(1))
    return None


def extract_quality(text):
    m = re.search(r"MEASUREMENT QUALITY[:\s]+([A-Z]+)", text, re.IGNORECASE)
    return m.group(1).upper() if m else None


def normalize_date(date_str):
    if not date_str:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return pd.to_datetime(date_str, format=fmt).strftime("%Y-%m-%d")
        except Exception:
            pass
    try:
        return pd.to_datetime(date_str).strftime("%Y-%m-%d")
    except Exception:
        return None


def normalize_time(time_str):
    if not time_str:
        return None
    t = str(time_str).replace(".", ":").replace("-", ":")
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return pd.to_datetime(t, format=fmt).strftime("%H:%M:%S")
        except Exception:
            pass
    try:
        return pd.to_datetime(t).strftime("%H:%M:%S")
    except Exception:
        return None


def extract_datetime_from_filename(filename):
    stem = Path(filename).stem
    for pattern in DATE_PATTERNS:
        m = pattern.search(stem)
        if m:
            date_value = normalize_date(m.groupdict().get("date"))
            time_value = normalize_time(m.groupdict().get("time"))
            return date_value, time_value
    return None, None


def imread_unicode(path):
    path = Path(path)
    try:
        data = np.fromfile(str(path), dtype=np.uint8)
        if data.size == 0:
            return None
        return cv2.imdecode(data, cv2.IMREAD_COLOR)
    except Exception:
        return None




# ============================================================
# DETECCIÓN V33 MODIFICADA CON LÓGICA KUBIOS V14
# Solo sustituye el motor de lectura OCR. No cambia menús, Excel,
# Word, gráficos, acumulativos ni estructura general de la V33.
# Principio: modo estricto anti-invención, sin lectura visual de
# burbujas/badges; PNS/SNS/edad fisiológica solo si aparecen como
# texto explícito en la imagen.
# ============================================================

V33_DETECTION_ENGINE = "V33_DETECCION_KUBIOS_V14_ESTRICTA_SIN_BURBUJAS"
V33_REQUIRE_UNIT_FOR_MS_VARIABLES = True
V33_STRICT_NO_GUESS_MODE = True
V33_MIN_SINGLE_CANDIDATE_CONF = 0.74

V33_V14_VARIABLES = [
    "Heart rate", "RMSSD", "PNS index", "SNS index", "Physiological age",
    "Mean RR", "SDNN", "Poincare SD1", "Poincare SD2", "Stress index", "Respiratory rate",
    "LF power", "HF power", "LF power (n.u.)", "HF power (n.u.)", "LF/HF ratio"
]

V33_V14_TO_INTERNAL = {
    "Heart rate": "HR",
    "RMSSD": "RMSSD",
    "PNS index": "PNS_INDEX",
    "SNS index": "SNS_INDEX",
    "Physiological age": "PHYSIO_AGE",
    "Mean RR": "RR_MEAN",
    "SDNN": "SDNN",
    "Poincare SD1": "SD1",
    "Poincare SD2": "SD2",
    "Stress index": "STRESS_INDEX",
    "Respiratory rate": "BREATH_RATE",
    "LF power": "LF_POWER",
    "HF power": "HF_POWER",
    "LF power (n.u.)": "LF_NU",
    "HF power (n.u.)": "HF_NU",
    "LF/HF ratio": "LF_HF",
}

V33_V14_PLAUSIBLE = {
    "Heart rate": (30, 220),
    "RMSSD": (0, 300),
    "PNS index": (-10, 10),
    "SNS index": (-10, 10),
    "Physiological age": (1, 120),
    "Mean RR": (250, 2000),
    "SDNN": (0, 300),
    "Poincare SD1": (0, 300),
    "Poincare SD2": (0, 500),
    "Stress index": (0, 1000),
    "Respiratory rate": (3, 45),
    "LF power": (0, 100000),
    "HF power": (0, 100000),
    "LF power (n.u.)": (0, 100),
    "HF power (n.u.)": (0, 100),
    "LF/HF ratio": (0, 50),
}

V33_V14_PATTERNS = {
    "Heart rate": [r"Heart\s*rate\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)\s*bpm", r"\bHR\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)\s*bpm"],
    "RMSSD": [r"RMSSD\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)\s*ms"],
    "PNS index": [r"PNS\s*index\s*[:=]?\s*(-?[0-9]+(?:[\.,][0-9]+)?)"],
    "SNS index": [r"SNS\s*index\s*[:=]?\s*(-?[0-9]+(?:[\.,][0-9]+)?)"],
    "Physiological age": [r"Physiological\s*age\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)"],
    "Mean RR": [r"Mean\s*RR\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)\s*ms"],
    "SDNN": [r"\bSDNN\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)\s*ms"],
    "Poincare SD1": [r"Poincar[eéè]?\s*SD1\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)\s*ms", r"Poincar\s*e?\s*SD1\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)\s*ms"],
    "Poincare SD2": [r"Poincar[eéè]?\s*SD2\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)\s*ms", r"Poincar\s*e?\s*SD2\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)\s*ms"],
    "Stress index": [r"Stress\s*index\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)"],
    "Respiratory rate": [r"Respiratory\s*rate\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)\s*breaths\s*/\s*min", r"Respiratory\s*rate\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)"],
    "LF power": [r"LF\s*power\s*(?!\(\s*n\.?u\.?)\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)\s*ms", r"LF\s*power\s*(?!\(\s*n\.?u\.?)\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)"],
    "HF power": [r"HF\s*power\s*(?!\(\s*n\.?u\.?)\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)\s*ms", r"HF\s*power\s*(?!\(\s*n\.?u\.?)\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)"],
    "LF power (n.u.)": [r"LF\s*power\s*\(\s*n\.?u\.?\s*\)\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)\s*%"],
    "HF power (n.u.)": [r"HF\s*power\s*\(\s*n\.?u\.?\s*\)\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)\s*%"],
    "LF/HF ratio": [r"LF\s*/\s*HF\s*ratio\s*[:=]?\s*([0-9]+(?:[\.,][0-9]+)?)"],
}

V33_V14_LABELS = {
    "Heart rate": ["Heart rate", "HR"],
    "RMSSD": ["RMSSD"],
    "PNS index": ["PNS index"],
    "SNS index": ["SNS index"],
    "Physiological age": ["Physiological age"],
    "Mean RR": ["Mean RR"],
    "SDNN": ["SDNN"],
    "Poincare SD1": ["Poincare SD1", "Poincaré SD1"],
    "Poincare SD2": ["Poincare SD2", "Poincaré SD2"],
    "Stress index": ["Stress index"],
    "Respiratory rate": ["Respiratory rate"],
    "LF power": ["LF power"],
    "HF power": ["HF power"],
    "LF power (n.u.)": ["LF power n.u", "LF power (n.u.)"],
    "HF power (n.u.)": ["HF power n.u", "HF power (n.u.)"],
    "LF/HF ratio": ["LF/HF ratio", "LF / HF ratio"],
}

def _v33_v14_normalize_text(t):
    t = unicodedata.normalize("NFKC", t or "")
    replacements = {
        "Poincaré": "Poincare", "Poincarè": "Poincare", "Poincaré": "Poincare",
        "S D N N": "SDNN", "R M S S D": "RMSSD", "R MSSD": "RMSSD",
        "LF / HF": "LF/HF", "LF/ HF": "LF/HF", "LF /HF": "LF/HF",
        "n. u.": "n.u.", "n u": "n.u.", "m5": "ms", "mS": "ms", "rn5": "ms",
        "Respiratoryrate": "Respiratory rate", "Stressindex": "Stress index",
    }
    for a, b in replacements.items():
        t = t.replace(a, b)
    t = re.sub(r"ms[²2]", "ms2", t, flags=re.I)
    t = re.sub(r"[^\S\r\n]+", " ", t)
    return t

def _v33_v14_resize_max(img, max_side=2200):
    h, w = img.shape[:2]
    m = max(h, w)
    if m <= max_side:
        return img
    scale = max_side / float(m)
    return cv2.resize(img, (int(w * scale), int(h * scale)), interpolation=cv2.INTER_AREA)

def _v33_v14_order_points(pts):
    pts = pts.reshape(4, 2).astype("float32")
    s = pts.sum(axis=1)
    diff = np.diff(pts, axis=1).ravel()
    rect = np.zeros((4, 2), dtype="float32")
    rect[0] = pts[np.argmin(s)]
    rect[2] = pts[np.argmax(s)]
    rect[1] = pts[np.argmin(diff)]
    rect[3] = pts[np.argmax(diff)]
    return rect

def _v33_v14_rectify_kubios_screen(img):
    img = _v33_v14_resize_max(img, 2200)
    h, w = img.shape[:2]
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    blur = cv2.GaussianBlur(gray, (5, 5), 0)
    edges = cv2.Canny(blur, 40, 120)
    edges = cv2.dilate(edges, np.ones((3, 3), np.uint8), iterations=1)
    contours, _ = cv2.findContours(edges, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    best, best_area = None, 0
    img_area = h * w
    for c in contours:
        area = cv2.contourArea(c)
        if area < img_area * 0.18 or area > img_area * 0.98:
            continue
        peri = cv2.arcLength(c, True)
        approx = cv2.approxPolyDP(c, 0.025 * peri, True)
        if len(approx) == 4:
            x, y, bw, bh = cv2.boundingRect(approx)
            ratio = bw / max(bh, 1)
            if 0.35 <= ratio <= 1.25 and area > best_area:
                best, best_area = approx, area
    if best is None:
        return img
    rect = _v33_v14_order_points(best)
    tl, tr, br, bl = rect
    max_w = int(max(np.linalg.norm(br - bl), np.linalg.norm(tr - tl)))
    max_h = int(max(np.linalg.norm(tr - br), np.linalg.norm(tl - bl)))
    if max_w < 300 or max_h < 500:
        return img
    dst = np.array([[0, 0], [max_w - 1, 0], [max_w - 1, max_h - 1], [0, max_h - 1]], dtype="float32")
    M = cv2.getPerspectiveTransform(rect, dst)
    return cv2.warpPerspective(img, M, (max_w, max_h))

def _v33_v14_prepare_text_image(img, scale=1.55):
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY) if img.ndim == 3 else img.copy()
    gray = cv2.resize(gray, None, fx=scale, fy=scale, interpolation=cv2.INTER_CUBIC)
    gray = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(gray)
    return gray

def _v33_v14_crop_hrv_parameters_area(img):
    h, w = img.shape[:2]
    return img[int(h * 0.34):h, 0:w]

def _v33_v14_extract_global_text(img):
    img = _v33_v14_rectify_kubios_screen(img)
    h, w = img.shape[:2]
    crops = [
        _v33_v14_crop_hrv_parameters_area(img),
        img[0:int(h * 0.42), 0:w],
    ]
    texts = []
    for crop in crops:
        if crop.size == 0:
            continue
        prep = _v33_v14_prepare_text_image(crop, 1.55)
        texts.append(ocr_text(prep, psm=6))
    return "\n".join(texts)

def _v33_v14_label_present_in_line(line, label_variants):
    from difflib import SequenceMatcher
    low = _v33_v14_normalize_text(line).lower()
    low_clean = re.sub(r"[^a-z0-9/(). ]", " ", low)
    low_clean = re.sub(r"\s+", " ", low_clean).strip()
    best_score, best_label = 0.0, ""
    for lab in label_variants:
        lab_norm = _v33_v14_normalize_text(lab).lower()
        lab_clean = re.sub(r"[^a-z0-9/(). ]", " ", lab_norm)
        lab_clean = re.sub(r"\s+", " ", lab_clean).strip()
        if not lab_clean:
            continue
        if lab_clean in low_clean:
            return True, 1.0, lab_clean
        if lab_clean in {"poincare sd1", "poincare sd2"}:
            continue
        if len(lab_clean) >= 8:
            tokens = low_clean.split()
            n = len(lab_clean.split())
            windows = [" ".join(tokens[i:i+n]) for i in range(max(1, len(tokens)-n+1))]
            for win in windows:
                score = SequenceMatcher(None, lab_clean, win).ratio()
                if score > best_score:
                    best_score, best_label = score, lab_clean
    return (best_score >= 0.91), best_score, best_label

def _v33_v14_extract_number_after_label(line, label_used, var):
    low = _v33_v14_normalize_text(line)
    low_search = re.sub(r"\s+", " ", low)
    label_rx = re.escape(label_used).replace("\\ ", r"\s+")
    m = re.search(label_rx, low_search, flags=re.I)
    segment = low_search[m.end():] if m else low_search
    stop_words = ["heart rate", "rmssd", "mean rr", "sdnn", "poincare", "stress index",
                  "respiratory rate", "lf power", "hf power", "lf/hf", "pns index", "sns index",
                  "physiological age", "measurement quality"]
    seg_low = segment.lower()
    stops = [seg_low.find(w) for w in stop_words if seg_low.find(w) > 0]
    if stops:
        segment = segment[:min(stops)]
    ms_vars = {"RMSSD", "Mean RR", "SDNN", "Poincare SD1", "Poincare SD2"}
    if V33_REQUIRE_UNIT_FOR_MS_VARIABLES and var in ms_vars and not re.search(r"\bms\b", segment, flags=re.I):
        return None
    nums = [clean_num(x) for x in re.findall(r"-?\d+(?:[\.,]\d+)?", segment)]
    lo, hi = V33_V14_PLAUSIBLE[var]
    nums = [float(x) for x in nums if x is not None and lo <= float(x) <= hi]
    if not nums:
        return None
    return round(float(nums[0]), 2)

def _v33_v14_best_line_value(lines, label_variants, var):
    norm_lines = [_v33_v14_normalize_text(x).strip() for x in lines if str(x).strip()]
    best = {"value": None, "confidence": 0.0, "method": "line_strict", "detail": ""}
    for i, line in enumerate(norm_lines):
        present, score, label_used = _v33_v14_label_present_in_line(line, label_variants)
        if not present:
            continue
        val = _v33_v14_extract_number_after_label(line, label_used, var)
        detail = line[:180]
        if val is None and i + 1 < len(norm_lines):
            nxt = norm_lines[i + 1].strip()
            if not any(_v33_v14_label_present_in_line(nxt, labs)[0] for labs in V33_V14_LABELS.values()):
                nums = [clean_num(x) for x in re.findall(r"-?\d+(?:[\.,]\d+)?", nxt)]
                lo, hi = V33_V14_PLAUSIBLE[var]
                nums = [float(x) for x in nums if x is not None and lo <= float(x) <= hi]
                if len(nums) == 1:
                    val = round(nums[0], 2)
                    detail = (line + " | " + nxt)[:180]
        if val is not None:
            conf = 0.82 if score >= 0.98 else 0.74
            if conf > best["confidence"]:
                best = {"value": val, "confidence": conf, "method": "line_strict", "detail": detail}
    return best

def _v33_v14_extract_by_patterns(text):
    t = _v33_v14_normalize_text(text)
    values = {v: None for v in V33_V14_VARIABLES}
    confs = {v: 0.0 for v in V33_V14_VARIABLES}
    audit = []
    for key, pats in V33_V14_PATTERNS.items():
        for pat in pats:
            m = re.search(pat, t, re.I)
            if m:
                val = clean_num(m.group(1))
                if val is not None and V33_V14_PLAUSIBLE[key][0] <= val <= V33_V14_PLAUSIBLE[key][1]:
                    values[key] = round(float(val), 2)
                    confs[key] = 0.91
                    audit.append((key, values[key], 0.91, "regex_explicit_label", pat))
                    break
    lines = t.splitlines()
    for var, labs in V33_V14_LABELS.items():
        if values.get(var) is None:
            source_lines = [ln for ln in lines if "n.u" not in ln.lower()] if var in {"LF power", "HF power"} else lines
            cand = _v33_v14_best_line_value(source_lines, labs, var)
            if cand["value"] is not None:
                values[var] = cand["value"]
                confs[var] = cand["confidence"]
                audit.append((var, cand["value"], cand["confidence"], cand["method"], cand["detail"]))
    return values, confs, audit

def _v33_v14_apply_consistency_guard(values, confs, audit):
    def f(name):
        try:
            return float(values.get(name)) if values.get(name) is not None else None
        except Exception:
            return None
    sdnn, sd1, sd2 = f("SDNN"), f("Poincare SD1"), f("Poincare SD2")
    reject_reasons = []
    if sd2 is not None:
        if sd1 is not None and sd2 < sd1:
            reject_reasons.append(f"SD2<{sd1:g} SD1")
        if sdnn is not None and sd2 < max(3.0, sdnn * 0.70):
            reject_reasons.append(f"SD2 incompatible con SDNN={sdnn:g}")
        if sd2 <= 3 and (sd1 is not None or sdnn is not None):
            reject_reasons.append("SD2 residual <=3")
    if reject_reasons:
        audit.append(("Poincare SD2", values.get("Poincare SD2"), 0.0, "rechazado_consistencia_fisiologica", "; ".join(reject_reasons)))
        values["Poincare SD2"] = None
        confs["Poincare SD2"] = 0.0

def _v33_v14_extract_measurement_quality(text):
    t = _v33_v14_normalize_text(text)
    m = re.search(r"MEASUREMENT\s*QUALITY\s*[:\s]*(EXCELLENT|GOOD|POOR|LOW|BAD|ACCEPTABLE|OK)", t, re.I)
    return m.group(1).upper() if m else None

def robust_metric_pack_from_image(img):
    text = _v33_v14_extract_global_text(img)
    values, confs, audit = _v33_v14_extract_by_patterns(text)
    _v33_v14_apply_consistency_guard(values, confs, audit)
    metrics = {internal: None for internal in V33_V14_TO_INTERNAL.values()}
    metrics["READINESS"] = extract_metric_robust(text, "READINESS")
    for var, internal in V33_V14_TO_INTERNAL.items():
        val = values.get(var)
        if val is not None and var == "Physiological age":
            val = int(round(float(val)))
        metrics[internal] = val
    metrics["_CONFIDENCE"] = {V33_V14_TO_INTERNAL.get(k, k): v for k, v in confs.items()}
    metrics["_AUDIT"] = audit
    metrics["_ENGINE"] = V33_DETECTION_ENGINE
    return metrics, text


def process_image(path):
    img = imread_unicode(path)
    if img is None:
        raise ValueError(f"No se pudo abrir la imagen: {path}")

    robust_metrics, text = robust_metric_pack_from_image(img)
    fecha_imagen, hora_imagen = extract_datetime_from_filename(Path(path).name)

    data = {
        "FILE": Path(path).name,
        "FECHA_IMAGEN": fecha_imagen,
        "HORA_IMAGEN": hora_imagen,
        "READINESS": robust_metrics.get("READINESS"),
        "HR": robust_metrics.get("HR"),
        "RMSSD": robust_metrics.get("RMSSD"),
        "RR_MEAN": robust_metrics.get("RR_MEAN"),
        "SDNN": robust_metrics.get("SDNN"),
        "SD1": robust_metrics.get("SD1"),
        "SD2": robust_metrics.get("SD2"),
        "STRESS_INDEX": robust_metrics.get("STRESS_INDEX"),
        "BREATH_RATE": robust_metrics.get("BREATH_RATE"),
        "LF_POWER": robust_metrics.get("LF_POWER"),
        "HF_POWER": robust_metrics.get("HF_POWER"),
        "LF_NU": robust_metrics.get("LF_NU"),
        "HF_NU": robust_metrics.get("HF_NU"),
        "LF_HF": robust_metrics.get("LF_HF"),
        "MEASUREMENT_QUALITY": _v33_v14_extract_measurement_quality(text) or extract_quality(text),
        "PNS_INDEX": robust_metrics.get("PNS_INDEX"),
        "SNS_INDEX": robust_metrics.get("SNS_INDEX"),
        "PHYSIO_AGE": robust_metrics.get("PHYSIO_AGE"),
    }

    debug = []
    confs = robust_metrics.get("_CONFIDENCE", {}) if isinstance(robust_metrics, dict) else {}
    audit = robust_metrics.get("_AUDIT", []) if isinstance(robust_metrics, dict) else []

    # Se conserva la hoja debug original, pero ya no se leen burbujas.
    # BUBBLE_* queda vacío deliberadamente; OCR_RAW recoge el método/detalle V14.
    for var, val, conf, method, detail in audit:
        key = V33_V14_TO_INTERNAL.get(var, var)
        debug.append({
            "FILE": Path(path).name,
            "FECHA_IMAGEN": fecha_imagen,
            "HORA_IMAGEN": hora_imagen,
            "KEY": key,
            "BUBBLE_X": None,
            "BUBBLE_Y": None,
            "BUBBLE_W": None,
            "BUBBLE_H": None,
            "OCR_RAW": f"{method} | conf={conf:.3f} | {detail}",
            "OCR_VALUE": val,
        })

    # Añadir una línea de auditoría para PNS/SNS/edad aunque no aparezcan: evita falsa lectura por burbuja.
    for key in ["PNS_INDEX", "SNS_INDEX", "PHYSIO_AGE"]:
        if not any(r.get("KEY") == key for r in debug):
            debug.append({
                "FILE": Path(path).name,
                "FECHA_IMAGEN": fecha_imagen,
                "HORA_IMAGEN": hora_imagen,
                "KEY": key,
                "BUBBLE_X": None,
                "BUBBLE_Y": None,
                "BUBBLE_W": None,
                "BUBBLE_H": None,
                "OCR_RAW": "no_detectado_textualmente_motor_v14_sin_burbujas",
                "OCR_VALUE": data.get(key),
            })

    missing_core = [k for k in CORE_OCR_VARIABLES if data.get(k) is None]
    missing_indices = [k for k in ["PNS_INDEX", "SNS_INDEX", "PHYSIO_AGE"] if data.get(k) is None]
    missing = missing_core + missing_indices
    data["ESTADO_EXTRACCION"] = "OK" if not missing_core else "REVISAR_MANUAL"
    data["FALTAN"] = ", ".join(missing)
    return data, debug, text

def _expand_file_field_to_names(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return []
    txt = str(value).strip()
    if not txt:
        return []
    parts = re.split(r"\s*\|\s*", txt)
    return [p.strip() for p in parts if p and p.strip()]


def load_existing_sheets(output_path):
    output = Path(output_path)
    if not output.exists():
        return {}, set()
    existing = {}
    try:
        xls = pd.ExcelFile(output)
        for sheet in ["datos_kubios", "debug_extraccion", "ocr_texto", "resumen", "metadata_informe", "variables_informe"]:
            if sheet in xls.sheet_names:
                existing[sheet] = pd.read_excel(output, sheet_name=sheet)
    except Exception:
        return {}, set()
    processed_files = set()
    if "ocr_texto" in existing and "FILE" in existing["ocr_texto"].columns:
        processed_files.update(existing["ocr_texto"]["FILE"].dropna().astype(str).tolist())
    elif "datos_kubios" in existing and "FILE" in existing["datos_kubios"].columns:
        for value in existing["datos_kubios"]["FILE"].dropna().tolist():
            processed_files.update(_expand_file_field_to_names(value))
    return existing, processed_files


def ensure_columns(df, required_columns):
    if df is None or df.empty:
        return pd.DataFrame(columns=required_columns)
    for col in required_columns:
        if col not in df.columns:
            df[col] = None
    return df[required_columns]


def merge_dataframes(old_df, new_df, key_columns=None):
    if old_df is None or old_df.empty:
        return new_df.copy()
    if new_df is None or new_df.empty:
        return old_df.copy()
    merged = pd.concat([old_df, new_df], ignore_index=True)
    if key_columns:
        available_keys = [c for c in key_columns if c in merged.columns]
        if available_keys:
            merged = merged.drop_duplicates(subset=available_keys, keep="last")
    return merged


def build_datetime_series(df):
    if df is None or df.empty:
        return pd.Series(dtype="datetime64[ns]")
    fecha = df["FECHA_IMAGEN"].astype(str) if "FECHA_IMAGEN" in df.columns else pd.Series("", index=df.index)
    hora = df["HORA_IMAGEN"].astype(str) if "HORA_IMAGEN" in df.columns else pd.Series("", index=df.index)
    return pd.to_datetime(fecha + " " + hora, errors="coerce")


def build_session_id(fecha, dt_value, file_value=None):
    if pd.notna(dt_value):
        return pd.to_datetime(dt_value).strftime("%Y-%m-%d %H:%M:%S")
    if fecha not in [None, "", "nan"]:
        file_part = safe_filename_text(file_value) if file_value not in [None, "", "nan"] else "sin_archivo"
        return f"{fecha}__{file_part}"
    return safe_filename_text(file_value) if file_value not in [None, "", "nan"] else "sesion_sin_fecha"


SESSION_MERGE_PRIORITY = [
    "READINESS", "HR", "RMSSD", "RR_MEAN", "SDNN", "SD1", "SD2",
    "STRESS_INDEX", "BREATH_RATE", "LF_POWER", "HF_POWER", "LF_NU", "HF_NU", "LF_HF",
    "MEASUREMENT_QUALITY", "PNS_INDEX", "SNS_INDEX", "PHYSIO_AGE",
]


def consolidate_session_rows(df, gap_seconds=SESSION_GROUP_GAP_SECONDS):
    if df is None or df.empty:
        return df.copy() if df is not None else pd.DataFrame()

    out = df.copy()
    out["DATETIME_IMAGEN"] = build_datetime_series(out)
    sort_cols = [c for c in ["FECHA_IMAGEN", "DATETIME_IMAGEN", "HORA_IMAGEN", "FILE"] if c in out.columns]
    out = out.sort_values(sort_cols, na_position="last").reset_index(drop=True)

    sessions = []
    current_idxs = []
    current_date = None
    current_dt = None

    for idx, row in out.iterrows():
        row_date = row.get("FECHA_IMAGEN")
        row_dt = row.get("DATETIME_IMAGEN")

        if not current_idxs:
            current_idxs = [idx]
            current_date = row_date
            current_dt = row_dt
            continue

        same_date = row_date == current_date and row_date not in [None, "", "nan"]
        close_in_time = False
        if pd.notna(row_dt) and pd.notna(current_dt):
            close_in_time = abs((row_dt - current_dt).total_seconds()) <= gap_seconds

        if same_date and close_in_time:
            current_idxs.append(idx)
            current_dt = row_dt
        else:
            sessions.append(current_idxs)
            current_idxs = [idx]
            current_date = row_date
            current_dt = row_dt

    if current_idxs:
        sessions.append(current_idxs)

    merged_rows = []
    for idxs in sessions:
        grp = out.loc[idxs].copy().reset_index(drop=True)
        merged = {}
        files = [str(x).strip() for x in grp.get("FILE", pd.Series(dtype=object)).tolist() if str(x).strip() and str(x).strip().lower() != "nan"]
        merged["FILE"] = " | ".join(dict.fromkeys(files))
        merged["FECHA_IMAGEN"] = next((x for x in grp.get("FECHA_IMAGEN", pd.Series(dtype=object)).tolist() if pd.notna(x) and str(x).strip() != ""), None)
        non_null_dt = grp["DATETIME_IMAGEN"].dropna() if "DATETIME_IMAGEN" in grp.columns else pd.Series(dtype="datetime64[ns]")
        first_dt = non_null_dt.iloc[0] if len(non_null_dt) else pd.NaT
        merged["HORA_IMAGEN"] = pd.to_datetime(first_dt).strftime("%H:%M:%S") if pd.notna(first_dt) else next((x for x in grp.get("HORA_IMAGEN", pd.Series(dtype=object)).tolist() if pd.notna(x) and str(x).strip() != ""), None)
        merged["DATETIME_IMAGEN"] = first_dt if pd.notna(first_dt) else pd.NaT
        merged["SESSION_ID"] = build_session_id(merged.get("FECHA_IMAGEN"), merged.get("DATETIME_IMAGEN"), merged.get("FILE"))
        merged["N_FILES_SESION"] = len(files) if files else len(grp)

        for col in SESSION_MERGE_PRIORITY:
            if col not in grp.columns:
                merged[col] = None
                continue
            values = [v for v in grp[col].tolist() if pd.notna(v) and str(v).strip() != ""]
            if not values:
                merged[col] = None
            else:
                merged[col] = values[0]

        faltantes = []
        for col in ["PNS_INDEX", "SNS_INDEX", "PHYSIO_AGE"]:
            if merged.get(col) is None:
                faltantes.append(col)
        merged["ESTADO_EXTRACCION"] = "OK" if not faltantes else "REVISAR_MANUAL"
        merged["FALTAN"] = ", ".join(faltantes)
        merged_rows.append(merged)

    return pd.DataFrame(merged_rows)


def add_derived_columns(df):
    if df is None or df.empty:
        return df
    df = df.copy()
    if "DATETIME_IMAGEN" not in df.columns:
        df["DATETIME_IMAGEN"] = build_datetime_series(df)
    else:
        dt_existing = pd.to_datetime(df["DATETIME_IMAGEN"], errors="coerce")
        dt_from_fields = build_datetime_series(df)
        df["DATETIME_IMAGEN"] = dt_existing.where(dt_existing.notna(), dt_from_fields)
    sort_cols = [c for c in ["DATETIME_IMAGEN", "FILE"] if c in df.columns]
    df = df.sort_values(sort_cols, na_position="last").reset_index(drop=True)
    df["ORDEN_DIA"] = df.groupby("FECHA_IMAGEN").cumcount() + 1
    if "SESSION_ID" not in df.columns:
        df["SESSION_ID"] = [build_session_id(fecha, dt, file_value) for fecha, dt, file_value in zip(df.get("FECHA_IMAGEN", pd.Series(index=df.index)), df.get("DATETIME_IMAGEN", pd.Series(index=df.index)), df.get("FILE", pd.Series(index=df.index)))]
    if "N_FILES_SESION" not in df.columns:
        df["N_FILES_SESION"] = df.get("FILE", pd.Series("", index=df.index)).astype(str).apply(lambda x: len([p for p in re.split(r"\s*\|\s*", x) if p and p.lower() != "nan"]))
    rmssd_num = pd.to_numeric(df["RMSSD"], errors="coerce") if "RMSSD" in df.columns else pd.Series(np.nan, index=df.index)
    df["LnRMSSD"] = np.where(rmssd_num > 0, np.log(rmssd_num), np.nan)
    cols = list(df.columns)
    if "ORDEN_DIA" in cols:
        cols.remove("ORDEN_DIA")
    idx = cols.index("HORA_IMAGEN") + 1 if "HORA_IMAGEN" in cols else 2
    cols.insert(idx, "ORDEN_DIA")
    return df[cols]


def build_summary(df, n_new):
    if df is None or df.empty:
        return pd.DataFrame([{"N_IMAGENES": 0, "N_OK": 0, "N_REVISAR": 0, "N_NUEVAS_EN_ESTA_EJECUCION": n_new}])
    return pd.DataFrame([{
        "N_IMAGENES": len(df),
        "N_OK": int((df["ESTADO_EXTRACCION"] == "OK").sum()) if "ESTADO_EXTRACCION" in df.columns else 0,
        "N_REVISAR": int((df["ESTADO_EXTRACCION"] != "OK").sum()) if "ESTADO_EXTRACCION" in df.columns else 0,
        "N_NUEVAS_EN_ESTA_EJECUCION": n_new,
    }])


def safe_filename_text(texto):
    texto = str(texto).strip()
    if not texto:
        return "Sin_nombre"
    texto = unicodedata.normalize("NFKC", texto)
    texto = re.sub(r"\s+", "_", texto)
    texto = re.sub(r'[\\/:*?"<>|]+', "", texto)
    texto = re.sub(r"_+", "_", texto).strip("._")
    return texto or "Sin_nombre"


def build_output_paths(nombre_sujeto, fecha_informe, tipo_informe_clave):
    """Rutas cortas y compatibles con Windows.

    Estructura generada:
    Resultados/
      Participante/
        Excel/
        Word/
        Graficos/
    """
    INFORMES_DIR.mkdir(parents=True, exist_ok=True)
    nombre_limpio = safe_filename_text(nombre_sujeto)

    subject_dir = INFORMES_DIR / nombre_limpio
    word_dir = subject_dir / "Word"
    excel_dir = subject_dir / "Excel"
    graph_dir = subject_dir / "Graficos"

    word_dir.mkdir(parents=True, exist_ok=True)
    excel_dir.mkdir(parents=True, exist_ok=True)
    graph_dir.mkdir(parents=True, exist_ok=True)

    sufijo = "ENT" if tipo_informe_clave == "entrenador" else "PAR"
    output_docx = word_dir / f"{nombre_limpio}_{sufijo}.docx"
    output_xlsx = excel_dir / f"{nombre_limpio}.xlsx"

    return output_docx, output_xlsx, subject_dir, graph_dir


def build_incremented_output_path(output_docx):
    output_docx = Path(output_docx)
    for i in range(2, 1000):
        candidate = output_docx.with_name(f"{output_docx.stem}_v{i}{output_docx.suffix}")
        if not candidate.exists():
            return candidate
    timestamp = pd.Timestamp.now().strftime("%Hh%Mm%Ss")
    return output_docx.with_name(f"{output_docx.stem}_{timestamp}{output_docx.suffix}")


def decidir_guardado_informe(output_docx, global_mode=None):
    output_docx = Path(output_docx)
    if not output_docx.exists():
        return output_docx, "nuevo"

    if global_mode == "nueva_version":
        nuevo = build_incremented_output_path(output_docx)
        print("")
        print("Ya existe un informe Word con este nombre:")
        print(output_docx)
        print(f"Se aplicará la política global: generar nueva versión -> {nuevo.name}")
        return nuevo, "nueva_version"

    if global_mode == "reescrito":
        print("")
        print("Ya existe un informe Word con este nombre:")
        print(output_docx)
        print("Se aplicará la política global: reescribir el informe existente.")
        return output_docx, "reescrito"

    print("")
    print("Ya existe un informe Word con este nombre:")
    print(output_docx)
    print("¿Qué deseas hacer?")
    print("1. Reescribir el informe existente")
    print("2. Generar un nuevo informe con sufijo incremental")
    while True:
        entrada = input("Selecciona 1 o 2 [Enter = 2]: ").strip()
        if entrada in {"", "2"}:
            nuevo = build_incremented_output_path(output_docx)
            print(f"Se generará una nueva versión del informe: {nuevo.name}")
            return nuevo, "nueva_version"
        if entrada == "1":
            print("Se intentará reescribir el informe existente.")
            return output_docx, "reescrito"
        print("Entrada no válida. Introduce 1 para reescribir o 2 para generar una nueva versión.")

def save_excel(df_final, df_debug_final, df_ocr_final, df_summary, df_metadata_informe, df_variables_informe, output_xlsx):
    output_xlsx = Path(output_xlsx)
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        df_final.to_excel(writer, sheet_name="datos_kubios", index=False)
        df_debug_final.to_excel(writer, sheet_name="debug_extraccion", index=False)
        df_ocr_final.to_excel(writer, sheet_name="ocr_texto", index=False)
        df_summary.to_excel(writer, sheet_name="resumen", index=False)
        df_metadata_informe.to_excel(writer, sheet_name="metadata_informe", index=False)
        df_variables_informe.to_excel(writer, sheet_name="variables_informe", index=False)

def get_definition_text(var, definition_mode):
    meta = VARIABLE_META.get(var, {})
    simple = SIMPLE_INTERPRETATION.get(var, "")
    if definition_mode == "detallada":
        base = meta.get("explicacion", "")
        interpret = meta.get("interpretacion", "")
        partes = [p for p in [base, interpret, simple] if p]
        return " ".join(partes)
    if definition_mode == "breve":
        base = SHORT_VARIABLE_INFO.get(var, meta.get("nombre", var))
        return f"{base} {simple}".strip()
    return ""


def get_definition_label(definition_mode):
    return {
        "breve": "Breve",
        "detallada": "Detallada",
        "directa": "Directa sin definiciones previas",
    }.get(definition_mode, str(definition_mode))


def build_report_metadata(run_id, nombre_sujeto, tipo_informe_label, fecha_informe, fecha_desde, momento, definition_mode, selected_vars, n_registros_informe, short_window=7, long_window=28):
    short_window, long_window = safe_window_pair(short_window, long_window)
    return pd.DataFrame([{
        "RUN_ID": run_id,
        "FECHA_EJECUCION": fecha_informe,
        "NOMBRE_SUJETO": nombre_sujeto,
        "TIPO_INFORME": tipo_informe_label,
        "FECHA_DESDE": str(fecha_desde),
        "ORDEN_REGISTRO": str(momento),
        "MODO_DEFINICION": definition_mode,
        "MODO_DEFINICION_ETIQUETA": get_definition_label(definition_mode),
        "N_VARIABLES_SELECCIONADAS": len(selected_vars),
        "VARIABLES_SELECCIONADAS": ", ".join(selected_vars),
        "N_REGISTROS_INFORME": int(n_registros_informe),
        "VENTANA_RECIENTE": int(short_window),
        "VENTANA_REFERENCIA": int(long_window),
    }])


def build_variables_metadata(run_id, nombre_sujeto, tipo_informe_label, fecha_informe, definition_mode, selected_vars):
    rows = []
    for orden, var in enumerate(selected_vars, start=1):
        meta = VARIABLE_META.get(var, {})
        rows.append({
            "RUN_ID": run_id,
            "FECHA_EJECUCION": fecha_informe,
            "NOMBRE_SUJETO": nombre_sujeto,
            "TIPO_INFORME": tipo_informe_label,
            "ORDEN_VARIABLE": orden,
            "VARIABLE": var,
            "NOMBRE_VARIABLE": meta.get("nombre", var),
            "MODO_DEFINICION": definition_mode,
            "MODO_DEFINICION_ETIQUETA": get_definition_label(definition_mode),
            "DEFINICION_MOSTRADA": get_definition_text(var, definition_mode) if definition_mode in {"breve", "detallada"} else "",
            "INTERPRETACION_GENERAL": meta.get("interpretacion", ""),
        })
    return pd.DataFrame(rows)


def set_margins(doc):
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)


def add_footer_date(doc, fecha_texto):
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.text = f"Fecha de realización del informe: {fecha_texto}"


def set_paragraph_bottom_border(paragraph, size=8, space=1, color="000000"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)


def set_cell_bottom_border(cell, size=8, color="000000"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    bottom = tc_borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        tc_borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), color)


def add_document_header(doc, header_text, logo_path=None):
    logo_inserted = False
    for section in doc.sections:
        section.different_first_page_header_footer = False
        header = section.header

        # Limpiar el encabezado existente
        for paragraph in list(header.paragraphs):
            p_el = paragraph._element
            p_el.getparent().remove(p_el)
        for table in list(header.tables):
            t_el = table._element
            t_el.getparent().remove(t_el)

        table = header.add_table(rows=1, cols=2, width=Inches(6.5))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders_none(table)

        try:
            table.columns[0].width = Inches(5.2)
            table.columns[1].width = Inches(1.3)
        except Exception:
            pass

        left_cell = table.cell(0, 0)
        right_cell = table.cell(0, 1)
        left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_bottom_border(left_cell, size=8, color="000000")
        set_cell_bottom_border(right_cell, size=8, color="000000")

        p_left = left_cell.paragraphs[0]
        p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_left.paragraph_format.space_before = Pt(0)
        p_left.paragraph_format.space_after = Pt(0)
        run_left = p_left.add_run(header_text)
        run_left.bold = True
        run_left.font.size = Pt(10)

        if logo_path and Path(logo_path).exists():
            p_right = right_cell.paragraphs[0]
            p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_right.paragraph_format.space_before = Pt(0)
            p_right.paragraph_format.space_after = Pt(0)
            run_logo = p_right.add_run()
            run_logo.add_picture(str(logo_path), width=Inches(1.0))
            logo_inserted = True
        else:
            right_cell.paragraphs[0].text = ""

    return logo_inserted


def moving_average(series, window=3, min_periods=None):
    if min_periods is None:
        min_periods = max(1, int(np.ceil(window * 0.7))) if window and window > 1 else 1
    return series.rolling(window=window, min_periods=min_periods).mean()


def safe_window_pair(short_window, long_window):
    try:
        short_window = int(short_window)
    except Exception:
        short_window = 7
    try:
        long_window = int(long_window)
    except Exception:
        long_window = 28

    short_window = max(2, short_window)
    long_window = max(short_window + 1, long_window)
    return short_window, long_window


def compute_window_comparison(series, short_window=7, long_window=28):
    short_window, long_window = safe_window_pair(short_window, long_window)
    s = pd.to_numeric(series, errors="coerce").dropna()
    if s.empty:
        return {
            "short_window": short_window,
            "long_window": long_window,
            "short_mean": None,
            "long_mean": None,
            "delta_abs": None,
            "delta_pct": None,
            "enough_short": False,
            "enough_long": False,
            "n_valid": 0,
        }

    short_slice = s.tail(short_window)
    long_slice = s.tail(long_window)

    short_mean = float(short_slice.mean()) if len(short_slice) >= max(2, int(np.ceil(short_window * 0.7))) else None
    long_mean = float(long_slice.mean()) if len(long_slice) >= max(4, int(np.ceil(long_window * 0.7))) else None

    delta_abs = None
    delta_pct = None
    if short_mean is not None and long_mean is not None and np.isfinite(long_mean):
        delta_abs = float(short_mean - long_mean)
        if long_mean != 0:
            delta_pct = float((delta_abs / abs(long_mean)) * 100.0)

    return {
        "short_window": short_window,
        "long_window": long_window,
        "short_mean": short_mean,
        "long_mean": long_mean,
        "delta_abs": delta_abs,
        "delta_pct": delta_pct,
        "enough_short": short_mean is not None,
        "enough_long": long_mean is not None,
        "n_valid": int(len(s)),
    }


def relative_state_direction(var):
    favorable_high = {"RMSSD", "LnRMSSD", "SDNN", "SD1", "SD2", "PNS_INDEX", "RR_MEAN", "HF_POWER", "HF_NU"}
    favorable_low = {"HR", "STRESS_INDEX", "SNS_INDEX", "LF_HF", "LF_POWER", "LF_NU", "BREATH_RATE", "PHYSIO_AGE"}
    if var in favorable_high:
        return "higher_is_better"
    if var in favorable_low:
        return "lower_is_better"
    return "neutral"


def interpret_window_comparison(var, comparison):
    if not comparison:
        return "No hay comparación temporal disponible."

    short_window = comparison.get("short_window")
    long_window = comparison.get("long_window")
    short_mean = comparison.get("short_mean")
    long_mean = comparison.get("long_mean")
    delta_pct = comparison.get("delta_pct")

    if short_mean is None or long_mean is None:
        return (
            f"No hay suficientes registros válidos para comparar la media reciente de {short_window} registros "
            f"con la media de los últimos {long_window} registros."
        )

    if delta_pct is None:
        return (
            f"La media reciente de los últimos {short_window} registros y la media de referencia de los últimos {long_window} registros están disponibles, "
            f"pero no puede calcularse una diferencia porcentual robusta."
        )

    abs_pct = abs(delta_pct)
    if abs_pct < 3:
        magnitud = "muy próxima"
    elif abs_pct < 8:
        magnitud = "moderadamente distinta"
    else:
        magnitud = "claramente distinta"

    direction = relative_state_direction(var)
    if delta_pct > 0:
        relation = "por encima"
    elif delta_pct < 0:
        relation = "por debajo"
    else:
        relation = "prácticamente al mismo nivel que"

    base = (
        f"La media reciente de {short_window} registros ({short_mean:.2f}) está {relation} "
        f"la media de los últimos {long_window} registros ({long_mean:.2f}), con una diferencia del {delta_pct:.1f}%. "
        f"La separación entre ambas es {magnitud}."
    )

    if direction == "higher_is_better":
        if delta_pct >= 3:
            meaning = "En esta variable, que los valores recientes estén por encima de tu referencia más amplia suele ser compatible con una respuesta aguda favorable o una adaptación positiva reciente."
        elif delta_pct <= -3:
            meaning = "En esta variable, que los valores recientes estén por debajo de tu referencia más amplia puede ser compatible con fatiga, estrés fisiológico o una recuperación menos favorable."
        else:
            meaning = "En esta variable, la cercanía entre ambas medias sugiere estabilidad respecto a tu referencia reciente-amplia."
    elif direction == "lower_is_better":
        if delta_pct <= -3:
            meaning = "En esta variable, que los valores recientes estén por debajo de tu referencia más amplia suele ser compatible con menor activación relativa o mejor recuperación."
        elif delta_pct >= 3:
            meaning = "En esta variable, que los valores recientes estén por encima de tu referencia más amplia puede ser compatible con mayor carga interna, activación o necesidad de recuperar mejor."
        else:
            meaning = "En esta variable, la cercanía entre ambas medias sugiere estabilidad respecto a tu referencia reciente-amplia."
    else:
        meaning = "La dirección del cambio debe interpretarse con prudencia y junto con el resto de variables del perfil autonómico."

    return f"{base} {meaning}"


def recommendation_from_window_comparison(var, comparison):
    if not comparison or comparison.get("short_mean") is None or comparison.get("long_mean") is None:
        return "Recomendación adicional: mantener el seguimiento hasta acumular suficientes registros para comparar el comportamiento reciente con la referencia más amplia."

    delta_pct = comparison.get("delta_pct")
    if delta_pct is None:
        return "Recomendación adicional: revisar el contexto de medición y seguir acumulando registros para interpretar mejor la diferencia entre ventanas."

    direction = relative_state_direction(var)

    if direction == "higher_is_better" and delta_pct <= -5:
        return "Recomendación adicional: los valores recientes están por debajo de la referencia amplia; conviene revisar carga, sueño, estrés y sensaciones antes de aumentar la exigencia."
    if direction == "higher_is_better" and delta_pct >= 5:
        return "Recomendación adicional: los valores recientes superan la referencia amplia; puede ser una señal favorable si también acompañan buenas sensaciones y estabilidad del resto de indicadores."
    if direction == "lower_is_better" and delta_pct >= 5:
        return "Recomendación adicional: los valores recientes están por encima de la referencia amplia; conviene vigilar la evolución en próximos registros y contrastarla con carga, descanso y percepción subjetiva."
    if direction == "lower_is_better" and delta_pct <= -5:
        return "Recomendación adicional: los valores recientes están por debajo de la referencia amplia; esto puede ser favorable si el resto del patrón acompaña y la persona se encuentra bien."
    return "Recomendación adicional: la diferencia entre ventanas no es grande; lo más útil es seguir observando si esta estabilidad se mantiene con el tiempo."




def compute_reference_metrics(series, short_window=7, long_window=28, swc_factor=0.5):
    short_window, long_window = safe_window_pair(short_window, long_window)
    comparison = compute_window_comparison(series, short_window=short_window, long_window=long_window)
    s = pd.to_numeric(series, errors="coerce").dropna()

    current_value = float(s.iloc[-1]) if len(s) > 0 and pd.notna(s.iloc[-1]) else None
    ratio = None
    zscore = None
    reference_sd = None

    long_slice = s.tail(long_window)
    if len(long_slice) >= max(4, int(np.ceil(long_window * 0.7))):
        reference_sd = float(long_slice.std(ddof=0)) if pd.notna(long_slice.std(ddof=0)) else None

    short_mean = comparison.get("short_mean")
    long_mean = comparison.get("long_mean")

    if short_mean is not None and long_mean is not None and np.isfinite(long_mean) and long_mean != 0:
        ratio = float(short_mean / long_mean)

    if current_value is not None and long_mean is not None and reference_sd is not None and np.isfinite(reference_sd) and reference_sd > 0:
        zscore = float((current_value - long_mean) / reference_sd)

    swc = None
    swc_relevant = None
    if reference_sd is not None and np.isfinite(reference_sd):
        swc = float(abs(swc_factor) * reference_sd)
        delta_abs = comparison.get("delta_abs")
        if delta_abs is not None and np.isfinite(delta_abs):
            swc_relevant = bool(abs(delta_abs) > swc)

    return {
        "comparison": comparison,
        "current_value": current_value,
        "ratio": ratio,
        "zscore": zscore,
        "reference_sd": reference_sd,
        "swc_factor": float(swc_factor),
        "swc": swc,
        "swc_relevant": swc_relevant,
        "short_window": short_window,
        "long_window": long_window,
    }


def interpret_ratio(var, ratio):
    if ratio is None or not np.isfinite(ratio):
        return "No hay suficientes registros válidos para calcular el ratio reciente/referencia."

    direction = relative_state_direction(var)
    if 0.95 <= ratio <= 1.05:
        return "El comportamiento reciente se mantiene muy próximo a la referencia amplia, lo que sugiere estabilidad."

    if direction == "higher_is_better":
        if ratio > 1.05:
            return "El comportamiento reciente está por encima de la referencia amplia; en esta variable eso suele ser una señal favorable dentro del periodo analizado."
        return "El comportamiento reciente está por debajo de la referencia amplia; en esta variable eso puede sugerir un momento de mayor exigencia o menor frescura."
    if direction == "lower_is_better":
        if ratio < 0.95:
            return "El comportamiento reciente está por debajo de la referencia amplia; en esta variable eso suele ser una señal favorable dentro del periodo analizado."
        return "El comportamiento reciente está por encima de la referencia amplia; en esta variable eso puede sugerir más activación o una etapa más exigente."
    return "El ratio debe leerse con prudencia y siempre en conjunto con el resto del patrón observado."


def interpret_zscore_swc(var, metrics):
    if not metrics:
        return "No hay datos suficientes para calcular Z-score y SWC."

    zscore = metrics.get("zscore")
    swc = metrics.get("swc")
    swc_relevant = metrics.get("swc_relevant")

    if zscore is None or not np.isfinite(zscore):
        base = "No hay datos suficientes para ubicar el último registro dentro de su rango de referencia."
    else:
        if zscore >= 1:
            base = "El último registro se sitúa claramente por encima de su rango habitual reciente."
        elif zscore <= -1:
            base = "El último registro se sitúa claramente por debajo de su rango habitual reciente."
        else:
            base = "El último registro se sitúa dentro de un rango próximo a lo habitual."

    if swc is None or not np.isfinite(swc):
        return base + " No hay base suficiente para decidir si el cambio supera un umbral práctico de relevancia."

    if swc_relevant is True:
        return base + " Además, la diferencia entre la media reciente y la referencia supera el umbral de cambio relevante (SWC), por lo que el cambio podría tener interés práctico."
    if swc_relevant is False:
        return base + " Además, la diferencia entre la media reciente y la referencia no supera el umbral de cambio relevante (SWC), por lo que podría estar dentro de la variación esperable."
    return base + " El umbral SWC está disponible, pero no puede aplicarse de forma robusta al cambio observado."


def emotional_interpretation_for_variable(var, stats, metrics):
    slope = (stats or {}).get("slope")
    comparison = (metrics or {}).get("comparison") or {}
    delta_pct = comparison.get("delta_pct")
    direction = relative_state_direction(var)

    if slope is None and delta_pct is None:
        return "Lectura emocional orientativa: con los registros disponibles esta variable transmite una sensación de continuidad, sin señales claras de cambio. Es una lectura descriptiva y no médica."

    favorable = 0
    demanding = 0

    if slope is not None:
        if direction == "higher_is_better":
            favorable += int(slope > 0)
            demanding += int(slope < 0)
        elif direction == "lower_is_better":
            favorable += int(slope < 0)
            demanding += int(slope > 0)

    if delta_pct is not None:
        if direction == "higher_is_better":
            favorable += int(delta_pct > 3)
            demanding += int(delta_pct < -3)
        elif direction == "lower_is_better":
            favorable += int(delta_pct < -3)
            demanding += int(delta_pct > 3)

    if favorable >= demanding + 1 and favorable >= 1:
        return "esta variable transmite un momento de mayor soltura, mejor disposición o sensación de margen en tu proceso. No es una interpretación médica, sino una forma práctica de resumir la señal."
    if demanding >= favorable + 1 and demanding >= 1:
        return "esta variable sugiere una etapa de mayor exigencia, más tensión interna o necesidad de ir con algo más de calma. No es una interpretación médica, sino una lectura práctica del comportamiento reciente."
    return "esta variable transmite estabilidad y continuidad, sin cambios muy marcados hacia un sentido u otro. No es una interpretación médica, sino una lectura práctica de la evolución observada."


def format_metric_value(value, decimals=2, suffix=""):
    if value is None or (isinstance(value, float) and not np.isfinite(value)):
        return "No disponible"
    if isinstance(value, (int, np.integer)):
        return f"{int(value)}{suffix}"
    if isinstance(value, (float, np.floating)):
        return f"{value:.{decimals}f}{suffix}"
    return f"{value}{suffix}"


def build_interpretation_table_rows(var, stats, metrics, selected_blocks):
    rows = []
    if not stats:
        return [("Estado", "Sin datos suficientes", "No fue posible construir una lectura robusta para esta variable.")]

    slope = stats.get("slope")
    change_pct = stats.get("change_pct")
    trend_label = "Ascendente" if slope is not None and slope > 0 else "Descendente" if slope is not None and slope < 0 else "Estable / no concluyente"
    rows.append(("Tendencia lineal", trend_label, automatic_interpretation(var, stats, {var: slope} if slope is not None else {})))

    rows.append(("Cambio inicio-fin", format_metric_value(change_pct, 1, "%"), "Cambio porcentual aproximado entre el inicio y el final del tramo representado."))

    comparison = (metrics or {}).get("comparison") or {}
    if "comparacion_ventanas" in selected_blocks:
        rows.append((
            f"Media reciente ({comparison.get('short_window', '')})",
            format_metric_value(comparison.get("short_mean"), 2),
            "Promedio de los registros más recientes dentro de la ventana corta."
        ))
        rows.append((
            f"Media referencia ({comparison.get('long_window', '')})",
            format_metric_value(comparison.get("long_mean"), 2),
            "Promedio de la ventana de referencia amplia."
        ))
        rows.append((
            "Comparación de ventanas",
            format_metric_value(comparison.get("delta_pct"), 1, "%"),
            interpret_window_comparison(var, comparison)
        ))

    if "ratio" in selected_blocks:
        rows.append(("Ratio reciente/referencia", format_metric_value((metrics or {}).get("ratio"), 3), interpret_ratio(var, (metrics or {}).get("ratio"))))

    if "zscore_swc" in selected_blocks:
        rows.append(("Z-score", format_metric_value((metrics or {}).get("zscore"), 2), "Ubica el último registro respecto a su rango habitual reciente."))
        rows.append(("SWC", format_metric_value((metrics or {}).get("swc"), 2), interpret_zscore_swc(var, metrics)))

    return rows


def add_interpretation_table(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = table.rows[0].cells
    headers[0].text = "Indicador"
    headers[1].text = "Valor"
    headers[2].text = "Lectura breve"
    for label, value, meaning in rows:
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)
        cells[2].text = str(meaning)
    style_table_font_8(table)
    doc.add_paragraph("")


def block_title_text(block):
    mapping = {
        "grafico": "Bloque gráfico",
        "interpretacion": "Lectura de la tendencia",
        "recomendaciones": "Orientación práctica",
        "comparacion_ventanas": "Bloque de comparación entre ventanas",
        "ratio": "Bloque de ratio reciente/referencia",
        "zscore_swc": "Bloque de Z-score y SWC",
        "tabla_interpretacion": "Bloque de tabla de interpretación",
        "lectura_emocional": "Bloque emocional",
    }
    return mapping.get(block, "Bloque analítico")


def block_intro_text(block, short_window=7, long_window=28, swc_factor=0.5):
    mapping = {
        "grafico": "Permite ver de forma visual cómo evoluciona la variable a lo largo del periodo analizado.",
        "interpretacion": "Resume de forma breve la dirección general que sigue la variable durante el periodo analizado.",
        "recomendaciones": "Sugerencia práctica para esta variable.",
        "comparacion_ventanas": f"Compara la media reciente de los últimos {short_window} registros con la media de referencia de los últimos {long_window} registros.",
        "ratio": f"Expresa en una sola cifra la relación entre la media reciente de {short_window} registros y la media de referencia de {long_window} registros.",
        "zscore_swc": f"Ubica el último valor respecto a su rango habitual y valora si el cambio reciente supera un umbral práctico de relevancia usando un factor SWC de {swc_factor:.2f}.",
        "tabla_interpretacion": "Resume en una tabla los principales indicadores calculados para facilitar una lectura rápida.",
        "lectura_emocional": "",
    }
    return mapping.get(block, "Bloque analítico complementario.")


def block_unavailable_text(block, var, stats=None, metrics=None, short_window=7, long_window=28):
    stats = stats or {}
    metrics = metrics or {}
    comparison = (metrics.get("comparison") or {}) if isinstance(metrics, dict) else {}
    n_valid = comparison.get("n_valid")
    short_mean = comparison.get("short_mean")
    long_mean = comparison.get("long_mean")
    zscore = metrics.get("zscore") if isinstance(metrics, dict) else None
    swc = metrics.get("swc") if isinstance(metrics, dict) else None

    if block == "grafico":
        return "No se ha podido realizar el análisis por falta de datos suficientes o por imposibilidad de generar el gráfico con los registros disponibles."

    if block == "interpretacion":
        return "No se ha podido realizar el análisis por falta de datos suficientes para estimar una tendencia lineal robusta en esta variable."

    if block == "recomendaciones":
        return "No se ha podido realizar el análisis por falta de datos suficientes para generar una sugerencia práctica fiable en esta variable."

    if block == "comparacion_ventanas":
        if n_valid in [0, None]:
            return "No se ha podido realizar el análisis por falta de datos suficientes: no hay registros válidos de esta variable en el periodo analizado."
        if short_mean is None and long_mean is None:
            return f"No se ha podido realizar el análisis por falta de datos suficientes: todavía no hay suficientes registros válidos para construir ni la media reciente de {short_window} registros ni la media de referencia de {long_window} registros."
        if long_mean is None:
            return f"No se ha podido realizar el análisis por falta de datos suficientes: sí hay señal reciente, pero aún no se alcanza una base suficientemente estable para la referencia amplia de {long_window} registros."
        return f"No se ha podido realizar el análisis por falta de datos suficientes: la media reciente de {short_window} registros no alcanza el mínimo de consistencia requerido."

    if block == "ratio":
        if short_mean is None or long_mean is None:
            return f"No se ha podido realizar el análisis por falta de datos suficientes: el ratio necesita disponer tanto de una media reciente válida ({short_window} registros) como de una media de referencia válida ({long_window} registros)."
        return "No se ha podido realizar el análisis por falta de datos suficientes para una estimación robusta."

    if block == "zscore_swc":
        if zscore is None and swc is None:
            return f"No se ha podido realizar el análisis por falta de datos suficientes: para ubicar el último registro y valorar el cambio relevante hace falta una referencia amplia más estable dentro de la ventana de {long_window} registros."
        if zscore is None:
            return "No se ha podido completar el análisis por falta de datos suficientes: el SWC puede estimarse parcialmente, pero aún no hay base suficiente para ubicar el último valor con un Z-score robusto."
        if swc is None:
            return "No se ha podido completar el análisis por falta de datos suficientes: el último valor puede ubicarse de forma orientativa, pero aún no hay base suficiente para calcular un SWC robusto."
        return "No se ha podido realizar el análisis por falta de datos suficientes para una estimación robusta."

    if block == "tabla_interpretacion":
        return "No se ha podido realizar el análisis por falta de datos suficientes para construir una tabla interpretativa completa de esta variable."

    if block == "lectura_emocional":
        return "No se ha podido realizar una lectura emocional sólida por falta de datos suficientes; por ahora solo puede hablarse de continuidad y observación."

    return "No se ha podido realizar el análisis por falta de datos suficientes."
def build_execution_summary_columns():
    cols = [
        "RUN_ID", "FECHA_REALIZACION", "HORA_REALIZACION", "TIMESTAMP_REALIZACION",
        "PARTICIPANTE", "CARPETA_PARTICIPANTE", "TIPO_INFORME",
        "FECHA_DESDE", "ORDEN_REGISTRO", "MODO_DEFINICION",
        "VARIABLES_SELECCIONADAS", "BLOQUES_SELECCIONADOS",
        "VENTANA_RECIENTE", "VENTANA_REFERENCIA", "FACTOR_SWC",
        "N_REGISTROS_INFORME", "N_IMAGENES_TOTALES_EXCEL", "N_REGISTROS_NUEVOS",
        "WORD_GENERADO", "EXCEL_PARTICIPANTE"
    ]
    for var in VARIABLE_META.keys():
        pref = f"{var}__"
        cols.extend([
            pref + "SELECCIONADA",
            pref + "PENDIENTE",
            pref + "CAMBIO_PCT",
            pref + "MEDIA_RECIENTE",
            pref + "MEDIA_REFERENCIA",
            pref + "DELTA_PCT",
            pref + "RATIO",
            pref + "ZSCORE",
            pref + "SWC",
            pref + "SWC_RELEVANTE",
            pref + "LECTURA_EMOCIONAL",
        ])
    return cols


def build_execution_summary_row(run_id, fecha_informe, hora_informe, nombre_sujeto, tipo_informe_label,
                                fecha_desde, momento, definition_mode, selected_vars, selected_blocks,
                                short_window, long_window, swc_factor, df_informe, df_final, n_new_rows,
                                graph_info, output_docx, output_xlsx):
    row = {
        "RUN_ID": run_id,
        "FECHA_REALIZACION": fecha_informe,
        "HORA_REALIZACION": hora_informe,
        "TIMESTAMP_REALIZACION": f"{fecha_informe} {hora_informe}",
        "PARTICIPANTE": nombre_sujeto,
        "CARPETA_PARTICIPANTE": BASE_DIR.name,
        "TIPO_INFORME": tipo_informe_label,
        "FECHA_DESDE": str(fecha_desde),
        "ORDEN_REGISTRO": str(momento),
        "MODO_DEFINICION": definition_mode,
        "VARIABLES_SELECCIONADAS": ", ".join(selected_vars),
        "BLOQUES_SELECCIONADOS": ", ".join(selected_blocks),
        "VENTANA_RECIENTE": int(short_window),
        "VENTANA_REFERENCIA": int(long_window),
        "FACTOR_SWC": float(swc_factor),
        "N_REGISTROS_INFORME": int(len(df_informe)) if df_informe is not None else 0,
        "N_IMAGENES_TOTALES_EXCEL": int(len(df_final)) if df_final is not None else 0,
        "N_REGISTROS_NUEVOS": int(n_new_rows),
        "WORD_GENERADO": str(output_docx),
        "EXCEL_PARTICIPANTE": str(output_xlsx),
    }
    selected_set = set(selected_vars)
    for var in VARIABLE_META.keys():
        pref = f"{var}__"
        row[pref + "SELECCIONADA"] = 1 if var in selected_set else 0
        row[pref + "PENDIENTE"] = None
        row[pref + "CAMBIO_PCT"] = None
        row[pref + "MEDIA_RECIENTE"] = None
        row[pref + "MEDIA_REFERENCIA"] = None
        row[pref + "DELTA_PCT"] = None
        row[pref + "RATIO"] = None
        row[pref + "ZSCORE"] = None
        row[pref + "SWC"] = None
        row[pref + "SWC_RELEVANTE"] = None
        row[pref + "LECTURA_EMOCIONAL"] = None
        if var not in selected_set:
            continue
        stats = (graph_info.get(var, {}) or {}).get("stats") or {}
        metrics = compute_reference_metrics(df_informe[var] if (df_informe is not None and var in df_informe.columns) else pd.Series(dtype=float),
                                            short_window=short_window, long_window=long_window, swc_factor=swc_factor)
        comparison = metrics.get("comparison") or {}
        row[pref + "PENDIENTE"] = stats.get("slope")
        row[pref + "CAMBIO_PCT"] = stats.get("change_pct")
        row[pref + "MEDIA_RECIENTE"] = comparison.get("short_mean")
        row[pref + "MEDIA_REFERENCIA"] = comparison.get("long_mean")
        row[pref + "DELTA_PCT"] = comparison.get("delta_pct")
        row[pref + "RATIO"] = metrics.get("ratio")
        row[pref + "ZSCORE"] = metrics.get("zscore")
        row[pref + "SWC"] = metrics.get("swc")
        row[pref + "SWC_RELEVANTE"] = metrics.get("swc_relevant")
        row[pref + "LECTURA_EMOCIONAL"] = emotional_interpretation_for_variable(var, stats, metrics) if ((stats and stats.get("slope") is not None) or comparison.get("delta_pct") is not None) else None
    return row


def append_execution_summary_excel(row, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    cols = build_execution_summary_columns()
    if output_path.exists():
        try:
            df_old = pd.read_excel(output_path)
        except Exception:
            df_old = pd.DataFrame(columns=cols)
    else:
        df_old = pd.DataFrame(columns=cols)
    for col in cols:
        if col not in df_old.columns:
            df_old[col] = None
    df_new = pd.DataFrame([row])
    for col in cols:
        if col not in df_new.columns:
            df_new[col] = None
    df_final = pd.concat([df_old[cols], df_new[cols]], ignore_index=True)
    df_final.to_excel(output_path, index=False)


def get_global_summary_excel_path():
    out_dir = SCRIPT_DIR / RESULTADOS_DIR_NAME / "Seguimiento_global"
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir / "todos_los_participantes_acumulado.xlsx"


def get_global_tracking_dir():
    out_dir = SCRIPT_DIR / RESULTADOS_DIR_NAME / "Seguimiento_global"
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir


def get_global_all_data_excel_path():
    return get_global_tracking_dir() / "todos_los_datos_kubios_acumulado.xlsx"


def update_global_all_data_excel(nombre_sujeto, carpeta_participante, df_final):
    output_path = get_global_all_data_excel_path()
    if df_final is None or df_final.empty:
        return output_path

    df_new = df_final.copy()
    df_new.insert(0, "CARPETA_PARTICIPANTE", str(carpeta_participante))
    df_new.insert(0, "PARTICIPANTE", str(nombre_sujeto))

    if output_path.exists():
        try:
            df_old = pd.read_excel(output_path, sheet_name="datos_kubios_todos")
        except Exception:
            df_old = pd.DataFrame()
    else:
        df_old = pd.DataFrame()

    all_cols = list(dict.fromkeys(list(df_new.columns) + list(df_old.columns)))
    for col in all_cols:
        if col not in df_new.columns:
            df_new[col] = None
        if col not in df_old.columns:
            df_old[col] = None

    df_all = pd.concat([df_old[all_cols], df_new[all_cols]], ignore_index=True)
    dedupe_cols = [c for c in ["PARTICIPANTE", "CARPETA_PARTICIPANTE", "SESSION_ID"] if c in df_all.columns]
    if len(dedupe_cols) >= 3:
        df_all = df_all.drop_duplicates(subset=dedupe_cols, keep="last")
    else:
        dedupe_cols = [c for c in ["PARTICIPANTE", "CARPETA_PARTICIPANTE", "FILE", "FECHA_IMAGEN", "HORA_IMAGEN"] if c in df_all.columns]
        if dedupe_cols:
            df_all = df_all.drop_duplicates(subset=dedupe_cols, keep="last")

    sort_cols = [c for c in ["PARTICIPANTE", "DATETIME_IMAGEN", "FECHA_IMAGEN", "HORA_IMAGEN"] if c in df_all.columns]
    if sort_cols:
        df_all = df_all.sort_values(sort_cols, na_position="last")

    resumen = (
        df_all.groupby(["PARTICIPANTE", "CARPETA_PARTICIPANTE"], dropna=False)
        .size()
        .reset_index(name="N_REGISTROS")
        .sort_values("PARTICIPANTE")
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_all.to_excel(writer, sheet_name="datos_kubios_todos", index=False)
        resumen.to_excel(writer, sheet_name="resumen_por_sujeto", index=False)

    return output_path


def get_batch_status_excel_path():
    return get_global_tracking_dir() / "estado_ultimo_lote.xlsx"


def save_batch_status_excel(rows):
    output_path = get_batch_status_excel_path()
    df = pd.DataFrame(rows)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_excel(output_path, index=False)
    return output_path


def ejecutar_informe_seguro(tipo_informe_clave, tipo_informe_label, default_name, preconfig):
    try:
        ok = ejecutar_informe(
            tipo_informe_clave,
            tipo_informe_label,
            default_name=default_name,
            preconfig=preconfig,
        )
        return bool(ok), ""
    except Exception:
        return False, traceback.format_exc()


def build_total_results_from_global_summary(fecha_desde_lote=None):
    global_xlsx = get_global_summary_excel_path()
    global_all_data_xlsx = get_global_all_data_excel_path()
    out_dir = get_global_tracking_dir()
    fecha_archivo = pd.Timestamp.now().strftime("%Y_%m_%d")
    total_xlsx = out_dir / f"{fecha_archivo}_resultados_totales.xlsx"
    total_docx = out_dir / f"{fecha_archivo}_resultados_totales.docx"

    if not global_xlsx.exists():
        print("No se pudo generar resultados totales: todavía no existe el Excel global.")
        return None, None

    try:
        df = pd.read_excel(global_xlsx)
    except Exception as exc:
        print(f"No se pudo leer el Excel global para resultados totales: {exc}")
        return None, None

    if df is None or df.empty:
        print("No se pudo generar resultados totales: el Excel global está vacío.")
        return None, None

    df = df.copy()
    if "TIMESTAMP_REALIZACION" in df.columns:
        df["_TS_SORT"] = pd.to_datetime(df["TIMESTAMP_REALIZACION"], errors="coerce", dayfirst=True)
    else:
        df["_TS_SORT"] = pd.NaT
    df = df.sort_values(["PARTICIPANTE", "TIPO_INFORME", "_TS_SORT"], na_position="last")

    last_cols = [c for c in ["PARTICIPANTE", "TIPO_INFORME"] if c in df.columns]
    df_ultimos = df.drop_duplicates(last_cols, keep="last").copy() if last_cols else df.copy()
    df_export = df.drop(columns=["_TS_SORT"], errors="ignore")
    df_ultimos_export = df_ultimos.drop(columns=["_TS_SORT"], errors="ignore")

    resumen_cols = [
        "PARTICIPANTE", "CARPETA_PARTICIPANTE", "TIPO_INFORME", "FECHA_REALIZACION",
        "FECHA_DESDE", "ORDEN_REGISTRO", "N_REGISTROS_INFORME",
        "N_IMAGENES_TOTALES_EXCEL", "N_REGISTROS_NUEVOS", "WORD_GENERADO", "EXCEL_PARTICIPANTE",
    ]
    resumen_cols = [c for c in resumen_cols if c in df_ultimos_export.columns]
    df_resumen = df_ultimos_export[resumen_cols].copy() if resumen_cols else df_ultimos_export.copy()

    metric_rows = []
    for _, row in df_ultimos_export.iterrows():
        base = {
            "PARTICIPANTE": row.get("PARTICIPANTE"),
            "TIPO_INFORME": row.get("TIPO_INFORME"),
            "FECHA_REALIZACION": row.get("FECHA_REALIZACION"),
        }
        for var, meta in VARIABLE_META.items():
            pref = f"{var}__"
            if row.get(pref + "SELECCIONADA") != 1:
                continue
            metric_rows.append({
                **base,
                "VARIABLE": var,
                "NOMBRE_VARIABLE": meta.get("nombre", var),
                "PENDIENTE": row.get(pref + "PENDIENTE"),
                "CAMBIO_PCT": row.get(pref + "CAMBIO_PCT"),
                "MEDIA_RECIENTE": row.get(pref + "MEDIA_RECIENTE"),
                "MEDIA_REFERENCIA": row.get(pref + "MEDIA_REFERENCIA"),
                "DELTA_PCT": row.get(pref + "DELTA_PCT"),
                "RATIO": row.get(pref + "RATIO"),
                "ZSCORE": row.get(pref + "ZSCORE"),
                "SWC": row.get(pref + "SWC"),
                "SWC_RELEVANTE": row.get(pref + "SWC_RELEVANTE"),
                "LECTURA_EMOCIONAL": row.get(pref + "LECTURA_EMOCIONAL"),
            })
    df_metricas = pd.DataFrame(metric_rows)

    df_all_data = pd.DataFrame()
    if global_all_data_xlsx.exists():
        try:
            df_all_data = pd.read_excel(global_all_data_xlsx, sheet_name="datos_kubios_todos")
        except Exception:
            df_all_data = pd.DataFrame()

    with pd.ExcelWriter(total_xlsx, engine="openpyxl") as writer:
        df_resumen.to_excel(writer, sheet_name="resumen_sujetos", index=False)
        df_metricas.to_excel(writer, sheet_name="metricas_variables", index=False)
        df_all_data.to_excel(writer, sheet_name="datos_kubios_todos", index=False)
        df_export.to_excel(writer, sheet_name="historico_global", index=False)

    doc = Document()
    set_default_font(doc)
    add_footer_date(doc, pd.Timestamp.now().strftime("%d/%m/%Y"))
    doc.add_heading("Resultados totales", level=0)
    add_paragraph(doc, f"Fecha de realización: {pd.Timestamp.now().strftime('%d/%m/%Y %H:%M:%S')}")
    if fecha_desde_lote is not None:
        add_paragraph(doc, f"Periodo analizado desde: {fecha_desde_lote.strftime('%Y-%m-%d')}")
    add_paragraph(doc, f"Sujetos incluidos: {df_ultimos_export.get('PARTICIPANTE', pd.Series(dtype=object)).nunique()}")
    add_paragraph(doc, f"Informes incluidos en el resumen: {len(df_ultimos_export)}")
    add_paragraph(doc, f"Registros acumulados de todos los sujetos: {len(df_all_data)}")

    doc.add_heading("Resumen por sujeto", level=1)
    if df_resumen.empty:
        add_paragraph(doc, "No hay filas disponibles para resumir.")
    else:
        table_cols = [c for c in ["PARTICIPANTE", "TIPO_INFORME", "FECHA_REALIZACION", "N_REGISTROS_INFORME", "N_IMAGENES_TOTALES_EXCEL"] if c in df_resumen.columns]
        table = doc.add_table(rows=1, cols=len(table_cols))
        table.style = "Table Grid"
        for j, col in enumerate(table_cols):
            table.rows[0].cells[j].text = col
        for _, row in df_resumen[table_cols].iterrows():
            cells = table.add_row().cells
            for j, col in enumerate(table_cols):
                value = row.get(col)
                cells[j].text = "" if pd.isna(value) else str(value)

    doc.add_heading("Archivos generados", level=1)
    add_paragraph(doc, f"Excel total: {total_xlsx}")
    add_paragraph(doc, f"Excel acumulativo global de todos los datos: {global_all_data_xlsx}")
    add_paragraph(doc, f"Excel global acumulado: {global_xlsx}")
    doc.save(total_docx)

    return total_xlsx, total_docx



DEFAULT_ALARM_FREQ_DAYS = 7


def get_alarmas_excel_path():
    out_dir = SCRIPT_DIR / RESULTADOS_DIR_NAME / "Seguimiento_global"
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir / "alarmas_informes.xlsx"


def update_alarmas_excel(participante_carpeta, fecha_informe, hora_informe, default_freq_days=DEFAULT_ALARM_FREQ_DAYS):
    alarmas_path = get_alarmas_excel_path()
    cols = ["PARTICIPANTE", "FRECUENCIA_DIAS", "FECHA_ULTIMO_INFORME", "HORA_ULTIMO_INFORME"]
    if alarmas_path.exists():
        try:
            df = pd.read_excel(alarmas_path)
        except Exception:
            df = pd.DataFrame(columns=cols)
    else:
        df = pd.DataFrame(columns=cols)

    for col in cols:
        if col not in df.columns:
            df[col] = ""

    # Forzar tipos seguros para evitar errores al escribir texto en columnas vacías
    for col in ["PARTICIPANTE", "FECHA_ULTIMO_INFORME", "HORA_ULTIMO_INFORME"]:
        df[col] = df[col].astype("object")
        df[col] = df[col].where(~pd.isna(df[col]), "")

    if "FRECUENCIA_DIAS" in df.columns:
        df["FRECUENCIA_DIAS"] = pd.to_numeric(df["FRECUENCIA_DIAS"], errors="coerce")

    participante_carpeta = str(participante_carpeta).strip()
    fecha_informe = str(fecha_informe).strip()
    hora_informe = str(hora_informe).strip()
    mask = df["PARTICIPANTE"].astype(str).str.strip() == participante_carpeta

    if mask.any():
        idx = df.index[mask][0]
        freq_actual = df.at[idx, "FRECUENCIA_DIAS"]
        try:
            freq_actual = int(freq_actual)
        except Exception:
            freq_actual = int(default_freq_days)
        df.at[idx, "FRECUENCIA_DIAS"] = freq_actual
        df.at[idx, "FECHA_ULTIMO_INFORME"] = fecha_informe
        df.at[idx, "HORA_ULTIMO_INFORME"] = hora_informe
    else:
        df = pd.concat([df, pd.DataFrame([{
            "PARTICIPANTE": participante_carpeta,
            "FRECUENCIA_DIAS": int(default_freq_days),
            "FECHA_ULTIMO_INFORME": fecha_informe,
            "HORA_ULTIMO_INFORME": hora_informe,
        }])], ignore_index=True)

    df = df[cols].sort_values("PARTICIPANTE", key=lambda s: s.astype(str).str.lower()).reset_index(drop=True)
    df.to_excel(alarmas_path, index=False)
    return alarmas_path

def linear_trend(x, y):
    mask = np.isfinite(x) & np.isfinite(y)
    if mask.sum() < 2:
        return None
    x2 = x[mask].astype(float)
    y2 = y[mask].astype(float)
    slope, intercept = np.polyfit(x2, y2, 1)
    yhat = slope * x2 + intercept
    return {"slope": float(slope), "intercept": float(intercept), "x": x2, "y": y2, "yhat": yhat}


def pct_change(first_vals, last_vals):
    if len(first_vals) == 0 or len(last_vals) == 0:
        return None
    a = np.nanmean(first_vals)
    b = np.nanmean(last_vals)
    if not np.isfinite(a) or a == 0 or not np.isfinite(b):
        return None
    return ((b - a) / abs(a)) * 100.0


def magnitude_text(change_pct):
    if change_pct is None:
        return "sin magnitud cuantificable"
    abs_pct = abs(change_pct)
    if abs_pct < 5:
        return "cambio pequeño"
    if abs_pct < 15:
        return "cambio moderado"
    return "cambio marcado"


def coherent_message(slopes):
    hr_s = slopes.get("HR")
    rmssd_s = slopes.get("RMSSD")
    pns_s = slopes.get("PNS_INDEX")
    sns_s = slopes.get("SNS_INDEX")
    coherent_recovery = (
        rmssd_s is not None and rmssd_s > 0 and
        pns_s is not None and pns_s > 0 and
        ((hr_s is not None and hr_s < 0) or hr_s is None) and
        ((sns_s is not None and sns_s < 0) or sns_s is None)
    )
    coherent_stress = (
        ((rmssd_s is not None and rmssd_s < 0) or rmssd_s is None) and
        ((pns_s is not None and pns_s < 0) or pns_s is None) and
        ((hr_s is not None and hr_s > 0) or hr_s is None) and
        ((sns_s is not None and sns_s > 0) or sns_s is None)
    )
    if coherent_recovery:
        return "El patrón global muestra coherencia fisiológica con una evolución hacia mejor recuperación autonómica: aumento de indicadores vagales y descenso de activación relativa."
    if coherent_stress:
        return "El patrón global muestra coherencia fisiológica con mayor carga o peor recuperación: descenso de indicadores vagales y aumento de activación relativa."
    return "La coherencia entre HR, RMSSD, PNS y SNS no es completamente uniforme, por lo que la interpretación debe hacerse con prudencia y siempre considerando el contexto aplicado."


def coherent_state(slopes):
    hr_s = slopes.get("HR")
    rmssd_s = slopes.get("RMSSD")
    pns_s = slopes.get("PNS_INDEX")
    sns_s = slopes.get("SNS_INDEX")

    coherent_recovery = (
        rmssd_s is not None and rmssd_s > 0 and
        pns_s is not None and pns_s > 0 and
        ((hr_s is not None and hr_s < 0) or hr_s is None) and
        ((sns_s is not None and sns_s < 0) or sns_s is None)
    )
    coherent_stress = (
        ((rmssd_s is not None and rmssd_s < 0) or rmssd_s is None) and
        ((pns_s is not None and pns_s < 0) or pns_s is None) and
        ((hr_s is not None and hr_s > 0) or hr_s is None) and
        ((sns_s is not None and sns_s > 0) or sns_s is None)
    )

    if coherent_recovery:
        return "recuperacion"
    if coherent_stress:
        return "carga"
    return "mixto"


def automatic_interpretation(var, stats, slopes):
    if stats is None:
        return "No hay suficientes datos válidos para construir una interpretación automática robusta. Se recomienda ampliar el número de registros y mantener condiciones de medición estables."
    slope = stats.get("slope")
    change_pct = stats.get("change_pct")
    mag = magnitude_text(change_pct)
    if slope is None:
        trend_text = "sin tendencia lineal robusta"
    elif slope > 0:
        trend_text = "con tendencia positiva"
    elif slope < 0:
        trend_text = "con tendencia negativa"
    else:
        trend_text = "sin cambios lineales relevantes"

    if var in {"RMSSD", "LnRMSSD", "SDNN", "SD1", "PNS_INDEX"}:
        base = f"La variable {trend_text} muestra un {mag}."
        detail = "Desde el punto de vista fisiológico, esto es compatible con una mejor modulación vagal o una recuperación más favorable." if slope and slope > 0 else ("Desde el punto de vista fisiológico, esto puede ser compatible con fatiga acumulada, estrés fisiológico o menor recuperación." if slope and slope < 0 else "La estabilidad sugiere un comportamiento relativamente constante del sistema en el periodo analizado.")
    elif var in {"HR", "STRESS_INDEX", "SNS_INDEX", "LF_HF"}:
        base = f"La variable {trend_text} muestra un {mag}."
        detail = "Desde el punto de vista fisiológico, esto puede reflejar mayor activación relativa o una respuesta compatible con mayor carga interna." if slope and slope > 0 else ("Desde el punto de vista fisiológico, esto puede ser compatible con una reducción de activación fisiológica relativa y mejor recuperación." if slope and slope < 0 else "La estabilidad sugiere ausencia de cambios relevantes en la activación relativa durante el periodo analizado.")
    else:
        base = f"La variable {trend_text} muestra un {mag}."
        detail = "Su interpretación debe integrarse con el resto de indicadores autonómicos."
    return f"{base} {detail} {coherent_message(slopes)}"


def recommendations_from_auto(var, stats):
    if stats is None:
        return "Se recomienda ampliar el seguimiento y revisar la calidad de los registros antes de tomar decisiones aplicadas."
    slope = stats.get("slope")
    if var in {"RMSSD", "LnRMSSD", "SDNN", "SD1", "PNS_INDEX"} and slope is not None and slope < 0:
        return "Recomendación: valorar una reducción temporal de la carga, reforzar el sueño y revisar factores de estrés externo antes de incrementar la exigencia."
    if var in {"HR", "STRESS_INDEX", "SNS_INDEX", "LF_HF"} and slope is not None and slope > 0:
        return "Recomendación: monitorizar de cerca la evolución en los próximos registros y considerar ajustes de carga o de recuperación si el patrón persiste."
    return "Recomendación: mantener la monitorización longitudinal y contrastar estos cambios con la percepción subjetiva, el rendimiento y la carga semanal."


def make_graph(df, var, out_png, momento, short_window=7, long_window=28):
    short_window, long_window = safe_window_pair(short_window, long_window)
    sub = df[["ORDEN_DIA", "FECHA_IMAGEN", var]].copy()
    sub[var] = pd.to_numeric(sub[var], errors="coerce")
    sub = sub.dropna()

    fig, ax = plt.subplots(figsize=(8.2, 4.6))
    ax.set_title(VARIABLE_META.get(var, {}).get("nombre", var), fontweight="bold")

    if str(momento) == "1":
        x = np.arange(len(sub), dtype=float)
        x_labels = sub["FECHA_IMAGEN"].astype(str).tolist()
        ax.set_xlabel("Fecha")
    else:
        x = sub["ORDEN_DIA"].to_numpy(dtype=float)
        x_labels = None
        ax.set_xlabel("Orden de registro")

    ax.set_ylabel(var)
    ax.grid(True, alpha=0.25)

    if len(sub) == 0:
        ax.text(0.5, 0.5, "Sin datos válidos", ha="center", va="center", transform=ax.transAxes)
        fig.tight_layout()
        fig.savefig(out_png, bbox_inches="tight")
        plt.close(fig)
        return None

    y = sub[var].to_numpy(dtype=float)
    # Mantener el comportamiento visual original del gráfico: suavizado ligero de 3 registros.
    ma_visual = moving_average(sub[var], 3, min_periods=1)
    ax.plot(x, ma_visual, color="black", linewidth=1.8, marker="o", markersize=5.5,
            markerfacecolor="white", markeredgecolor="black", markeredgewidth=1.1)

    for xi, yi in zip(x, y):
        ax.annotate(f"{yi:.2f}", (xi, yi), textcoords="offset points", xytext=(0, 8), ha="center", fontsize=8)

    tr = linear_trend(x, y)
    eq_text = "Tendencia no disponible"
    slope = None
    intercept = None
    if tr is not None:
        slope = tr["slope"]
        intercept = tr["intercept"]
        trend_color = "green" if slope > 0 else "red" if slope < 0 else "black"
        ax.plot(tr["x"], tr["yhat"], color=trend_color, linewidth=2.0, linestyle="--")
        eq_text = f"y = {slope:.4f}x + {intercept:.4f}"

    if x_labels is not None:
        ax.set_xticks(x)
        ax.set_xticklabels(x_labels, rotation=45, ha="right")

    ax.text(0.99, 0.03, eq_text, transform=ax.transAxes, ha="right", va="bottom", fontsize=9,
            bbox=dict(boxstyle="round,pad=0.25", facecolor="white", edgecolor="black", alpha=0.9))

    fig.tight_layout()
    fig.savefig(out_png, bbox_inches="tight")
    plt.close(fig)

    first_vals = y[:min(3, len(y))]
    last_vals = y[-min(3, len(y)):]
    comparison = compute_window_comparison(sub[var], short_window=short_window, long_window=long_window)
    return {
        "slope": slope,
        "intercept": intercept,
        "mean": float(np.nanmean(y)) if len(y) else None,
        "change_pct": pct_change(first_vals, last_vals),
        "short_window": short_window,
        "long_window": long_window,
        "recent_mean": comparison.get("short_mean"),
        "baseline_mean": comparison.get("long_mean"),
        "recent_vs_baseline_abs": comparison.get("delta_abs"),
        "recent_vs_baseline_pct": comparison.get("delta_pct"),
        "window_comparison": comparison,
    }


def generate_graphs(df, momento, selected_vars, graph_dir, short_window=7, long_window=28):
    short_window, long_window = safe_window_pair(short_window, long_window)
    graph_dir.mkdir(parents=True, exist_ok=True)
    graph_info = {}
    for var in selected_vars:
        if var not in df.columns:
            continue
        png = graph_dir / f"{var.lower()}.png"
        graph_info[var] = {
            "png": str(png),
            "stats": make_graph(df, var, png, momento, short_window=short_window, long_window=long_window)
        }
    return graph_info


def add_paragraph(doc, text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=10):
    p = doc.add_paragraph()
    p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    return p


def style_table_font_8(table):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)


def generate_session_report(df, graph_info):
    n = len(df)
    if n == 0:
        return "No se dispone de sesiones analizadas para elaborar un informe final."
    ok_n = int((df["ESTADO_EXTRACCION"] == "OK").sum()) if "ESTADO_EXTRACCION" in df.columns else 0
    slopes = {k: (graph_info.get(k, {}).get("stats") or {}).get("slope") for k in ["HR", "RMSSD", "LnRMSSD", "PNS_INDEX", "SNS_INDEX"]}
    pattern = []
    for key in ["RMSSD", "LnRMSSD", "HR", "PNS_INDEX", "SNS_INDEX"]:
        s = slopes.get(key)
        if s is None:
            continue
        estado = "ascendente" if s > 0 else "descendente" if s < 0 else "estable"
        pattern.append(f"{key} {estado}")
    pattern_text = ", ".join(pattern) if pattern else "sin patrón principal claramente identificable"
    return (
        f"Se analizaron {n} sesiones, de las cuales {ok_n} presentaron una extracción marcada como OK. "
        f"En conjunto, el perfil longitudinal muestra {pattern_text}. {coherent_message(slopes)} "
        f"Desde un punto de vista aplicado, este patrón debe interpretarse junto con la carga reciente, la percepción subjetiva, "
        f"la calidad del sueño y las condiciones en las que se realizaron las mediciones. "
        f"El seguimiento repetido permite detectar cambios con más valor práctico que la interpretación de un único registro aislado."
    )


def participant_overall_status(slopes):
    hr_s = slopes.get("HR")
    rmssd_s = slopes.get("RMSSD")
    pns_s = slopes.get("PNS_INDEX")
    stress_s = slopes.get("STRESS_INDEX")
    sns_s = slopes.get("SNS_INDEX")

    favorable = 0
    unfavorable = 0

    if hr_s is not None:
        favorable += int(hr_s < 0)
        unfavorable += int(hr_s > 0)
    if rmssd_s is not None:
        favorable += int(rmssd_s > 0)
        unfavorable += int(rmssd_s < 0)
    if pns_s is not None:
        favorable += int(pns_s > 0)
        unfavorable += int(pns_s < 0)
    if stress_s is not None:
        favorable += int(stress_s < 0)
        unfavorable += int(stress_s > 0)
    if sns_s is not None:
        favorable += int(sns_s < 0)
        unfavorable += int(sns_s > 0)

    if favorable >= max(2, unfavorable + 1):
        return "favorable"
    if unfavorable >= max(2, favorable + 1):
        return "desfavorable"
    return "estable"


def participant_summary_text(slopes):
    status = participant_overall_status(slopes)
    if status == "favorable":
        return (
            "Durante este periodo, tus registros muestran una evolución global compatible con una mejor recuperación. "
            "En conjunto, varios de tus indicadores se mueven en una dirección positiva y eso sugiere que tu cuerpo podría estar respondiendo bien al programa de ejercicio."
        )
    if status == "desfavorable":
        return (
            "Durante este periodo, tus registros muestran una evolución compatible con una recuperación algo menos favorable. "
            "Esto no significa necesariamente que algo vaya mal, pero sí que conviene prestar atención a cómo te sientes y a cómo estás tolerando la carga de ejercicio."
        )
    return (
        "Durante este periodo, tus registros muestran una evolución bastante estable. "
        "No se observan cambios muy marcados en conjunto, lo que sugiere un comportamiento relativamente constante en tus mediciones."
    )


def participant_meaning_text(slopes):
    status = participant_overall_status(slopes)
    if status == "favorable":
        return (
            "En conjunto, estos cambios sugieren que tu cuerpo podría estar adaptándose bien al programa de ejercicio. "
            "Esto es una señal positiva, ya que apunta a una posible mejora en tu capacidad de recuperación y en el equilibrio general de tu organismo."
        )
    if status == "desfavorable":
        return (
            "En conjunto, estos cambios sugieren que tu cuerpo podría necesitar algo más de tiempo para recuperarse o que estás acumulando cierta carga. "
            "No debe interpretarse como algo alarmante, pero sí como una señal útil para seguir observando tu evolución con calma."
        )
    return (
        "En conjunto, estos cambios sugieren una situación bastante estable. "
        "Eso puede ser una buena señal cuando va acompañado de buenas sensaciones, aunque siempre es importante seguir observando la evolución con el paso de las semanas."
    )


def participant_next_steps_text(slopes):
    status = participant_overall_status(slopes)
    if status == "favorable":
        return (
            "Seguiremos observando tus registros en el tiempo y los compararemos con cómo te sientes, para adaptar el programa de la forma más adecuada a ti."
        )
    if status == "desfavorable":
        return (
            "Seguiremos observando tus registros en el tiempo y los compararemos con cómo te sientes, para ajustar el programa de forma progresiva y adaptada a ti."
        )
    return (
        "Seguiremos observando tus registros en el tiempo y los compararemos con cómo te sientes, para confirmar si esta estabilidad se mantiene y adaptar el programa cuando sea necesario."
    )


def build_participant_highlights(slopes):
    mensajes = []
    if slopes.get("HR") is not None:
        if slopes["HR"] < 0:
            mensajes.append("Tu frecuencia cardiaca en reposo tiende a disminuir ligeramente.")
        elif slopes["HR"] > 0:
            mensajes.append("Tu frecuencia cardiaca en reposo tiende a aumentar ligeramente.")
    if slopes.get("PNS_INDEX") is not None:
        if slopes["PNS_INDEX"] > 0:
            mensajes.append("El indicador de recuperación del sistema nervioso muestra una tendencia favorable.")
        elif slopes["PNS_INDEX"] < 0:
            mensajes.append("El indicador de recuperación del sistema nervioso muestra una tendencia algo menos favorable.")
    if slopes.get("STRESS_INDEX") is not None:
        if slopes["STRESS_INDEX"] < 0:
            mensajes.append("Además, el índice de estrés fisiológico ha descendido.")
        elif slopes["STRESS_INDEX"] > 0:
            mensajes.append("Además, el índice de estrés fisiológico ha aumentado ligeramente.")
    if not mensajes:
        mensajes.append("En este periodo no se aprecian cambios muy marcados en los indicadores principales.")
    return mensajes


def add_bullet_list(doc, items, size=10):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(item)
        r.font.size = Pt(size)


# ===========================
# MÓDULO EMOCIONAL BASADO EN DATOS
# ===========================

def safe_numeric_series(df, col):
    if df is None or df.empty or col not in df.columns:
        return pd.Series(dtype=float)
    return pd.to_numeric(df[col], errors="coerce").dropna()


def emotional_data_validity(df):
    resultado = {
        "es_valido": False,
        "n_registros": 0,
        "n_ok": 0,
        "pct_ok": 0.0,
        "calidad_favorable_pct": None,
        "motivos": [],
    }

    if df is None or df.empty:
        resultado["motivos"].append("No hay registros en el periodo seleccionado.")
        return resultado

    n_registros = len(df)
    n_ok = int((df["ESTADO_EXTRACCION"] == "OK").sum()) if "ESTADO_EXTRACCION" in df.columns else 0
    pct_ok = (n_ok / n_registros * 100.0) if n_registros > 0 else 0.0

    resultado["n_registros"] = n_registros
    resultado["n_ok"] = n_ok
    resultado["pct_ok"] = pct_ok

    if "MEASUREMENT_QUALITY" in df.columns:
        q = df["MEASUREMENT_QUALITY"].fillna("").astype(str).str.upper().str.strip()
        if len(q) > 0:
            favorables = q.isin(["GOOD", "ACCEPTABLE", "OK"]).sum()
            resultado["calidad_favorable_pct"] = float(favorables / len(q) * 100.0)

    if n_registros < 4:
        resultado["motivos"].append("Hay menos de 4 registros válidos para una lectura longitudinal prudente.")
    if pct_ok < 70:
        resultado["motivos"].append("Menos del 70% de los registros presentan extracción marcada como OK.")

    resultado["es_valido"] = len(resultado["motivos"]) == 0
    return resultado


def classify_autonomic_state(slopes):
    vars_consideradas = {
        "HR": ("down", "up"),
        "RMSSD": ("up", "down"),
        "PNS_INDEX": ("up", "down"),
        "SNS_INDEX": ("down", "up"),
        "STRESS_INDEX": ("down", "up"),
    }

    favorable = 0
    unfavorable = 0
    disponibles = 0

    for var, (fav_dir, unf_dir) in vars_consideradas.items():
        s = slopes.get(var)
        if s is None or not np.isfinite(s):
            continue
        disponibles += 1
        if s > 0:
            direction = "up"
        elif s < 0:
            direction = "down"
        else:
            direction = "flat"

        favorable += int(direction == fav_dir)
        unfavorable += int(direction == unf_dir)

    if disponibles < 2:
        return "incierto", favorable, unfavorable, disponibles

    score = favorable - unfavorable
    if score >= 2:
        return "favorable", favorable, unfavorable, disponibles
    if score <= -2:
        return "desfavorable", favorable, unfavorable, disponibles
    return "estable", favorable, unfavorable, disponibles


def emotional_confidence(df, slopes):
    validity = emotional_data_validity(df)
    base = 0

    n_registros = validity["n_registros"]
    if n_registros >= 8:
        base += 35
    elif n_registros >= 6:
        base += 28
    elif n_registros >= 4:
        base += 20

    pct_ok = validity["pct_ok"]
    if pct_ok >= 90:
        base += 30
    elif pct_ok >= 80:
        base += 24
    elif pct_ok >= 70:
        base += 18

    estado_coherencia = coherent_state(slopes)
    if estado_coherencia in {"recuperacion", "carga"}:
        base += 25
    else:
        base += 10

    _, _, _, disponibles = classify_autonomic_state(slopes)
    if disponibles >= 5:
        base += 10
    elif disponibles >= 3:
        base += 6

    return min(100, int(round(base)))


def confidence_label(score):
    if score >= 80:
        return "alta"
    if score >= 60:
        return "moderada"
    return "baja"


def generate_emotional_message(slopes, df):
    validity = emotional_data_validity(df)
    if not validity["es_valido"]:
        frase = "Todavía no hay suficientes datos consistentes para describir tu evolución con claridad."
        explicacion = (
            "Las mediciones disponibles en este periodo no alcanzan un nivel suficiente de estabilidad para emitir un mensaje emocional basado en datos. "
            "Lo más prudente es seguir acumulando registros en condiciones similares antes de extraer una lectura global."
        )
        apoyo = "Esto no implica una evolución negativa; solo indica que aún no hay base suficiente para resumirla con solidez."
        confianza = 0
        return {
            "frase": frase,
            "explicacion": explicacion,
            "apoyo": apoyo,
            "confianza": confianza,
            "confianza_texto": "insuficiente",
            "estado": "incierto",
            "validez": validity,
        }

    estado, favorable, unfavorable, disponibles = classify_autonomic_state(slopes)
    confianza = emotional_confidence(df, slopes)
    confianza_texto = confidence_label(confianza)
    coherencia = coherent_state(slopes)

    if estado == "favorable":
        frase = "Tus registros muestran una evolución consistente hacia una mejor recuperación en este periodo."
        if coherencia == "recuperacion":
            explicacion = (
                "Varios de los indicadores principales siguen una dirección compatible con una respuesta fisiológica más favorable, "
                "con señales de mejor recuperación autonómica en el conjunto del periodo analizado."
            )
        else:
            explicacion = (
                "Aunque no todas las variables se mueven exactamente en la misma dirección, el balance global de los datos apunta a una evolución fisiológica favorable."
            )
        apoyo = (
            "Esto no sustituye una valoración clínica ni describe por sí solo tu estado de salud general, "
            "pero sí refleja una señal real y medible de mejor recuperación dentro de los registros analizados."
        )

    elif estado == "desfavorable":
        frase = "En este periodo, tus registros indican que tu organismo podría estar necesitando más tiempo para recuperarse."
        if coherencia == "carga":
            explicacion = (
                "El patrón conjunto de las variables es compatible con mayor carga fisiológica relativa o una recuperación menos favorable "
                "durante este tramo del seguimiento."
            )
        else:
            explicacion = (
                "Algunas variables relevantes se desplazan en una dirección compatible con más activación o menor recuperación, "
                "aunque siempre deben interpretarse con prudencia y junto con cómo te sientes."
            )
        apoyo = (
            "Esto no debe leerse como algo alarmante ni como un juicio sobre tu proceso; "
            "leer a tiempo esta señal también es una forma de cuidar y adaptar mejor el programa."
        )

    elif estado == "estable":
        frase = "Tus registros muestran una evolución bastante estable a lo largo del periodo analizado."
        explicacion = (
            "No se observan cambios muy marcados en el conjunto de los indicadores principales, "
            "lo que sugiere un comportamiento fisiológico relativamente constante entre las mediciones seleccionadas."
        )
        apoyo = (
            "Esa estabilidad también tiene valor, porque ayuda a entender tu proceso con más calma y a detectar con mayor claridad cualquier cambio futuro."
        )

    else:
        frase = "Los datos disponibles no permiten describir una tendencia clara en este periodo."
        explicacion = (
            "Las variables analizadas no muestran un patrón suficientemente consistente entre sí, "
            "o bien la información disponible no es todavía lo bastante completa para resumir una dirección global con seguridad."
        )
        apoyo = (
            "Seguir acumulando mediciones en condiciones similares será lo más útil para construir una lectura más sólida de tu evolución."
        )

    detalle = f"Indicadores valorados: {disponibles}. Señales favorables: {favorable}. Señales menos favorables: {unfavorable}."
    return {
        "frase": frase,
        "explicacion": explicacion,
        "apoyo": apoyo,
        "detalle": detalle,
        "confianza": confianza,
        "confianza_texto": confianza_texto,
        "estado": estado,
        "validez": validity,
    }



def set_table_borders_none(table):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = tbl_borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tbl_borders.append(element)
        element.set(qn("w:val"), "nil")



def add_intro_block_with_optional_photo(doc, tipo_informe, nombre_sujeto, orden_registro, fecha_desde, fecha_informe, foto_participante_path=None, short_window=None, long_window=None, **kwargs):
    # Parámetros extra admitidos para mantener compatibilidad con futuras ampliaciones del informe.
    textos = [
        f"Tipo de informe: {tipo_informe}",
        f"Nombre: {nombre_sujeto}",
        f"Orden del registro dentro del día: {orden_registro}",
    ]
    if fecha_desde:
        textos.append(f"Periodo analizado desde: {fecha_desde}")
    textos.append(f"Fecha de realización: {fecha_informe}")

    foto_ok = bool(foto_participante_path) and Path(foto_participante_path).exists()

    # Siempre construir el bloque con tabla invisible para mantener una maquetación estable
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders_none(table)

    try:
        table.columns[0].width = Inches(5.7)
        table.columns[1].width = Inches(1.7)
    except Exception:
        pass

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)

    for idx, linea in enumerate(textos):
        p = left_cell.paragraphs[0] if idx == 0 else left_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(linea)
        run.font.size = Pt(12)

    if foto_ok:
        p_img = right_cell.paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_img = p_img.add_run()
        run_img.add_picture(str(foto_participante_path), width=Inches(1.45))
        return True

    # Si no hay foto, dejar la celda derecha vacía y seguir sin error
    right_cell.paragraphs[0].text = ""
    return False


def find_named_image(images_dir, base_name):
    images_dir = Path(images_dir)
    if not images_dir.exists() or not images_dir.is_dir():
        return None
    target = str(base_name).strip().lower()
    candidates = []
    for ext in IMAGE_EXTENSIONS:
        candidates.extend(images_dir.glob(ext))
    for path in sorted(candidates):
        if path.stem.strip().lower() == target:
            return path
    return None



def get_shared_images_dir():
    return SHARED_IMAGES_DIR if SHARED_IMAGES_DIR.exists() and SHARED_IMAGES_DIR.is_dir() else None



def resolve_logo_path():
    shared_dir = get_shared_images_dir()
    if shared_dir is None:
        return None
    return find_named_image(shared_dir, "logo_catedra")



def resolve_participant_photo_path():
    shared_dir = get_shared_images_dir()
    if shared_dir is None:
        return None
    participant_folder_name = BASE_DIR.name.strip()
    if not participant_folder_name:
        return None
    return find_named_image(shared_dir, participant_folder_name)


def generate_word(df, summary_df, nombre_sujeto, orden_registro, graph_info, selected_vars, fecha_informe, output_docx, tipo_informe, fecha_desde=None, definition_mode="directa", logo_path=None, foto_participante_path=None, short_window=7, long_window=28, selected_blocks=None, swc_factor=0.5):
    doc = Document()
    set_margins(doc)
    logo_inserted = add_document_header(doc, "Cátedra de Ejercicio, Educación y Cáncer. San Javier", logo_path=logo_path)
    add_footer_date(doc, fecha_informe)

    styles = doc.styles["Normal"]
    styles.font.name = "Calibri"
    styles.font.size = Pt(10)

    if selected_blocks is None:
        selected_blocks = ["grafico", "interpretacion", "recomendaciones"]

    all_blocks_order = [
        "grafico",
        "interpretacion",
        "recomendaciones",
        "comparacion_ventanas",
        "ratio",
        "zscore_swc",
        "tabla_interpretacion",
        "lectura_emocional",
    ]
    selected_blocks = [b for b in all_blocks_order if b in list(dict.fromkeys(selected_blocks))]

    es_participante = "participante" in str(tipo_informe).lower()
    short_window, long_window = safe_window_pair(short_window, long_window)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo = "Informe de tu evolución" if es_participante else "Análisis Longitudinal de la Variabilidad de la Frecuencia Cardiaca"
    r = p.add_run(titulo)
    r.bold = True
    r.font.size = Pt(18)

    foto_insertada = add_intro_block_with_optional_photo(
        doc,
        tipo_informe=tipo_informe,
        nombre_sujeto=nombre_sujeto,
        orden_registro=orden_registro,
        fecha_desde=fecha_desde,
        fecha_informe=fecha_informe,
        foto_participante_path=foto_participante_path,
        short_window=short_window,
        long_window=long_window,
        selected_blocks=selected_blocks,
        swc_factor=swc_factor,
    )

    if not summary_df.empty:
        rs = summary_df.iloc[0]
        if es_participante:
            add_paragraph(doc, f"En este informe se han incluido {int(rs.get('N_IMAGENES', 0))} registros válidos para ayudarte a entender mejor tu evolución.", align=WD_ALIGN_PARAGRAPH.LEFT)
        else:
            add_paragraph(doc, f"Resumen: {int(rs.get('N_IMAGENES', 0))} imágenes, {int(rs.get('N_OK', 0))} OK, {int(rs.get('N_REVISAR', 0))} para revisar y {int(rs.get('N_NUEVAS_EN_ESTA_EJECUCION', 0))} nuevas en esta ejecución.", align=WD_ALIGN_PARAGRAPH.LEFT)

    if any(x in selected_blocks for x in ["comparacion_ventanas", "ratio", "zscore_swc"]):
        add_paragraph(doc, f"Análisis por ventanas activado: referencia reciente de {short_window} registros y referencia amplia de {long_window} registros.", align=WD_ALIGN_PARAGRAPH.LEFT)
    if "zscore_swc" in selected_blocks:
        add_paragraph(doc, f"SWC activado con factor {swc_factor:.2f} sobre la desviación estándar de la ventana amplia.", align=WD_ALIGN_PARAGRAPH.LEFT)
    if not selected_blocks:
        add_paragraph(doc, "No se han seleccionado bloques analíticos opcionales para las variables. El informe se limitará a la estructura general y a la información básica disponible.", align=WD_ALIGN_PARAGRAPH.LEFT)

    slopes = {var: (graph_info.get(var, {}).get("stats") or {}).get("slope") for var in selected_vars}
    emotional_pack = generate_emotional_message(slopes, df) if es_participante else None

    definitions_added = add_selected_variable_definitions_section(
        doc,
        selected_vars=selected_vars,
        definition_mode=definition_mode,
        section_number=1,
    )
    next_heading_offset = 1 if definitions_added else 0

    def render_variable_block(doc, var, heading_prefix):
        if var not in VARIABLE_META or var not in df.columns:
            return
        meta = VARIABLE_META[var]
        doc.add_heading(f"{heading_prefix}. {meta['nombre']}", level=2)
        info = graph_info.get(var, {})
        png = info.get("png")
        stats = info.get("stats") or {}
        metrics = compute_reference_metrics(df[var], short_window=short_window, long_window=long_window, swc_factor=swc_factor)
        comparison = metrics.get("comparison") or {}

        if not es_participante:
            add_paragraph(doc, f"Explicación: {meta['explicacion']}")
            add_paragraph(doc, f"Interpretación general: {meta['interpretacion']}")

        if "grafico" in selected_blocks:
            add_paragraph(doc, f"{block_title_text('grafico')}: {block_intro_text('grafico', short_window=short_window, long_window=long_window, swc_factor=swc_factor)}", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
            if png and Path(png).exists():
                doc.add_picture(str(png), width=Inches(6.3 if es_participante else 6.4))
                add_paragraph(doc, "Figura. Evolución temporal de la variable con su línea de tendencia.", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
            else:
                add_paragraph(doc, block_unavailable_text('grafico', var, stats=stats, metrics=metrics, short_window=short_window, long_window=long_window))

        if "interpretacion" in selected_blocks:
            add_paragraph(doc, f"{block_title_text('interpretacion')}: {block_intro_text('interpretacion', short_window=short_window, long_window=long_window, swc_factor=swc_factor)}", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
            if stats and stats.get("slope") is not None:
                texto = automatic_interpretation(var, stats, slopes)
                if es_participante:
                    texto = texto.replace("Desde el punto de vista fisiológico, ", "En términos sencillos, ")
                add_paragraph(doc, texto)
            else:
                add_paragraph(doc, block_unavailable_text('interpretacion', var, stats=stats, metrics=metrics, short_window=short_window, long_window=long_window))

        if "recomendaciones" in selected_blocks:
            add_paragraph(doc, f"{block_title_text('recomendaciones')}: {block_intro_text('recomendaciones', short_window=short_window, long_window=long_window, swc_factor=swc_factor)}", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
            if stats and stats.get("slope") is not None:
                rec = recommendations_from_auto(var, stats)
                if es_participante:
                    rec = rec.replace("Recomendación: ", "")
                add_paragraph(doc, rec)
            else:
                add_paragraph(doc, block_unavailable_text('recomendaciones', var, stats=stats, metrics=metrics, short_window=short_window, long_window=long_window))

        if "comparacion_ventanas" in selected_blocks:
            add_paragraph(doc, f"{block_title_text('comparacion_ventanas')}: {block_intro_text('comparacion_ventanas', short_window=short_window, long_window=long_window, swc_factor=swc_factor)}", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
            if comparison.get("short_mean") is not None and comparison.get("long_mean") is not None:
                add_paragraph(doc, f"Comparación entre ventanas: {interpret_window_comparison(var, comparison)}")
            else:
                add_paragraph(doc, block_unavailable_text('comparacion_ventanas', var, stats=stats, metrics=metrics, short_window=short_window, long_window=long_window))

        if "ratio" in selected_blocks:
            add_paragraph(doc, f"{block_title_text('ratio')}: {block_intro_text('ratio', short_window=short_window, long_window=long_window, swc_factor=swc_factor)}", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
            ratio = metrics.get("ratio")
            if ratio is not None and np.isfinite(ratio):
                ratio_txt = format_metric_value(ratio, 3)
                add_paragraph(doc, f"Ratio reciente/referencia: {ratio_txt}. {interpret_ratio(var, ratio)}")
            else:
                add_paragraph(doc, block_unavailable_text('ratio', var, stats=stats, metrics=metrics, short_window=short_window, long_window=long_window))

        if "zscore_swc" in selected_blocks:
            add_paragraph(doc, f"{block_title_text('zscore_swc')}: {block_intro_text('zscore_swc', short_window=short_window, long_window=long_window, swc_factor=swc_factor)}", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
            if metrics.get("zscore") is not None or metrics.get("swc") is not None:
                z_txt = format_metric_value(metrics.get("zscore"), 2)
                swc_txt = format_metric_value(metrics.get("swc"), 2)
                add_paragraph(doc, f"Z-score individual: {z_txt}. SWC: {swc_txt}. {interpret_zscore_swc(var, metrics)}")
            else:
                add_paragraph(doc, block_unavailable_text('zscore_swc', var, stats=stats, metrics=metrics, short_window=short_window, long_window=long_window))

        if "tabla_interpretacion" in selected_blocks:
            add_paragraph(doc, f"{block_title_text('tabla_interpretacion')}: {block_intro_text('tabla_interpretacion', short_window=short_window, long_window=long_window, swc_factor=swc_factor)}", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
            rows = build_interpretation_table_rows(var, stats, metrics, selected_blocks)
            add_interpretation_table(doc, rows)

        if "lectura_emocional" in selected_blocks:
            add_paragraph(doc, f"{block_title_text('lectura_emocional')}: {block_intro_text('lectura_emocional', short_window=short_window, long_window=long_window, swc_factor=swc_factor)}", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
            if (stats and stats.get("slope") is not None) or comparison.get("delta_pct") is not None:
                add_paragraph(doc, emotional_interpretation_for_variable(var, stats, metrics))
            else:
                add_paragraph(doc, block_unavailable_text('lectura_emocional', var, stats=stats, metrics=metrics, short_window=short_window, long_window=long_window))

    if not es_participante:
        doc.add_heading(f"{1 + next_heading_offset}. Tabla básica de registros", level=1)
        basic_cols = [c for c in ["FILE", "FECHA_IMAGEN", "HORA_IMAGEN", "ORDEN_DIA"] if c in df.columns]
        if basic_cols:
            table = doc.add_table(rows=1, cols=len(basic_cols))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = table.rows[0].cells
            for i, col in enumerate(basic_cols):
                hdr[i].text = {"FILE": "Nombre imagen", "FECHA_IMAGEN": "Fecha", "HORA_IMAGEN": "Hora", "ORDEN_DIA": "Orden"}.get(col, col)
            for _, row in df.iterrows():
                cells = table.add_row().cells
                for i, col in enumerate(basic_cols):
                    value = row.get(col, "")
                    cells[i].text = "" if pd.isna(value) else str(value)
            style_table_font_8(table)

        doc.add_heading(f"{2 + next_heading_offset}. Variables seleccionadas", level=1)
        if not selected_blocks:
            add_paragraph(doc, "No se seleccionaron bloques de análisis para las variables.")
        for idx, var in enumerate(selected_vars, start=1):
            render_variable_block(doc, var, f"{2 + next_heading_offset}.{idx}")

        doc.add_heading(f"{3 + next_heading_offset}. Informe de las sesiones analizadas", level=1)
        add_paragraph(doc, generate_session_report(df, graph_info))

        doc.add_heading(f"{4 + next_heading_offset}. Referencias", level=1)
        refs_used, seen = [], set()
        for var in selected_vars:
            if var not in VARIABLE_META:
                continue
            for ref_key in VARIABLE_META[var]["refs"]:
                if ref_key not in seen:
                    seen.add(ref_key)
                    refs_used.append(APA_REFERENCES[ref_key])
        for ref in refs_used:
            add_paragraph(doc, ref, align=WD_ALIGN_PARAGRAPH.LEFT)
    else:
        doc.add_heading(f"{1 + next_heading_offset}. Mensaje sobre tu evolución", level=1)
        add_paragraph(doc, emotional_pack["frase"], bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, size=11)
        add_paragraph(doc, emotional_pack["explicacion"])
        add_paragraph(doc, emotional_pack["apoyo"])
        if emotional_pack.get("detalle"):
            add_paragraph(doc, emotional_pack["detalle"], italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, size=9)
        add_paragraph(doc, f"Nivel de confianza del mensaje: {emotional_pack['confianza_texto']} ({emotional_pack['confianza']}%).", align=WD_ALIGN_PARAGRAPH.LEFT)

        doc.add_heading(f"{2 + next_heading_offset}. Resumen de tu evolución", level=1)
        add_paragraph(doc, participant_summary_text(slopes))
        for msg in build_participant_highlights(slopes):
            add_paragraph(doc, msg, align=WD_ALIGN_PARAGRAPH.LEFT)

        doc.add_heading(f"{3 + next_heading_offset}. ¿Qué significa esto?", level=1)
        add_paragraph(doc, participant_meaning_text(slopes))

        doc.add_heading(f"{4 + next_heading_offset}. Importante tener en cuenta", level=1)
        add_paragraph(doc, "Estos resultados no deben interpretarse de forma aislada. Es importante considerar también:")
        add_bullet_list(doc, [
            "Cómo te sientes en tu día a día.",
            "La calidad de tu sueño.",
            "Tu nivel de cansancio o fatiga.",
            "La presencia de dolor o molestias.",
            "El estrés emocional.",
            "Las condiciones en las que realizaste la medición.",
        ])
        add_paragraph(doc, "Un valor puntual puede variar por muchos factores. Por eso, lo más importante es observar la evolución a lo largo de varias semanas.")

        doc.add_heading(f"{5 + next_heading_offset}. ¿Qué haremos a partir de ahora?", level=1)
        add_paragraph(doc, participant_next_steps_text(slopes))
        add_paragraph(doc, "Te recomendamos:")
        add_bullet_list(doc, [
            "Realizar las mediciones en condiciones similares siempre que sea posible.",
            "Comunicar cómo te encuentras: energía, sueño, molestias y estado de ánimo.",
            "No preocuparte por cambios puntuales en los valores.",
        ])

        if selected_vars:
            doc.add_heading(f"{6 + next_heading_offset}. Evolución de tus variables registradas", level=1)
            if not selected_blocks:
                add_paragraph(doc, "No se seleccionaron bloques de análisis para las variables.")
            for idx, var in enumerate(selected_vars, start=1):
                render_variable_block(doc, var, f"{6 + next_heading_offset}.{idx}")

    if logo_path:
        print(f"Logo en encabezado: {'insertado' if logo_inserted else 'detectado pero no insertado'}")
    else:
        print("Logo en encabezado: no disponible")

    if foto_participante_path:
        print(f"Foto del participante en portada: {'insertada' if foto_insertada else 'detectada pero no insertada'}")
    else:
        print("Foto del participante en portada: no disponible")

    try:
        doc.save(output_docx)
        return output_docx
    except PermissionError:
        for i in range(2, 100):
            alt_path = output_docx.with_name(f"{output_docx.stem}_v{i}{output_docx.suffix}")
            try:
                doc.save(alt_path)
                print(f"El archivo Word original estaba abierto o bloqueado. Se guardó una copia en: {alt_path}")
                return alt_path
            except PermissionError:
                continue
        raise



def add_selected_variable_definitions_section(doc, selected_vars, definition_mode, section_number=1):
    """Añade al Word la sección con las variables seleccionadas y la definición mostrada antes de la selección."""
    if definition_mode not in {"breve", "detallada"} or not selected_vars:
        return False

    doc.add_heading(f"{section_number}. Variables seleccionadas y definición mostrada antes de la selección", level=1)
    add_paragraph(doc, f"Nivel de definición elegido antes de la selección: {get_definition_label(definition_mode)}.", align=WD_ALIGN_PARAGRAPH.LEFT)

    for idx, var in enumerate(selected_vars, start=1):
        if var not in VARIABLE_META:
            continue
        meta = VARIABLE_META[var]
        doc.add_heading(f"{section_number}.{idx}. {meta['nombre']}", level=2)
        definicion = get_definition_text(var, definition_mode).strip()
        if definicion:
            add_paragraph(doc, definicion, align=WD_ALIGN_PARAGRAPH.LEFT)
        else:
            add_paragraph(doc, meta.get('nombre', var), align=WD_ALIGN_PARAGRAPH.LEFT)
    return True


def parse_fecha_usuario(fecha_str):
    fecha_str = str(fecha_str).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return pd.to_datetime(fecha_str, format=fmt).normalize()
        except Exception:
            pass
    try:
        return pd.to_datetime(fecha_str, dayfirst=True).normalize()
    except Exception:
        return None


def pedir_fecha_desde():
    while True:
        entrada = input("Fecha desde la que quieres analizar imágenes (dd/mm/aaaa o aaaa-mm-dd): ").strip()
        fecha = parse_fecha_usuario(entrada)
        if fecha is not None:
            return fecha
        print("Fecha no válida. Ejemplos correctos: 01/03/2026 o 2026-03-01")


def pedir_fecha_desde_lote():
    print("\nFecha inicial común para el lote:")
    print("Se buscarán datos desde esa fecha en todos los participantes seleccionados.")
    print("Si un participante no tiene datos válidos desde esa fecha, se omitirá y el proceso seguirá con el siguiente.")
    return pedir_fecha_desde()


def filtrar_imagenes_desde_fecha(files, fecha_desde):
    seleccionadas = []
    descartadas_sin_fecha = []
    descartadas_anteriores = []
    for f in files:
        fecha_txt, hora_txt = extract_datetime_from_filename(f.name)
        fecha_img = pd.to_datetime(fecha_txt, errors="coerce") if fecha_txt else pd.NaT
        if pd.isna(fecha_img):
            descartadas_sin_fecha.append(f)
            continue
        if fecha_img.normalize() >= fecha_desde.normalize():
            seleccionadas.append(f)
        else:
            descartadas_anteriores.append(f)
    return seleccionadas, descartadas_sin_fecha, descartadas_anteriores


def filtrar_dataframe_para_informe(df, fecha_desde, orden_dia):
    if df is None or df.empty:
        return df.copy()
    out = df.copy()
    fechas = pd.to_datetime(out["FECHA_IMAGEN"], errors="coerce")
    out = out.loc[fechas >= fecha_desde.normalize()].copy()
    if "ORDEN_DIA" in out.columns:
        out["ORDEN_DIA"] = pd.to_numeric(out["ORDEN_DIA"], errors="coerce")
        out = out.loc[out["ORDEN_DIA"] == int(orden_dia)].copy()
    out = out.sort_values([c for c in ["DATETIME_IMAGEN", "FILE"] if c in out.columns], na_position="last").reset_index(drop=True)
    return out


def get_image_files(images_dir, recursive=True):
    """Devuelve imágenes compatibles de forma robusta.

    Mejora clave: además de buscar en la carpeta exacta, permite búsqueda
    recursiva y extensiones en mayúsculas/minúsculas. Esto evita que el
    script se detenga cuando las imágenes están en subcarpetas o en carpetas
    con nombres distintos a "imagenes".
    """
    images_dir = Path(images_dir)
    if not images_dir.exists() or not images_dir.is_dir():
        return []

    valid_suffixes = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}
    iterator = images_dir.rglob("*") if recursive else images_dir.glob("*")
    files = [p for p in iterator if p.is_file() and p.suffix.lower() in valid_suffixes]
    return sorted(files, key=lambda x: str(x).lower())


def find_images_flexible(preferred_dir=None):
    """Busca imágenes en las ubicaciones habituales sin romper el flujo original."""
    candidatos = []
    for p in [
        preferred_dir,
        IMAGES_DIR,
        BASE_DIR / "imagenes",
        BASE_DIR / "Imagenes",
        BASE_DIR / "Datos_Imagenes",
        BASE_DIR / "DATOS_IMAGENES",
        BASE_DIR / "datos_imagenes",
        BASE_DIR / "Datos",
        BASE_DIR / "DATOS",
        BASE_DIR,
        SCRIPT_DIR / "imagenes",
        SCRIPT_DIR / "Imagenes",
        SCRIPT_DIR / "Datos_Imagenes",
        SCRIPT_DIR / "DATOS_IMAGENES",
        SCRIPT_DIR / "datos_imagenes",
        SCRIPT_DIR / "Datos",
        SCRIPT_DIR / "DATOS",
        SCRIPT_DIR,
        Path.cwd() / "imagenes",
        Path.cwd() / "Imagenes",
        Path.cwd() / "Datos_Imagenes",
        Path.cwd() / "DATOS_IMAGENES",
        Path.cwd() / "datos_imagenes",
        Path.cwd() / "Datos",
        Path.cwd() / "DATOS",
        Path.cwd(),
    ]:
        if p is None:
            continue
        try:
            rp = Path(p).expanduser().resolve()
        except Exception:
            rp = Path(p).expanduser().absolute()
        if rp not in candidatos:
            candidatos.append(rp)

    mejores_files = []
    mejor_dir = None
    for cand in candidatos:
        files = get_image_files(cand, recursive=True)
        if files and len(files) > len(mejores_files):
            mejores_files = files
            mejor_dir = cand

    return mejor_dir, mejores_files, candidatos


def obtener_candidatas_carpeta_imagenes():
    candidatos = []
    vistos = set()

    posibles = [
        IMAGES_DIR,
        Path.cwd() / "imagenes",
        Path.cwd(),
        BASE_DIR,
        BASE_DIR.parent / "imagenes",
    ]

    for p in posibles:
        try:
            rp = p.expanduser().resolve()
        except Exception:
            rp = p.expanduser().absolute()
        key = str(rp).lower()
        if key not in vistos:
            vistos.add(key)
            candidatos.append(rp)

    return candidatos


def pedir_carpeta_imagenes():
    candidatas = obtener_candidatas_carpeta_imagenes()
    ruta_defecto = next((p for p in candidatas if p.exists() and p.is_dir()), candidatas[0])

    while True:
        entrada = input(f"Ruta de la carpeta de imágenes [Enter = {ruta_defecto}]: ").strip()
        carpeta = Path(entrada) if entrada else ruta_defecto
        carpeta = carpeta.expanduser()
        try:
            carpeta = carpeta.resolve()
        except Exception:
            carpeta = carpeta.absolute()

        if carpeta.exists() and carpeta.is_dir():
            return carpeta

        print(f"La carpeta no existe o no es válida: {carpeta}")
        print("Sugerencias automáticas detectadas por el script:")
        for cand in candidatas:
            estado = "OK" if cand.exists() and cand.is_dir() else "No encontrada"
            print(f" - {cand} [{estado}]")
        print("Escribe la ruta completa de la carpeta donde están las imágenes de Kubios o pulsa Ctrl+C para salir.")


def pedir_tipo_informe():
    print("\nTipos de informe disponibles:")
    print("1. Informe entrenador")
    print("2. Informe participante")
    print("3. Ambos informes (entrenador y participante)")
    while True:
        entrada = input("Selecciona el tipo de informe (1, 2 o 3): ").strip()
        if entrada == "1":
            return "entrenador", "Informe entrenador"
        if entrada == "2":
            return "participante", "Informe participante"
        if entrada == "3":
            return "ambos", "Ambos informes"
        print("Entrada no válida. Introduce 1 para Informe entrenador, 2 para Informe participante o 3 para Ambos informes.")


def nombre_sugerido_desde_carpeta(base_dir=None):
    base = Path(base_dir) if base_dir is not None else BASE_DIR
    return nombre_sujeto_desde_carpeta(base)


def pedir_nombre(default_name=None):
    sugerido = default_name if default_name else nombre_sugerido_desde_carpeta()
    sugerido = str(sugerido).replace("_", " ").strip()
    return sugerido if sugerido else "Sin nombre"


def pedir_momento():
    while True:
        momento = input("Orden del registro dentro del día (1,2,3,4...): ").strip()
        if re.fullmatch(r"[1-9]\d*", momento):
            return momento
        print("Entrada no válida. Introduce solo un número entero positivo, por ejemplo: 1")


def pedir_modo_definiciones():
    print("\n¿Quieres ver una explicación de las variables antes de seleccionarlas?")
    print("1. Sí, explicación breve")
    print("2. Sí, explicación detallada")
    print("3. No, ir directamente a selección")
    while True:
        entrada = input("Selecciona una opción (1, 2 o 3): ").strip().lower()
        if entrada in {"1", "si", "sí", "breve"}:
            return "breve"
        if entrada in {"2", "detallada", "detalle"}:
            return "detallada"
        if entrada in {"3", "no", "directa", "directo"}:
            return "directa"
        print("Entrada no válida. Introduce 1, 2 o 3.")


def mostrar_definiciones_para_seleccion(definition_mode):
    if definition_mode not in {"breve", "detallada"}:
        return
    print("\nDefiniciones de las variables disponibles:")
    for i, var in enumerate(VARIABLE_META.keys(), start=1):
        meta = VARIABLE_META[var]
        print(f"\n{i}. {meta['nombre']} [{var}]")
        print(get_definition_text(var, definition_mode))


def elegir_variables():
    definition_mode = pedir_modo_definiciones()
    if definition_mode in {"breve", "detallada"}:
        mostrar_definiciones_para_seleccion(definition_mode)

    disponibles = list(VARIABLE_META.keys())
    print("\nVariables disponibles para incluir en el informe:")
    for i, var in enumerate(disponibles, start=1):
        print(f"{i}. {VARIABLE_META[var]['nombre']} [{var}]")
    print("Escribe números separados por comas, por ejemplo: 1,3,5")
    print("O escribe: todas")

    while True:
        entrada = input("Selección de variables: ").strip().lower()
        if entrada == "todas":
            return disponibles, definition_mode
        if not entrada:
            print("Debes elegir al menos una variable.")
            continue
        partes = [p.strip() for p in entrada.split(",") if p.strip()]
        if not partes:
            print("Entrada no válida.")
            continue
        if not all(re.fullmatch(r"\d+", p) for p in partes):
            print("Solo se admiten números separados por comas o la palabra 'todas'.")
            continue
        indices = []
        ok = True
        for p in partes:
            idx = int(p)
            if idx < 1 or idx > len(disponibles):
                ok = False
                break
            indices.append(idx)
        if not ok:
            print("Hay números fuera del rango disponible.")
            continue
        indices = sorted(set(indices))
        return [disponibles[i - 1] for i in indices], definition_mode


def pedir_entero_positivo(mensaje, minimo=2):
    while True:
        entrada = input(mensaje).strip()
        if not re.fullmatch(r"\d+", entrada):
            print("Entrada no válida. Introduce un número entero positivo.")
            continue
        valor = int(entrada)
        if valor < minimo:
            print(f"El valor debe ser igual o mayor que {minimo}.")
            continue
        return valor


def pedir_ventanas_analisis():
    short_window, long_window = 7, 28
    print("\nAnálisis adicional de referencia temporal:")
    print(f"Por defecto se comparará la media reciente de los últimos {short_window} registros con la media de los últimos {long_window} registros.")
    while True:
        entrada = input("¿Deseas cambiar las ventanas temporales? (s/n) [Enter = n]: ").strip().lower()
        if entrada in {"", "n", "no"}:
            return short_window, long_window
        if entrada in {"s", "si", "sí", "y", "yes"}:
            nueva_corta = pedir_entero_positivo("Introduce la ventana reciente: ", minimo=2)
            nueva_larga = pedir_entero_positivo("Introduce la ventana de referencia: ", minimo=3)
            nueva_corta, nueva_larga = safe_window_pair(nueva_corta, nueva_larga)
            if nueva_larga <= nueva_corta:
                print(f"La ventana de referencia debe ser mayor que la reciente. Se ajustará automáticamente a {nueva_larga}.")
            print(f"Se utilizarán ventanas de {nueva_corta} y {nueva_larga} registros.")
            return nueva_corta, nueva_larga
        print("Entrada no válida. Responde s/n.")




def ask_analysis_blocks():
    opciones = {
        "1": "grafico",
        "2": "interpretacion",
        "3": "recomendaciones",
        "4": "comparacion_ventanas",
        "5": "ratio",
        "6": "zscore_swc",
        "7": "tabla_interpretacion",
        "8": "lectura_emocional",
    }

    print("")
    print("Bloques de análisis disponibles:")
    print("1. Gráfico de evolución + línea de tendencia")
    print("2. Interpretación automática (basada en tendencia)")
    print("3. Recomendaciones automáticas")
    print("4. Comparación entre ventanas (7 vs 28 registros)")
    print("5. Ratio reciente/referencia (7/28)")
    print("6. Z-score individual y SWC (cambio relevante)")
    print("7. Tabla de interpretación")
    print("8. Lectura emocional por variable (no médica)")
    print("9. Todas")
    print("0. Ninguna")
    print("")
    print("Escribe números separados por comas, por ejemplo: 1,2,4")
    print("También puedes escribir 9 para todas o 0 para ninguna")

    entrada = input("Selección de bloques [Enter = 1,2,3]: ").strip().lower()

    if entrada == "":
        return ["grafico", "interpretacion", "recomendaciones"]

    if entrada in {"9", "todos", "todas", "todo", "all"}:
        return [
            "grafico",
            "interpretacion",
            "recomendaciones",
            "comparacion_ventanas",
            "ratio",
            "zscore_swc",
            "tabla_interpretacion",
            "lectura_emocional",
        ]

    if entrada in {"0", "ninguno", "ninguna", "none"}:
        return []

    seleccion = []
    for item in entrada.split(","):
        item = item.strip()
        if item in opciones:
            seleccion.append(opciones[item])

    if not seleccion:
        print("Selección no válida. Se aplicará por defecto (1,2,3).")
        return ["grafico", "interpretacion", "recomendaciones"]

    return list(dict.fromkeys(seleccion))


def ask_analysis_windows():
    short_window = 7
    long_window = 28

    print("")
    print("Para los análisis que usan ventanas, por defecto se utilizarán:")
    print(f"- ventana reciente: {short_window} registros")
    print(f"- ventana de referencia: {long_window} registros")

    cambiar = input("¿Deseas cambiar las ventanas temporales? (s/n) [Enter = n]: ").strip().lower()

    if cambiar in {"", "n", "no"}:
        return short_window, long_window

    if cambiar in {"s", "si", "sí", "y", "yes"}:
        try:
            sw = int(input("Introduce la ventana reciente: ").strip())
            lw = int(input("Introduce la ventana de referencia: ").strip())
            sw, lw = safe_window_pair(sw, lw)
            print(f"Se utilizarán ventanas de {sw} y {lw} registros.")
            return sw, lw
        except Exception:
            print("Entrada no válida. Se mantendrán las ventanas por defecto (7 y 28).")
            return short_window, long_window

    print("Entrada no válida. Se mantendrán las ventanas por defecto (7 y 28).")
    return short_window, long_window


def ask_swc_factor():
    swc_factor = 0.5
    print("")
    print("Para el SWC, por defecto se utilizará:")
    print("SWC = 0.5 * SD de la ventana de referencia")

    cambiar = input("¿Deseas cambiar el factor del SWC? (s/n) [Enter = n]: ").strip().lower()
    if cambiar in {"", "n", "no"}:
        return swc_factor

    if cambiar in {"s", "si", "sí", "y", "yes"}:
        try:
            valor = float(input("Introduce el factor del SWC (por ejemplo 0.2, 0.5, 1.0): ").strip().replace(",", "."))
            if valor > 0:
                return valor
            print("Valor no válido. Se mantendrá 0.5.")
            return swc_factor
        except Exception:
            print("Entrada no válida. Se mantendrá 0.5.")
            return swc_factor

    print("Entrada no válida. Se mantendrá 0.5.")
    return swc_factor

def ask_same_configuration_for_all(n_participantes):
    if n_participantes <= 1:
        return False
    print(f"\nHas seleccionado {n_participantes} participantes.")
    while True:
        entrada = input("¿Deseas aplicar la misma configuración a todos? (s/n) [Enter = s]: ").strip().lower()
        if entrada in {"", "s", "si", "sí", "y", "yes"}:
            return True
        if entrada in {"n", "no"}:
            return False
        print("Entrada no válida. Responde s o n.")


def ask_global_save_mode_for_existing_reports():
    print("\nGestión de informes ya existentes:")
    print("1. Reescribir siempre")
    print("2. Crear siempre nueva versión")
    while True:
        entrada = input("Selecciona una opción [Enter = 2]: ").strip()
        if entrada in {"", "2"}:
            return "nueva_version"
        if entrada == "1":
            return "reescrito"
        print("Entrada no válida. Introduce 1 o 2.")


def collect_execution_config(tipo_informe_clave, default_name=None, images_dir=None, global_save_mode=None, fecha_desde=None):
    nombre_sujeto = pedir_nombre(default_name=default_name)
    momento = pedir_momento()
    if fecha_desde is None:
        fecha_desde = pedir_fecha_desde()
    selected_vars, definition_mode = elegir_variables()
    selected_blocks = ask_analysis_blocks()

    short_window, long_window = 7, 28
    swc_factor = 0.5
    if any(x in selected_blocks for x in ["comparacion_ventanas", "ratio", "zscore_swc", "tabla_interpretacion", "lectura_emocional"]):
        short_window, long_window = ask_analysis_windows()
    if "zscore_swc" in selected_blocks:
        swc_factor = ask_swc_factor()

    return {
        "tipo_informe_clave": tipo_informe_clave,
        "nombre_sujeto": nombre_sujeto,
        "momento": momento,
        "fecha_desde": fecha_desde,
        "selected_vars": selected_vars,
        "definition_mode": definition_mode,
        "selected_blocks": selected_blocks,
        "short_window": short_window,
        "long_window": long_window,
        "swc_factor": swc_factor,
        "images_dir": images_dir,
        "global_save_mode": global_save_mode,
    }


def clone_execution_config(config, default_name=None, images_dir=None):
    out = dict(config)
    if default_name is not None:
        out["nombre_sujeto"] = pedir_nombre(default_name=default_name)
    if images_dir is not None:
        out["images_dir"] = images_dir
    return out


def ejecutar_informe(tipo_informe_clave, tipo_informe_label, default_name=None, preconfig=None):
    print("\n------------------------------------")
    print(f"REALIZANDO {tipo_informe_label.upper()}")
    print("------------------------------------")

    if preconfig is None:
        config = collect_execution_config(tipo_informe_clave, default_name=default_name, images_dir=IMAGES_DIR)
        if config.get("images_dir") is None:
            config["images_dir"] = IMAGES_DIR
    else:
        config = dict(preconfig)
        config["tipo_informe_clave"] = tipo_informe_clave
        config.setdefault("nombre_sujeto", pedir_nombre(default_name=default_name))
        if config.get("images_dir") is None:
            config["images_dir"] = IMAGES_DIR

    nombre_sujeto = config["nombre_sujeto"]
    momento = config["momento"]
    fecha_desde = config["fecha_desde"]
    selected_vars = config["selected_vars"]
    definition_mode = config["definition_mode"]
    selected_blocks = config["selected_blocks"]
    short_window = config.get("short_window", 7)
    long_window = config.get("long_window", 28)
    swc_factor = config.get("swc_factor", 0.5)
    images_dir = config.get("images_dir")
    global_save_mode = config.get("global_save_mode")
    now_exec = pd.Timestamp.now()
    fecha_informe = now_exec.strftime("%d/%m/%Y")
    hora_informe = now_exec.strftime("%H:%M:%S")

    output_docx_base, output_xlsx, report_dir_base, graph_dir_base = build_output_paths(nombre_sujeto, fecha_informe, tipo_informe_clave)

    output_docx, save_mode = decidir_guardado_informe(output_docx_base, global_mode=global_save_mode)
    # Aunque el Word se guarde como _v2, _v3, etc., mantenemos la estructura corta:
    # Participante/Word, Participante/Excel y Participante/Graficos.
    report_dir = report_dir_base
    graph_dir = graph_dir_base
    report_dir.mkdir(parents=True, exist_ok=True)
    graph_dir.mkdir(parents=True, exist_ok=True)

    logo_path = resolve_logo_path()
    foto_participante_path = resolve_participant_photo_path()

    tesseract_path = set_tesseract()
    print(f"Tesseract: {tesseract_path}")
    print(f"Carpeta del sujeto donde se buscarán imágenes: {images_dir}")
    print(f"Carpeta base de resultados: {INFORMES_DIR}")
    print(f"Carpeta del informe actual: {report_dir}")
    print(f"Excel acumulativo del sujeto: {output_xlsx}")
    print(f"Modo de definición seleccionado: {get_definition_label(definition_mode)}")
    print(f"Carpeta común de imágenes: {get_shared_images_dir() if get_shared_images_dir() else 'No encontrada'}")
    print(f"Logo de cátedra detectado: {logo_path if logo_path else 'No encontrado'}")
    print(f"Foto del participante detectada: {foto_participante_path if foto_participante_path else 'No encontrada'}")
    print(f"Bloques seleccionados: {', '.join(selected_blocks) if selected_blocks else 'ninguno'}")
    print(f"Ventanas temporales para análisis complementarios: reciente={short_window}, referencia={long_window}")
    print(f"Factor SWC: {swc_factor}")

    if Path(output_xlsx).exists():
        print("Estado del Excel: ya existe, se cargará y actualizará.")
    else:
        print("Estado del Excel: no existe todavía, se creará uno nuevo en esta ejecución.")

    images_dir, files, checked_dirs = resolver_carpeta_imagenes_participante(BASE_DIR)
    config["images_dir"] = images_dir
    print(f"Carpeta final de imágenes del participante: {images_dir}")

    if not files:
        print(f"No se encontraron imágenes compatibles para el participante: {nombre_sujeto}")
        print("Carpetas revisadas:")
        for cand in checked_dirs:
            print(f" - {cand}")
        print("Participante omitido porque no hay imágenes disponibles para procesar.")
        return False

    print(f"Imágenes detectadas para {nombre_sujeto}: {len(files)}")
    files_en_rango, files_sin_fecha, files_fuera_rango = filtrar_imagenes_desde_fecha(files, fecha_desde)
    print(f"Imágenes desde {fecha_desde.strftime('%Y-%m-%d')}: {len(files_en_rango)}")
    if files_sin_fecha:
        print(f"Imágenes omitidas por no poder extraer fecha del nombre: {len(files_sin_fecha)}")
    if files_fuera_rango:
        print(f"Imágenes omitidas por ser anteriores a la fecha indicada: {len(files_fuera_rango)}")

    required_main_cols = [
        "SESSION_ID", "FILE", "N_FILES_SESION", "FECHA_IMAGEN", "HORA_IMAGEN", "ORDEN_DIA", "DATETIME_IMAGEN",
        "READINESS", "HR", "RMSSD", "LnRMSSD", "RR_MEAN", "SDNN", "SD1", "SD2",
        "STRESS_INDEX", "BREATH_RATE", "LF_POWER", "HF_POWER", "LF_NU", "HF_NU", "LF_HF",
        "MEASUREMENT_QUALITY", "PNS_INDEX", "SNS_INDEX", "PHYSIO_AGE",
        "ESTADO_EXTRACCION", "FALTAN"
    ]
    required_debug_cols = ["FILE", "FECHA_IMAGEN", "HORA_IMAGEN", "KEY", "BUBBLE_X", "BUBBLE_Y", "BUBBLE_W", "BUBBLE_H", "OCR_RAW", "OCR_VALUE"]
    required_ocr_cols = ["FILE", "FECHA_IMAGEN", "HORA_IMAGEN", "OCR_TEXT"]
    required_metadata_cols = [
        "RUN_ID", "FECHA_EJECUCION", "NOMBRE_SUJETO", "TIPO_INFORME", "FECHA_DESDE",
        "ORDEN_REGISTRO", "MODO_DEFINICION", "MODO_DEFINICION_ETIQUETA",
        "N_VARIABLES_SELECCIONADAS", "VARIABLES_SELECCIONADAS", "N_REGISTROS_INFORME",
        "VENTANA_RECIENTE", "VENTANA_REFERENCIA"
    ]
    required_variables_cols = [
        "RUN_ID", "FECHA_EJECUCION", "NOMBRE_SUJETO", "TIPO_INFORME", "ORDEN_VARIABLE",
        "VARIABLE", "NOMBRE_VARIABLE", "MODO_DEFINICION", "MODO_DEFINICION_ETIQUETA",
        "DEFINICION_MOSTRADA", "INTERPRETACION_GENERAL"
    ]

    existing_sheets, processed_files = load_existing_sheets(output_xlsx)
    print(f"Registros históricos detectados en Excel: {len(processed_files)}")

    files_to_process = [f for f in files_en_rango if f.name not in processed_files]
    if files_to_process:
        print(f"Imágenes nuevas para procesar: {len(files_to_process)}")
    else:
        print("No hay imágenes nuevas para añadir. Se conservará el histórico ya guardado.")

    rows, debug_rows, ocr_rows = [], [], []
    for i, f in enumerate(files_to_process, start=1):
        print(f"[{i}/{len(files_to_process)}] Procesando: {f.name}")
        data, debug, text_ocr = process_image(f)
        rows.append(data)
        debug_rows.extend(debug)
        ocr_rows.append({"FILE": f.name, "FECHA_IMAGEN": data.get("FECHA_IMAGEN"), "HORA_IMAGEN": data.get("HORA_IMAGEN"), "OCR_TEXT": text_ocr})

    df_old = ensure_columns(existing_sheets.get("datos_kubios"), required_main_cols)
    df_debug_old = ensure_columns(existing_sheets.get("debug_extraccion"), required_debug_cols)
    df_ocr_old = ensure_columns(existing_sheets.get("ocr_texto"), required_ocr_cols)
    df_metadata_old = ensure_columns(existing_sheets.get("metadata_informe"), required_metadata_cols)
    df_variables_old = ensure_columns(existing_sheets.get("variables_informe"), required_variables_cols)

    df_new_raw = pd.DataFrame(rows)
    df_new = ensure_columns(consolidate_session_rows(df_new_raw), required_main_cols)
    df_debug_new = ensure_columns(pd.DataFrame(debug_rows), required_debug_cols)
    df_ocr_new = ensure_columns(pd.DataFrame(ocr_rows), required_ocr_cols)

    df_old = ensure_columns(consolidate_session_rows(df_old), required_main_cols)
    df_final = merge_dataframes(df_old, df_new, key_columns=["SESSION_ID"])
    df_debug_final = merge_dataframes(df_debug_old, df_debug_new, key_columns=["FILE", "KEY"])
    df_ocr_final = merge_dataframes(df_ocr_old, df_ocr_new, key_columns=["FILE"])

    df_final = add_derived_columns(df_final)
    df_summary = build_summary(df_final, len(df_new))

    df_informe = filtrar_dataframe_para_informe(df_final, fecha_desde, momento)
    nuevas_informe = 0
    if rows:
        df_nuevo_informe = filtrar_dataframe_para_informe(
            add_derived_columns(ensure_columns(consolidate_session_rows(df_new_raw), required_main_cols)),
            fecha_desde,
            momento
        )
        nuevas_informe = len(df_nuevo_informe)
    df_summary_informe = build_summary(df_informe, nuevas_informe)

    if df_informe is None or df_informe.empty:
        print("")
        print(f"Sin registros válidos desde {fecha_desde.strftime('%Y-%m-%d')} para el orden {momento}.")
        print("Se guardará/actualizará el Excel acumulativo si corresponde y se continuará con el siguiente participante.")
        save_excel(df_final, df_debug_final, df_ocr_final, df_summary, df_metadata_old, df_variables_old, output_xlsx)
        global_all_data_xlsx = update_global_all_data_excel(nombre_sujeto, BASE_DIR.name, df_final)
        print(f"Excel generado/actualizado: {output_xlsx}")
        print(f"Excel acumulativo global de todos los datos: {global_all_data_xlsx}")
        print("Participante omitido para informe Word por falta de datos en el periodo seleccionado.")
        return False

    run_id = pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")
    df_metadata_new = ensure_columns(
        build_report_metadata(
            run_id,
            nombre_sujeto,
            tipo_informe_label,
            fecha_informe,
            fecha_desde.strftime("%Y-%m-%d"),
            momento,
            definition_mode,
            selected_vars,
            len(df_informe),
            short_window=short_window,
            long_window=long_window,
        ),
        required_metadata_cols
    )
    df_variables_new = ensure_columns(
        build_variables_metadata(
            run_id,
            nombre_sujeto,
            tipo_informe_label,
            fecha_informe,
            definition_mode,
            selected_vars,
        ),
        required_variables_cols
    )

    df_metadata_final = merge_dataframes(df_metadata_old, df_metadata_new, key_columns=["RUN_ID"])
    df_variables_final = merge_dataframes(df_variables_old, df_variables_new, key_columns=["RUN_ID", "VARIABLE"])

    print("Guardando Excel acumulativo...")
    save_excel(df_final, df_debug_final, df_ocr_final, df_summary, df_metadata_final, df_variables_final, output_xlsx)
    global_all_data_xlsx = update_global_all_data_excel(nombre_sujeto, BASE_DIR.name, df_final)
    print("Excel guardado correctamente.")

    graph_info = generate_graphs(df_informe, momento, selected_vars, graph_dir, short_window=short_window, long_window=long_window)
    saved_docx = generate_word(
        df_informe,
        df_summary_informe,
        nombre_sujeto,
        momento,
        graph_info,
        selected_vars,
        fecha_informe,
        output_docx,
        tipo_informe_label,
        fecha_desde.strftime("%Y-%m-%d"),
        definition_mode,
        logo_path=logo_path,
        foto_participante_path=foto_participante_path,
        short_window=short_window,
        long_window=long_window,
        selected_blocks=selected_blocks,
        swc_factor=swc_factor,
    )

    summary_row = build_execution_summary_row(
        run_id=run_id,
        fecha_informe=fecha_informe,
        hora_informe=hora_informe,
        nombre_sujeto=nombre_sujeto,
        tipo_informe_label=tipo_informe_label,
        fecha_desde=fecha_desde.strftime("%Y-%m-%d"),
        momento=momento,
        definition_mode=definition_mode,
        selected_vars=selected_vars,
        selected_blocks=selected_blocks,
        short_window=short_window,
        long_window=long_window,
        swc_factor=swc_factor,
        df_informe=df_informe,
        df_final=df_final,
        n_new_rows=len(df_new),
        graph_info=graph_info,
        output_docx=saved_docx,
        output_xlsx=output_xlsx,
    )
    participant_summary_xlsx = Path(output_xlsx).parent / f"{safe_filename_text(nombre_sujeto)}_seguimiento.xlsx"
    append_execution_summary_excel(summary_row, participant_summary_xlsx)
    global_summary_xlsx = get_global_summary_excel_path()
    append_execution_summary_excel(summary_row, global_summary_xlsx)

    alarmas_xlsx = update_alarmas_excel(BASE_DIR.name, fecha_informe, hora_informe)

    print("")
    print(f"Tipo de informe generado: {tipo_informe_label}")
    print(f"Excel generado/actualizado: {output_xlsx}")
    print(f"Excel acumulativo global de todos los datos: {global_all_data_xlsx}")
    print(f"Excel incremental por participante: {participant_summary_xlsx}")
    print(f"Excel global de participantes: {global_summary_xlsx}")
    print(f"Excel de alarmas actualizado: {alarmas_xlsx}")
    print(f"Registros totales en Excel: {len(df_final)}")
    print(f"Nuevas sesiones procesadas en esta ejecución: {len(df_new)}")
    print(f"Modo de guardado del informe: {save_mode}")
    print(f"Word generado: {saved_docx}")
    print(f"Gráficos generados en: {graph_dir}")
    print(f"Registros incluidos en el informe (desde fecha y orden seleccionado): {len(df_informe)}")
    print("Proceso completado correctamente.")
    return True



def mostrar_preanalisis_seleccion_activa():
    """Muestra un resumen breve de la selección activa antes del análisis OCR."""
    if load_active_selection is None:
        return

    payload = load_active_selection(SCRIPT_DIR)
    participantes = payload.get("participants", []) if payload else []
    if not participantes:
        return

    print("\n" + "-" * 60)
    print(" SELECCIÓN ACTIVA DETECTADA")
    print("-" * 60)
    print(f"Participantes activos: {len(participantes)}")
    for idx, item in enumerate(participantes, start=1):
        nombre = item.get("name") or item.get("folder_name") or "Sin nombre"
        imagenes = item.get("images", 0)
        sesiones = item.get("sessions", 0)
        print(f"{idx:>3}. {nombre} | imágenes: {imagenes} | sesiones: {sesiones}")
    print("\nEl análisis podrá reutilizar esta selección si confirmas la opción al continuar.")

def ejecutar_analisis():
    print("====================================")
    print("KUBIOS TODO-EN-UNO ESTABLE v41")
    print("====================================")

    mostrar_preanalisis_seleccion_activa()
    carpetas_participantes = seleccionar_carpetas_participantes(SCRIPT_DIR)
    tipo_informe_clave, tipo_informe_label = pedir_tipo_informe()

    global_save_mode = ask_global_save_mode_for_existing_reports() if len(carpetas_participantes) > 1 else None
    shared_configs = {}
    fecha_desde_lote = pedir_fecha_desde_lote()

    configurar_directorios_base(carpetas_participantes[0])
    print("\nSe aplicará una configuración común a todos los participantes seleccionados.")
    print(f"Fecha inicial común: {fecha_desde_lote.strftime('%Y-%m-%d')}")
    if tipo_informe_clave == "ambos":
        print("\nConfiguración común para el informe entrenador")
        shared_configs["entrenador"] = collect_execution_config(
            "entrenador",
            default_name=nombre_sugerido_desde_carpeta(carpetas_participantes[0]),
            images_dir=IMAGES_DIR,
            global_save_mode=global_save_mode,
            fecha_desde=fecha_desde_lote,
        )
        print("\nConfiguración común para el informe participante")
        shared_configs["participante"] = collect_execution_config(
            "participante",
            default_name=nombre_sugerido_desde_carpeta(carpetas_participantes[0]),
            images_dir=IMAGES_DIR,
            global_save_mode=global_save_mode,
            fecha_desde=fecha_desde_lote,
        )
    else:
        shared_configs[tipo_informe_clave] = collect_execution_config(
            tipo_informe_clave,
            default_name=nombre_sugerido_desde_carpeta(carpetas_participantes[0]),
            images_dir=IMAGES_DIR,
            global_save_mode=global_save_mode,
            fecha_desde=fecha_desde_lote,
        )

    total = len(carpetas_participantes)
    informes_generados = 0
    participantes_sin_datos = 0
    participantes_con_error = 0
    estado_lote = []
    for idx, carpeta_participante in enumerate(carpetas_participantes, start=1):
        configurar_directorios_base(carpeta_participante)
        nombre_auto = nombre_sugerido_desde_carpeta(carpeta_participante)

        print(f"\n[{idx}/{total}] Carpeta de sujeto seleccionada: {BASE_DIR}")
        print(f"Sujeto detectado automáticamente: {nombre_auto.replace('_', ' ')}")
        print(f"Carpeta del sujeto para buscar imágenes: {IMAGES_DIR}")
        print(f"Carpeta donde se guardarán sus resultados: {INFORMES_DIR}")

        if tipo_informe_clave == "ambos":
            print("\nINICIANDO OPCIÓN 3: AMBOS INFORMES")
            cfg_ent = clone_execution_config(shared_configs["entrenador"], default_name=nombre_auto, images_dir=IMAGES_DIR)
            cfg_par = clone_execution_config(shared_configs["participante"], default_name=nombre_auto, images_dir=IMAGES_DIR)
            print("REALIZANDO PRIMERO EL INFORME DEL ENTRENADOR")
            ok_ent, error_ent = ejecutar_informe_seguro("entrenador", "Informe entrenador", nombre_auto, cfg_ent)
            if error_ent:
                print("Error generando informe entrenador para este participante:")
                print(error_ent)
            print("\nREALIZANDO AHORA EL INFORME DEL PARTICIPANTE")
            ok_par, error_par = ejecutar_informe_seguro("participante", "Informe participante", nombre_auto, cfg_par)
            if error_par:
                print("Error generando informe participante para este participante:")
                print(error_par)
            informes_generados += int(bool(ok_ent)) + int(bool(ok_par))
            if not ok_ent and not ok_par:
                if error_ent or error_par:
                    participantes_con_error += 1
                    print("\nNo se generaron informes para este participante por error. Se continúa con el siguiente.")
                else:
                    participantes_sin_datos += 1
                    print("\nNo se generaron informes para este participante por falta de datos en el periodo seleccionado.")
            else:
                print("\nProceso de informes completado para este participante.")
            estado_lote.append({
                "ORDEN": idx,
                "PARTICIPANTE": nombre_auto,
                "CARPETA_PARTICIPANTE": str(BASE_DIR),
                "TIPO_INFORME": "entrenador",
                "GENERADO": int(bool(ok_ent)),
                "ERROR": error_ent,
            })
            estado_lote.append({
                "ORDEN": idx,
                "PARTICIPANTE": nombre_auto,
                "CARPETA_PARTICIPANTE": str(BASE_DIR),
                "TIPO_INFORME": "participante",
                "GENERADO": int(bool(ok_par)),
                "ERROR": error_par,
            })
        else:
            print(f"\nREALIZANDO {tipo_informe_label.upper()}")
            cfg = clone_execution_config(shared_configs[tipo_informe_clave], default_name=nombre_auto, images_dir=IMAGES_DIR)
            ok, error = ejecutar_informe_seguro(tipo_informe_clave, tipo_informe_label, nombre_auto, cfg)
            if error:
                print("Error generando informe para este participante:")
                print(error)
            informes_generados += int(bool(ok))
            if not ok:
                if error:
                    participantes_con_error += 1
                else:
                    participantes_sin_datos += 1
            estado_lote.append({
                "ORDEN": idx,
                "PARTICIPANTE": nombre_auto,
                "CARPETA_PARTICIPANTE": str(BASE_DIR),
                "TIPO_INFORME": tipo_informe_label,
                "GENERADO": int(bool(ok)),
                "ERROR": error,
            })

        save_batch_status_excel(estado_lote)

    if total > 1:
        print(f"\nProceso completado correctamente para {total} participantes.")
        print(f"Informes generados: {informes_generados}")
        print(f"Participantes omitidos por falta de datos desde {fecha_desde_lote.strftime('%Y-%m-%d')}: {participantes_sin_datos}")
        print(f"Participantes con error: {participantes_con_error}")

    estado_lote_xlsx = save_batch_status_excel(estado_lote)
    print(f"Estado del lote: {estado_lote_xlsx}")
    total_xlsx, total_docx = build_total_results_from_global_summary(fecha_desde_lote=fecha_desde_lote)
    if total_xlsx and total_docx:
        print(f"Excel de resultados totales: {total_xlsx}")
        print(f"Word de resultados totales: {total_docx}")


def _progress_bar(porcentaje, ancho=20):
    try:
        porcentaje = int(porcentaje)
    except Exception:
        porcentaje = 0
    porcentaje = max(0, min(100, porcentaje))
    relleno = int(round((porcentaje / 100) * ancho))
    return "█" * relleno + "░" * (ancho - relleno)


def _count_files_by_suffix(root_dir, suffixes):
    root_dir = Path(root_dir)
    if not root_dir.exists():
        return 0
    suffixes = {s.lower() for s in suffixes}
    return sum(1 for p in root_dir.rglob("*") if p.is_file() and p.suffix.lower() in suffixes)


def _status_text(path):
    path = Path(path)
    return "OK" if path.exists() else "NO ENCONTRADO"


def _print_project_notes(notes):
    """Imprime notas del proyecto evitando que una tupla aparezca con paréntesis y comillas."""
    if isinstance(notes, (list, tuple)):
        for line in notes:
            print(str(line).strip())
    else:
        print(str(notes).strip())


def _count_subjects():
    datos_dir = SCRIPT_DIR / "Datos"
    if not datos_dir.exists() or not datos_dir.is_dir():
        return 0
    return len(listar_carpetas_participantes(datos_dir))


def _count_input_images():
    datos_dir = SCRIPT_DIR / "Datos"
    suffixes = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}
    return _count_files_by_suffix(datos_dir, suffixes)


def _count_processed_records_from_excels():
    resultados_dir = SCRIPT_DIR / RESULTADOS_DIR_NAME
    if not resultados_dir.exists():
        return 0

    total = 0
    for xlsx in resultados_dir.rglob("*_kubios_acumulado.xlsx"):
        try:
            df = pd.read_excel(xlsx, sheet_name="datos_kubios")
            total += len(df)
        except Exception:
            continue
    return total


def _safe_count_dirs(root_dir):
    root_dir = Path(root_dir)
    if not root_dir.exists():
        return 0
    return sum(1 for p in root_dir.rglob("*") if p.is_dir() and "__pycache__" not in p.parts and ".git" not in p.parts)


def _collect_python_project_stats(root_dir):
    root_dir = Path(root_dir)
    stats = {
        "py_files": 0,
        "lines": 0,
        "functions": 0,
        "classes": 0,
        "modules": 0,
        "folders": _safe_count_dirs(root_dir),
    }
    excluded = {".git", "__pycache__", ".venv", "venv", "env", "build", "dist"}

    py_files = []
    for py_file in root_dir.rglob("*.py"):
        if any(part in excluded for part in py_file.parts):
            continue
        py_files.append(py_file)

    stats["py_files"] = len(py_files)
    stats["modules"] = sum(1 for p in py_files if p.name != "__init__.py")

    for py_file in py_files:
        try:
            source = py_file.read_text(encoding="utf-8", errors="ignore")
            stats["lines"] += sum(1 for line in source.splitlines() if line.strip())
            tree = ast.parse(source)
            stats["functions"] += sum(isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) for node in ast.walk(tree))
            stats["classes"] += sum(isinstance(node, ast.ClassDef) for node in ast.walk(tree))
        except Exception:
            continue

    return stats


def _count_sessions_in_data(datos_dir):
    datos_dir = Path(datos_dir)
    if not datos_dir.exists():
        return 0
    session_files = {".xlsx", ".xls", ".csv", ".txt", ".fit", ".tcx"}
    return _count_files_by_suffix(datos_dir, session_files)


def _requirements_summary():
    """Resume el estado del archivo requirements.txt sin detener el programa."""
    req = SCRIPT_DIR / "requirements.txt"
    if not req.exists():
        return "NO ENCONTRADO"

    try:
        lines = [
            line.strip()
            for line in req.read_text(encoding="utf-8", errors="ignore").splitlines()
            if line.strip() and not line.strip().startswith("#")
        ]
    except Exception:
        return "NO LEGIBLE"

    return f"OK ({len(lines)} dependencias declaradas)"


def _git_status_summary():
    """Devuelve un estado simple de Git compatible con equipos sin Git instalado."""
    if not (SCRIPT_DIR / ".git").exists():
        return "NO INICIALIZADO"

    if shutil.which("git") is None:
        return "REPOSITORIO DETECTADO / GIT NO DISPONIBLE"

    try:
        result = subprocess.run(
            ["git", "status", "--short"],
            cwd=str(SCRIPT_DIR),
            capture_output=True,
            text=True,
            timeout=5,
        )
        if result.returncode != 0:
            return "REPOSITORIO DETECTADO / ESTADO NO DISPONIBLE"
        cambios = [line for line in result.stdout.splitlines() if line.strip()]
        if not cambios:
            return "OK / SIN CAMBIOS PENDIENTES"
        return f"OK / {len(cambios)} CAMBIOS PENDIENTES"
    except Exception:
        return "REPOSITORIO DETECTADO / ESTADO NO DISPONIBLE"


def _project_quality_status():
    tests_dir = SCRIPT_DIR / "tests"
    docs_dir = SCRIPT_DIR / "docs"
    modules_dir = SCRIPT_DIR / "modules"
    github_workflows = SCRIPT_DIR / ".github" / "workflows"
    return {
        "Git": (SCRIPT_DIR / ".git").exists(),
        "README": (SCRIPT_DIR / "README.md").exists(),
        "LICENSE": (SCRIPT_DIR / "LICENSE").exists() or (SCRIPT_DIR / "LICENSE.md").exists(),
        "CHANGELOG": (SCRIPT_DIR / "CHANGELOG.md").exists(),
        "Requirements": (SCRIPT_DIR / "requirements.txt").exists(),
        "Documentación docs/": docs_dir.exists() and any(docs_dir.rglob("*")),
        "Tests": tests_dir.exists() and any(tests_dir.rglob("test*.py")),
        "Módulos Python": modules_dir.exists() and any(p.suffix == ".py" for p in modules_dir.rglob("*.py")),
        "GitHub Actions": github_workflows.exists() and any(github_workflows.glob("*.yml")),
        "Módulo dashboard": (SCRIPT_DIR / "modules" / "dashboard.py").exists() or True,
    }


def _ok_symbol(value):
    return "OK" if value else "PENDIENTE"


def _calculate_dynamic_progress():
    values = []
    for _, pct in PROJECT_PROGRESS.items():
        try:
            values.append(int(pct))
        except Exception:
            continue
    if not values:
        return PROJECT_OVERALL_PROGRESS

    quality = _project_quality_status()
    quality_score = int(sum(1 for v in quality.values() if v) / max(len(quality), 1) * 100)
    return round((sum(values) / len(values) * 0.75) + (quality_score * 0.25))

def mostrar_dashboard_proyecto():
    resultados_dir = SCRIPT_DIR / RESULTADOS_DIR_NAME
    datos_dir = SCRIPT_DIR / "Datos"
    legacy_dir = SCRIPT_DIR / "legacy"
    assets_dir = SCRIPT_DIR / ASSETS_DIR_NAME

    sujetos = _count_subjects()
    imagenes_entrada = _count_input_images()
    registros_excel = _count_processed_records_from_excels()
    sesiones_detectadas = _count_sessions_in_data(datos_dir)

    excels_generados = _count_files_by_suffix(resultados_dir, {".xlsx", ".xls"})
    words_generados = _count_files_by_suffix(resultados_dir, {".docx"})
    pdf_generados = _count_files_by_suffix(resultados_dir, {".pdf"})
    graficos_generados = _count_files_by_suffix(resultados_dir, {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"})

    code_stats = _collect_python_project_stats(SCRIPT_DIR)
    quality = _project_quality_status()
    dynamic_progress = _calculate_dynamic_progress()

    print("\n" + "=" * 70)
    print(" DASHBOARD DEL PROYECTO")
    print("=" * 70)

    print(f"\nVersión actual: {PROJECT_VERSION}")
    print(f"Progreso global calculado: {dynamic_progress}%  {_progress_bar(dynamic_progress, 25)}")
    print(f"Próximo objetivo: {PROJECT_NEXT_OBJECTIVE}")
    print(f"Sistema operativo: {platform.system()} {platform.release()}")
    print(f"Python: {sys.version.split()[0]}")
    print(f"Directorio: {SCRIPT_DIR}")

    print("\n" + "-" * 70)
    print(" DATOS")
    print("-" * 70)
    print(f"Sujetos detectados:                 {sujetos}")
    print(f"Sesiones/archivos de datos:         {sesiones_detectadas}")
    print(f"Imágenes de entrada detectadas:     {imagenes_entrada}")
    print(f"Registros en Excel acumulativo:     {registros_excel}")

    print("\n" + "-" * 70)
    print(" RESULTADOS")
    print("-" * 70)
    print(f"Excels generados:                   {excels_generados}")
    print(f"Informes Word generados:            {words_generados}")
    print(f"PDF generados:                      {pdf_generados}")
    print(f"Gráficos generados:                 {graficos_generados}")

    print("\n" + "-" * 70)
    print(" CÓDIGO")
    print("-" * 70)
    print(f"Archivos Python:                    {code_stats['py_files']}")
    print(f"Módulos Python:                     {code_stats['modules']}")
    print(f"Líneas de código útiles:            {code_stats['lines']}")
    print(f"Funciones detectadas:               {code_stats['functions']}")
    print(f"Clases detectadas:                  {code_stats['classes']}")
    print(f"Carpetas del proyecto:              {code_stats['folders']}")

    print("\n" + "-" * 70)
    print(" CARPETAS Y ARCHIVOS CLAVE")
    print("-" * 70)
    print(f"Datos/:                              {_status_text(datos_dir)}")
    print(f"{RESULTADOS_DIR_NAME}/:                         {_status_text(resultados_dir)}")
    print(f"assets/:                             {_status_text(assets_dir)}")
    print(f"legacy/:                             {_status_text(legacy_dir)}")
    print(f"modules/:                            {_status_text(SCRIPT_DIR / 'modules')}")
    print(f"docs/:                               {_status_text(SCRIPT_DIR / 'docs')}")
    print(f"tests/:                              {_status_text(SCRIPT_DIR / 'tests')}")
    print(f"main.py:                             {_status_text(SCRIPT_DIR / 'main.py')}")
    print(f"config.py:                           {_status_text(SCRIPT_DIR / 'config.py')}")
    print(f"legacy/V41_kubios.py:                {_status_text(legacy_dir / 'V41_kubios.py')}")

    print("\n" + "-" * 70)
    print(" CALIDAD DEL PROYECTO")
    print("-" * 70)
    for item, ok in quality.items():
        print(f"{item:<38} {_ok_symbol(ok)}")

    print("\n" + "-" * 70)
    print(" ESTADO TÉCNICO")
    print("-" * 70)
    print(f"Git:                                  {_git_status_summary()}")
    print(f"Dependencias:                         {_requirements_summary()}")
    print(f"Ruta del proyecto:                    {SCRIPT_DIR}")
    print(f"Ejecutable Python:                    {sys.executable}")

    print("\n" + "-" * 70)
    print(" DESARROLLO")
    print("-" * 70)
    for area, porcentaje in PROJECT_PROGRESS.items():
        print(f"{area:<38} {porcentaje:>3}%  {_progress_bar(porcentaje)}")

    print("\nNota:")
    _print_project_notes(PROJECT_PROGRESS_NOTES)
    input("\nPulsa ENTER para volver al menú...")

def mostrar_porcentaje_proyecto():
    print("\n" + "=" * 60)
    print(" ESTADO DEL PROYECTO")
    print("=" * 60)
    print(f"\nVersión actual: {PROJECT_VERSION}")
    print(f"Progreso global estimado: {PROJECT_OVERALL_PROGRESS}%  {_progress_bar(PROJECT_OVERALL_PROGRESS)}\n")

    for area, porcentaje in PROJECT_PROGRESS.items():
        print(f"{area:<38} {porcentaje:>3}%  {_progress_bar(porcentaje)}")

    print("\nNota:")
    _print_project_notes(PROJECT_PROGRESS_NOTES)
    input("\nPulsa ENTER para volver al menú...")


def main():
    while True:
        print("\n" + "=" * 60)
        print(" HRV-LONGITUDINAL-ANALYZER")
        print("=" * 60)
        print("1. Ejecutar análisis Kubios OCR")
        print("2. Ver porcentaje del proyecto")
        print("3. Dashboard del proyecto")
        print("4. Gestor de participantes")
        print("5. Ver selección activa")
        print("6. Validación científica (SVF)")
        print("0. Salir")

        opcion = input("\nSelecciona una opción: ").strip()

        if opcion == "1":
            ejecutar_analisis()
            break

        if opcion == "2":
            mostrar_porcentaje_proyecto()
            continue

        if opcion == "3":
            mostrar_dashboard_proyecto()
            continue

        if opcion == "4":
            if show_participants_manager is None:
                print("El módulo de participantes no está disponible.")
                if PARTICIPANTS_IMPORT_ERROR:
                    print(f"Detalle técnico: {PARTICIPANTS_IMPORT_ERROR}")
                print("\nComprueba que exista el archivo: modules/participants.py")
                input("Pulsa ENTER para volver al menú...")
            else:
                show_participants_manager(SCRIPT_DIR / "Datos")
            continue

        if opcion == "5":
            if print_active_selection is None:
                print("El módulo de participantes no está disponible.")
            else:
                print_active_selection(SCRIPT_DIR)
            input("\nPulsa ENTER para volver al menú...")
            continue


        if opcion == "6":
            if show_validation_menu is None:
                print("El módulo de validación científica no está disponible.")
                if VALIDATION_IMPORT_ERROR:
                    print(f"Detalle técnico: {VALIDATION_IMPORT_ERROR}")
                print("\nComprueba que exista el archivo: modules/validation.py")
                input("Pulsa ENTER para volver al menú...")
            else:
                show_validation_menu(SCRIPT_DIR)
            continue

        if opcion == "0":
            print("Saliendo...")
            break

        print("Opción no válida.")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print("ERROR DURANTE LA EJECUCIÓN")
        print(str(e))
        print(traceback.format_exc())
        input("Presione una tecla para continuar . . .")
