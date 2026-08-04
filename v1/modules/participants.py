"""
Gestor de participantes para HRV-Longitudinal-Analyzer.

v1.3.5
- Selección activa persistente para conectar participantes con informes.
- Detección automática de participantes en Datos/.
- Recuento de archivos, imágenes y tablas.
- Extracción básica de fechas desde nombres de archivo.
- Detalle individual por participante.
- Exportación del listado a Excel en Resultados/Participantes/.
- Menú interactivo seguro: seleccionar, ver detalle, exportar o volver.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Sequence
import re

DATA_EXTENSIONS = {".xlsx", ".xls", ".csv", ".txt", ".png", ".jpg", ".jpeg", ".bmp", ".webp"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".webp"}
TABLE_EXTENSIONS = {".xlsx", ".xls", ".csv"}

DATE_RE = re.compile(r"(?P<date>20\d{2}[-_\.]\d{2}[-_\.]\d{2}|\d{2}[-_\.]\d{2}[-_\.]20\d{2})")
TIME_RE = re.compile(r"(?P<time>\d{2}[\.\-_:]\d{2}[\.\-_:]\d{2})")


@dataclass(frozen=True)
class Participant:
    """Información resumida de un participante detectado."""

    id: int
    name: str
    path: Path
    files: int
    images: int
    tables: int
    dates: tuple[str, ...] = field(default_factory=tuple)
    first_date: str = ""
    last_date: str = ""
    possible_sessions: int = 0
    duplicate_names: int = 0

    @property
    def display_name(self) -> str:
        return self.name.replace("_", " ").strip()


def _is_hidden(path: Path) -> bool:
    return any(part.startswith(".") or part == "__pycache__" for part in path.parts)


def _candidate_subject_dirs(data_dir: Path) -> List[Path]:
    if not data_dir.exists() or not data_dir.is_dir():
        return []
    return [child for child in sorted(data_dir.iterdir(), key=lambda p: p.name.lower()) if child.is_dir() and not _is_hidden(child)]


def _iter_valid_files(path: Path, extensions: Iterable[str] = DATA_EXTENSIONS) -> List[Path]:
    extensions = {ext.lower() for ext in extensions}
    if not path.exists():
        return []
    return [item for item in path.rglob("*") if item.is_file() and not _is_hidden(item) and item.suffix.lower() in extensions]


def _count_files(path: Path, extensions: Iterable[str]) -> int:
    return len(_iter_valid_files(path, extensions))


def _normalise_date(text: str) -> str:
    value = text.replace("_", "-").replace(".", "-")
    parts = value.split("-")
    if len(parts) != 3:
        return value
    if len(parts[0]) == 4:
        return f"{parts[0]}-{parts[1]}-{parts[2]}"
    return f"{parts[2]}-{parts[1]}-{parts[0]}"


def _extract_dates(files: Sequence[Path]) -> tuple[str, ...]:
    dates: set[str] = set()
    for file in files:
        match = DATE_RE.search(file.name)
        if match:
            dates.add(_normalise_date(match.group("date")))
    return tuple(sorted(dates))


def _duplicate_name_count(files: Sequence[Path]) -> int:
    seen: set[str] = set()
    duplicates = 0
    for file in files:
        key = file.stem.lower().strip()
        if key in seen:
            duplicates += 1
        else:
            seen.add(key)
    return duplicates


def _build_participant(idx: int, folder: Path) -> Participant:
    files = _iter_valid_files(folder, DATA_EXTENSIONS)
    dates = _extract_dates(files)
    return Participant(
        id=idx,
        name=folder.name,
        path=folder,
        files=len(files),
        images=sum(1 for p in files if p.suffix.lower() in IMAGE_EXTENSIONS),
        tables=sum(1 for p in files if p.suffix.lower() in TABLE_EXTENSIONS),
        dates=dates,
        first_date=dates[0] if dates else "",
        last_date=dates[-1] if dates else "",
        possible_sessions=len(dates),
        duplicate_names=_duplicate_name_count(files),
    )


def detect_participants(data_dir: str | Path = "Datos") -> List[Participant]:
    """Detecta participantes a partir de subcarpetas dentro de Datos/."""

    data_path = Path(data_dir).resolve()
    participants: List[Participant] = []

    for idx, folder in enumerate(_candidate_subject_dirs(data_path), start=1):
        participants.append(_build_participant(idx, folder))

    if not participants and data_path.exists():
        direct_files = _iter_valid_files(data_path, DATA_EXTENSIONS)
        if direct_files:
            dates = _extract_dates(direct_files)
            participants.append(
                Participant(
                    id=1,
                    name="Datos",
                    path=data_path,
                    files=len(direct_files),
                    images=sum(1 for p in direct_files if p.suffix.lower() in IMAGE_EXTENSIONS),
                    tables=sum(1 for p in direct_files if p.suffix.lower() in TABLE_EXTENSIONS),
                    dates=dates,
                    first_date=dates[0] if dates else "",
                    last_date=dates[-1] if dates else "",
                    possible_sessions=len(dates),
                    duplicate_names=_duplicate_name_count(direct_files),
                )
            )

    return participants


def parse_selection(selection: str, participants: Sequence[Participant]) -> List[Participant]:
    """Interpreta selecciones: 0, 1, 1,3,5 o 2-6."""

    text = (selection or "").strip()
    if not text:
        raise ValueError("Selección vacía.")
    if text == "0":
        return list(participants)

    max_id = len(participants)
    selected_ids: set[int] = set()

    for part in text.split(","):
        token = part.strip()
        if not token:
            continue
        if "-" in token:
            start_text, end_text = token.split("-", 1)
            start = int(start_text.strip())
            end = int(end_text.strip())
            if start > end:
                start, end = end, start
            selected_ids.update(range(start, end + 1))
        else:
            selected_ids.add(int(token))

    invalid = [idx for idx in selected_ids if idx < 1 or idx > max_id]
    if invalid:
        raise ValueError(f"Selección fuera de rango: {invalid}")

    return [participants[idx - 1] for idx in sorted(selected_ids)]


def print_participants_table(participants: Sequence[Participant]) -> None:
    print("\n" + "=" * 90)
    print(" GESTOR DE PARTICIPANTES")
    print("=" * 90)

    if not participants:
        print("\nNo se han detectado participantes en la carpeta Datos/.")
        return

    print("\nPuedes seleccionar:")
    print(" • Un participante      -> 1")
    print(" • Varios               -> 1,3,5")
    print(" • Intervalo            -> 5-12")
    print(" • Todos                -> 0")
    print(" • Detalle participante -> d1")
    print(" • Exportar a Excel     -> e")
    print(" • Ver selección activa -> s")
    print(" • Borrar selección     -> b")
    print(" • Volver al menú       -> ENTER")

    print("\n" + "-" * 90)
    print(f"{'ID':>3}  {'Participante':<30} {'Archivos':>8} {'Imágenes':>9} {'Tablas':>7} {'Sesiones':>8} {'Fechas':<23}")
    print("-" * 90)
    for p in participants:
        rango = ""
        if p.first_date and p.last_date:
            rango = p.first_date if p.first_date == p.last_date else f"{p.first_date} a {p.last_date}"
        print(f"{p.id:>3}  {p.display_name:<30} {p.files:>8} {p.images:>9} {p.tables:>7} {p.possible_sessions:>8} {rango:<23}")


def print_participant_detail(participant: Participant) -> None:
    print("\n" + "=" * 70)
    print(f" DETALLE: {participant.display_name}")
    print("=" * 70)
    print(f"Ruta:                       {participant.path}")
    print(f"Archivos totales:           {participant.files}")
    print(f"Imágenes:                   {participant.images}")
    print(f"Tablas/Excel/CSV:           {participant.tables}")
    print(f"Sesiones posibles por fecha:{participant.possible_sessions:>6}")
    print(f"Primera fecha detectada:    {participant.first_date or 'No detectada'}")
    print(f"Última fecha detectada:     {participant.last_date or 'No detectada'}")
    print(f"Nombres duplicados:         {participant.duplicate_names}")

    if participant.dates:
        print("\nFechas detectadas:")
        line = ""
        for idx, date in enumerate(participant.dates, start=1):
            line += f"{date}  "
            if idx % 5 == 0:
                print("  " + line.strip())
                line = ""
        if line:
            print("  " + line.strip())
    else:
        print("\nNo se han detectado fechas en los nombres de archivo.")


def export_participants_excel(participants: Sequence[Participant], project_root: str | Path | None = None) -> Path:
    """Exporta el listado a Excel. Requiere openpyxl, incluido habitualmente con pandas."""

    root = Path(project_root).resolve() if project_root else Path.cwd().resolve()
    output_dir = root / "Resultados" / "Participantes"
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"listado_participantes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        from openpyxl.utils import get_column_letter
    except Exception as exc:
        raise RuntimeError("No se pudo importar openpyxl. Instala la dependencia con: pip install openpyxl") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "Participantes"

    headers = [
        "ID",
        "Participante",
        "Ruta",
        "Archivos",
        "Imagenes",
        "Tablas",
        "Sesiones_posibles",
        "Primera_fecha",
        "Ultima_fecha",
        "Duplicados_nombre",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for p in participants:
        ws.append([
            p.id,
            p.display_name,
            str(p.path),
            p.files,
            p.images,
            p.tables,
            p.possible_sessions,
            p.first_date,
            p.last_date,
            p.duplicate_names,
        ])

    ws_dates = wb.create_sheet("Fechas")
    ws_dates.append(["ID", "Participante", "Fecha"])
    for cell in ws_dates[1]:
        cell.font = Font(bold=True)
    for p in participants:
        for date in p.dates:
            ws_dates.append([p.id, p.display_name, date])

    for sheet in (ws, ws_dates):
        for column_cells in sheet.columns:
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 12), 60)

    wb.save(output_path)
    return output_path


def _find_participant_by_command(command: str, participants: Sequence[Participant]) -> Participant:
    number_text = command[1:].strip()
    if not number_text:
        raise ValueError("Indica el número de participante. Ejemplo: d1")
    idx = int(number_text)
    if idx < 1 or idx > len(participants):
        raise ValueError("Participante fuera de rango.")
    return participants[idx - 1]



def active_selection_path(project_root: str | Path | None = None) -> Path:
    """Ruta del archivo JSON donde se guarda la selección activa."""
    root = Path(project_root).resolve() if project_root else Path.cwd().resolve()
    output_dir = root / "Resultados" / "Participantes"
    output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir / "seleccion_activa_participantes.json"


def save_active_selection(participants: Sequence[Participant], project_root: str | Path | None = None) -> Path:
    """Guarda la selección activa para reutilizarla desde el análisis/informes."""
    import json

    path = active_selection_path(project_root)
    payload = {
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "count": len(participants),
        "participants": [
            {
                "id": p.id,
                "name": p.display_name,
                "folder_name": p.name,
                "path": str(p.path),
                "files": p.files,
                "images": p.images,
                "tables": p.tables,
                "sessions": p.possible_sessions,
                "first_date": p.first_date,
                "last_date": p.last_date,
            }
            for p in participants
        ],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def load_active_selection(project_root: str | Path | None = None) -> dict | None:
    """Carga la selección activa si existe."""
    import json

    path = active_selection_path(project_root)
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


def clear_active_selection(project_root: str | Path | None = None) -> bool:
    """Elimina la selección activa guardada."""
    path = active_selection_path(project_root)
    if path.exists():
        path.unlink()
        return True
    return False


def selected_paths_from_active_selection(project_root: str | Path | None = None) -> List[Path]:
    """Devuelve rutas existentes de la selección activa."""
    payload = load_active_selection(project_root)
    if not payload:
        return []
    paths: List[Path] = []
    for item in payload.get("participants", []):
        candidate = Path(str(item.get("path", "")))
        if candidate.exists() and candidate.is_dir():
            paths.append(candidate.resolve())
    return paths


def print_active_selection(project_root: str | Path | None = None) -> None:
    """Muestra en consola la selección activa."""
    payload = load_active_selection(project_root)
    print("\n" + "=" * 70)
    print(" SELECCIÓN ACTIVA DE PARTICIPANTES")
    print("=" * 70)
    if not payload or not payload.get("participants"):
        print("\nNo hay selección activa guardada.")
        return
    print(f"\nFecha de selección: {payload.get('created_at', 'No disponible')}")
    print(f"Participantes:      {payload.get('count', len(payload.get('participants', [])))}")
    print("\n" + "-" * 70)
    for idx, item in enumerate(payload.get("participants", []), start=1):
        print(f"{idx:>3}. {item.get('name', '')} | imágenes: {item.get('images', 0)} | sesiones: {item.get('sessions', 0)}")
        print(f"     {item.get('path', '')}")


def show_participants_manager(data_dir: str | Path = "Datos") -> List[Participant]:
    """Muestra el gestor interactivo y devuelve participantes seleccionados."""

    data_path = Path(data_dir).resolve()
    project_root = data_path.parent
    participants = detect_participants(data_path)

    while True:
        print_participants_table(participants)

        if not participants:
            input("\nPulsa ENTER para volver al menú...")
            return []

        selection = input("\nSelección: ").strip()
        command = selection.lower()

        if command == "":
            return []

        if command in {"s", "seleccion", "selección", "activa"}:
            print_active_selection(project_root)
            input("\nPulsa ENTER para continuar...")
            continue

        if command in {"b", "borrar", "limpiar"}:
            removed = clear_active_selection(project_root)
            print("\nSelección activa eliminada." if removed else "\nNo había selección activa guardada.")
            input("\nPulsa ENTER para continuar...")
            continue

        if command in {"e", "excel", "exportar"}:
            try:
                output = export_participants_excel(participants, project_root)
                print(f"\nExcel generado correctamente:\n{output}")
            except Exception as exc:
                print(f"\nNo se pudo generar el Excel: {exc}")
            input("\nPulsa ENTER para continuar...")
            continue

        if command.startswith("d"):
            try:
                participant = _find_participant_by_command(command, participants)
                print_participant_detail(participant)
            except Exception as exc:
                print(f"\nNo se pudo mostrar el detalle: {exc}")
            input("\nPulsa ENTER para continuar...")
            continue

        try:
            selected = parse_selection(selection, participants)
        except Exception as exc:
            print(f"Selección no válida: {exc}")
            input("\nPulsa ENTER para continuar...")
            continue

        print("\nParticipantes seleccionados:")
        for p in selected:
            print(f" - {p.display_name} ({p.files} archivos, {p.possible_sessions} sesiones posibles)")

        active_path = save_active_selection(selected, project_root)
        print(f"\nSelección activa guardada en:\n{active_path}")
        print("Esta selección podrá reutilizarse al ejecutar informes/análisis.")

        input("\nPulsa ENTER para volver al menú...")
        return selected


def participants_summary(data_dir: str | Path = "Datos") -> dict:
    participants = detect_participants(data_dir)
    return {
        "participants": len(participants),
        "files": sum(p.files for p in participants),
        "images": sum(p.images for p in participants),
        "tables": sum(p.tables for p in participants),
        "sessions": sum(p.possible_sessions for p in participants),
    }
