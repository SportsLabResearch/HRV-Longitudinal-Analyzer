# -*- coding: utf-8 -*-
"""
Auditoría interna del proyecto HRV-Longitudinal-Analyzer.
Versión: v1.4.2-a

Genera:
- Resumen en pantalla.
- Carpeta Resultados_Auditoria/YYYY-MM-DD_HH-MM-SS/
- Informe Word (.docx).
- Informe Excel (.xlsx).
- Resumen TXT.

No modifica datos, análisis OCR ni resultados previos.
"""

from __future__ import annotations

import subprocess
import sys
import zipfile
from pathlib import Path
from datetime import datetime
from xml.sax.saxutils import escape

VERSION_AUDITORIA = "v1.4.2-a"
NOMBRE_PROYECTO = "HRV-Longitudinal-Analyzer"


def _root() -> Path:
    return Path(__file__).resolve().parents[1]


def _run_git(args: list[str]) -> str:
    try:
        result = subprocess.run(["git"] + args, cwd=_root(), capture_output=True, text=True, timeout=6)
        if result.returncode == 0:
            return result.stdout.strip()
        return "No disponible"
    except Exception:
        return "No disponible"


def _exists(path: str) -> bool:
    return (_root() / path).exists()


def _count_files(folder: str, extensions: tuple[str, ...] | None = None) -> int:
    p = _root() / folder
    if not p.exists() or not p.is_dir():
        return 0
    files = [x for x in p.rglob("*") if x.is_file()]
    if extensions:
        files = [x for x in files if x.suffix.lower() in extensions]
    return len(files)


def _check_python_file(path: str) -> tuple[bool, str]:
    p = _root() / path
    if not p.exists():
        return False, "No existe"
    try:
        subprocess.run(
            [sys.executable, "-m", "py_compile", str(p)],
            cwd=_root(), capture_output=True, text=True, timeout=12, check=True,
        )
        return True, "Compila correctamente"
    except subprocess.CalledProcessError as exc:
        msg = exc.stderr.strip().splitlines()[-1] if exc.stderr else "Error de compilación"
        return False, msg
    except Exception as exc:
        return False, f"No se pudo comprobar: {exc}"


def _score(ok: bool, peso: int, obs: str) -> dict:
    return {"ok": ok, "peso": peso, "puntos": peso if ok else 0, "obs": obs}


def auditar_proyecto() -> dict:
    root = _root()
    main_ok, main_obs = _check_python_file("main.py")
    config_ok, config_obs = _check_python_file("config.py") if _exists("config.py") else (False, "No existe")
    validation_ok, validation_obs = _check_python_file("modules/validation.py") if _exists("modules/validation.py") else (False, "No existe")
    audit_ok, audit_obs = _check_python_file("modules/auditoria.py") if _exists("modules/auditoria.py") else (False, "No existe")

    tests_py = _count_files("tests", (".py",))
    docs_files = _count_files("docs")
    validation_files = _count_files("Validation")
    release_files = len(list(root.glob("RELEASE*.md")))

    checks = [
        ("main.py", _score(main_ok, 15, main_obs)),
        ("config.py", _score(config_ok, 7, config_obs)),
        ("modules/", _score(_exists("modules"), 7, "Carpeta de módulos")),
        ("modules/validation.py", _score(validation_ok, 9, validation_obs)),
        ("modules/auditoria.py", _score(audit_ok, 9, audit_obs)),
        ("Validation/", _score(_exists("Validation"), 8, f"{validation_files} archivo(s) detectado(s)")),
        ("README.md", _score(_exists("README.md"), 7, "Documento principal")),
        ("CHANGELOG.md", _score(_exists("CHANGELOG.md"), 7, "Historial de cambios")),
        ("LICENSE", _score(_exists("LICENSE") or _exists("LICENSE.md"), 5, "Licencia")),
        ("RELEASE", _score(release_files > 0, 7, f"{release_files} archivo(s) RELEASE detectado(s)")),
        ("docs/", _score(_exists("docs") and docs_files > 0, 7, f"{docs_files} archivo(s) detectado(s)")),
        ("tests/", _score(_exists("tests") and tests_py > 0, 12, f"{tests_py} test(s) Python detectado(s)")),
    ]

    total = sum(c[1]["peso"] for c in checks)
    puntos = sum(c[1]["puntos"] for c in checks)
    porcentaje = round((puntos / total) * 100) if total else 0

    if porcentaje >= 90:
        conclusion = "LISTO PARA REVISIÓN FINAL DE VERSIÓN ESTABLE."
    elif porcentaje >= 75:
        conclusion = "AVANZADO, PERO NO CERRAR TODAVÍA COMO ESTABLE."
    else:
        conclusion = "NO LISTO PARA CIERRE ESTABLE."

    status = _run_git(["status", "--short"])
    return {
        "proyecto": NOMBRE_PROYECTO,
        "version": VERSION_AUDITORIA,
        "fecha": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "ruta": str(root),
        "rama": _run_git(["branch", "--show-current"]) or "No disponible",
        "estado_git": status if status else "Limpio",
        "checks": checks,
        "porcentaje": porcentaje,
        "conclusion": conclusion,
    }


def _lineas(data: dict) -> list[str]:
    lineas = [
        f"Proyecto: {data['proyecto']}",
        f"Versión auditoría: {data['version']}",
        f"Fecha: {data['fecha']}",
        f"Ruta: {data['ruta']}",
        f"Rama Git: {data['rama']}",
        "",
        "RESULTADO GLOBAL",
        f"Estado general: {data['porcentaje']} %",
        f"Conclusión: {data['conclusion']}",
        "",
        "COMPROBACIONES",
    ]
    for nombre, item in data["checks"]:
        estado = "OK" if item["ok"] else "PENDIENTE"
        lineas.append(f"- {nombre}: {estado} | {item['puntos']}/{item['peso']} | {item['obs']}")
    lineas += ["", "ESTADO GIT", data["estado_git"]]
    return lineas


def _crear_docx(path: Path, data: dict) -> None:
    try:
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        doc.add_heading("Auditoría del Proyecto", level=1)
        doc.add_paragraph(f"Proyecto: {data['proyecto']}")
        doc.add_paragraph(f"Versión auditoría: {data['version']}")
        doc.add_paragraph(f"Fecha: {data['fecha']}")
        doc.add_paragraph(f"Rama Git: {data['rama']}")
        doc.add_heading("Resultado global", level=2)
        p = doc.add_paragraph()
        r = p.add_run(f"Estado general: {data['porcentaje']} %")
        r.bold = True
        r.font.size = Pt(12)
        doc.add_paragraph(data["conclusion"])
        doc.add_heading("Comprobaciones", level=2)
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Bloque"
        hdr[1].text = "Estado"
        hdr[2].text = "Puntos"
        hdr[3].text = "Peso"
        hdr[4].text = "Observación"
        for nombre, item in data["checks"]:
            row = table.add_row().cells
            row[0].text = nombre
            row[1].text = "OK" if item["ok"] else "PENDIENTE"
            row[2].text = str(item["puntos"])
            row[3].text = str(item["peso"])
            row[4].text = item["obs"]
        doc.add_heading("Estado Git", level=2)
        doc.add_paragraph(data["estado_git"])
        doc.save(path)
    except Exception:
        _crear_docx_minimo(path, "AUDITORÍA DEL PROYECTO", _lineas(data))


def _crear_docx_minimo(path: Path, titulo: str, lineas: list[str]) -> None:
    paragraphs = []
    for line in [titulo, ""] + lineas:
        paragraphs.append(f"<w:p><w:r><w:t xml:space='preserve'>{escape(str(line))}</w:t></w:r></w:p>")
    document_xml = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>{''.join(paragraphs)}<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr></w:body></w:document>'''
    content_types = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'''
    rels = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'''
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document_xml)
        z.writestr("word/_rels/document.xml.rels", "<?xml version='1.0' encoding='UTF-8'?><Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'/>")


def _xlsx_col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _sheet_xml(rows: list[list[str]]) -> str:
    xml_rows = []
    for r_idx, row in enumerate(rows, start=1):
        cells = []
        for c_idx, value in enumerate(row, start=1):
            ref = f"{_xlsx_col(c_idx)}{r_idx}"
            cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{escape(str(value))}</t></is></c>')
        xml_rows.append(f'<row r="{r_idx}">{"".join(cells)}</row>')
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>{''.join(xml_rows)}</sheetData></worksheet>'''


def _crear_xlsx(path: Path, data: dict) -> None:
    rows = [["Bloque", "Estado", "Puntos", "Peso", "Observación"]]
    for nombre, item in data["checks"]:
        rows.append([nombre, "OK" if item["ok"] else "PENDIENTE", str(item["puntos"]), str(item["peso"]), item["obs"]])
    rows += [[""], ["Estado general", f"{data['porcentaje']} %"], ["Conclusión", data["conclusion"]], ["Rama Git", data["rama"]]]

    content_types = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>'''
    rels = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>'''
    wb = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Auditoria" sheetId="1" r:id="rId1"/></sheets></workbook>'''
    wb_rels = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>'''
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("xl/workbook.xml", wb)
        z.writestr("xl/_rels/workbook.xml.rels", wb_rels)
        z.writestr("xl/worksheets/sheet1.xml", _sheet_xml(rows))


def guardar_resultados(data: dict) -> Path:
    stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    out_dir = _root() / "Resultados_Auditoria" / stamp
    out_dir.mkdir(parents=True, exist_ok=True)
    lineas = _lineas(data)
    (out_dir / "Resumen_Auditoria.txt").write_text("\n".join(lineas), encoding="utf-8")
    _crear_docx(out_dir / f"Auditoria_Proyecto_{data['version']}.docx", data)
    _crear_xlsx(out_dir / f"Auditoria_Proyecto_{data['version']}.xlsx", data)
    return out_dir


def imprimir_resumen(data: dict, out_dir: Path) -> None:
    print("\n" + "=" * 60)
    print("AUDITORÍA DEL PROYECTO")
    print("=" * 60)
    print(f"Proyecto : {data['proyecto']}")
    print(f"Versión  : {data['version']}")
    print(f"Rama Git : {data['rama']}")
    print(f"Estado   : {data['porcentaje']} %")
    print("-" * 60)
    for nombre, item in data["checks"]:
        marca = "OK" if item["ok"] else "PENDIENTE"
        print(f"{marca:10} {nombre:24} {item['puntos']:>2}/{item['peso']:<2}  {item['obs']}")
    print("-" * 60)
    print(f"Conclusión: {data['conclusion']}")
    print(f"Resultados: {out_dir}")
    print("=" * 60 + "\n")


def ejecutar_auditoria_proyecto() -> None:
    data = auditar_proyecto()
    out_dir = guardar_resultados(data)
    imprimir_resumen(data, out_dir)
    input("Pulsa ENTER para volver al menú...")


if __name__ == "__main__":
    ejecutar_auditoria_proyecto()
